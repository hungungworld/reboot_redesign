"""Excel 내보내기 도우미 — 설정(작업 로그)과 인사이트(데이터 받기)가 함께 씀.

- '='로 시작하는 문자열을 수식이 아니라 글자로 저장 (수식 주입 막기)
- 머리글 굵게 · 첫 줄 고정 · 열 너비 대략 맞춤
"""
import io
from datetime import timezone

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

HEAD_FILL = PatternFill("solid", fgColor="F1F7E1")
BAD_SHEET_CHARS = set('[]:*?/\\')


def new_workbook():
    wb = Workbook()
    wb.remove(wb.active)
    return wb


def sheet_title(name):
    s = "".join("_" if c in BAD_SHEET_CHARS else c for c in str(name)).strip() or "Sheet"
    return s[:31]


def _cell_value(v):
    if v is None:
        return None
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return v
    if isinstance(v, bool):
        return "예" if v else "아니요"
    s = str(v)
    return s[:32000]  # 셀 한도 32,767자


def add_sheet(wb, title, header, rows, widths=None):
    ws = wb.create_sheet(sheet_title(title))
    ws.append(list(header))
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = HEAD_FILL
    for r in rows:
        ws.append([_cell_value(v) for v in r])
        for c in ws[ws.max_row]:
            if c.data_type == "f":  # '='로 시작하는 글자 → 수식이 아닌 글자로
                c.data_type = "s"
    ws.freeze_panes = "A2"
    for i, _ in enumerate(header, start=1):
        letter = ws.cell(row=1, column=i).column_letter
        w = (widths or {}).get(i) if widths else None
        if w is None:
            longest = max((len(str(ws.cell(row=r, column=i).value or "")) for r in range(1, min(ws.max_row, 200) + 1)),
                          default=8)
            w = min(60, max(8, longest + 2))
        ws.column_dimensions[letter].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            if isinstance(c.value, str) and len(c.value) > 60:
                c.alignment = Alignment(wrap_text=True, vertical="top")
    return ws


def to_bytes(wb):
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def local_text(dt, tz):
    """DB의 UTC 시각 → 현지 'YYYY-MM-DD HH:MM'."""
    if not dt:
        return ""
    return dt.replace(tzinfo=timezone.utc).astimezone(tz).strftime("%Y-%m-%d %H:%M")
