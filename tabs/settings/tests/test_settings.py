"""설정 파트 API — 수집 설정 · 연결(키 입력·연결 테스트) · 작업 로그 Excel (인터넷 · 실제 키 없이)."""
import io
import json
import smtplib

import pytest
import requests

from tabs.settings import connections


class FakeRes:
    def __init__(self, status=200, data=None, text=""):
        self.status_code = status
        self._data = data
        self.text = text or (json.dumps(data) if data is not None else "")

    def json(self):
        if self._data is None:
            raise ValueError("no json")
        return self._data


@pytest.fixture()
def calls(monkeypatch):
    """connections._get을 가짜로: routes[url 일부] = FakeRes 또는 함수."""
    seen = []
    routes = {}

    def fake_get(url, params=None, headers=None):
        seen.append({"url": url, "params": params or {}, "headers": headers or {}})
        for part, res in routes.items():
            if part in url:
                return res(url, params, headers) if callable(res) else res
        raise AssertionError(f"예상하지 못한 호출: {url}")

    monkeypatch.setattr(connections, "_get", fake_get)
    return {"seen": seen, "routes": routes}


# ───────── 로그인 · 화면 ─────────

def test_settings_api_requires_login(client):
    assert client.get("/api/settings/collect").status_code == 401
    assert client.get("/api/settings/connections").status_code == 401


def test_settings_page_loads_part_files(auth_client):
    html = auth_client.get("/dashboard/settings").get_data(as_text=True)
    assert "/parts/settings/settings.js" in html and "/parts/settings/settings.css" in html


# ───────── 수집 설정 ─────────

def test_collect_defaults(auth_client):
    data = auth_client.get("/api/settings/collect").get_json()
    cfg = data["config"]
    assert cfg["amazon_queries"] == ["hangover", "korean hangover"]
    assert set(cfg["search_sets"]) == {"A", "B", "K"} and cfg["video_language"] == "en"
    assert data["updatedAt"] is None


def test_collect_save_and_accessor(app, auth_client):
    cfg = auth_client.get("/api/settings/collect").get_json()["config"]
    cfg["amazon_queries"] = ["hangover", " Hangover ", "hangover jelly"]   # 중복은 하나로
    cfg["tracked_asins"] = [{"asin": "b0abcdefgh", "label": "자사 젤리", "own": True}]
    cfg["min_views"] = "25,000"
    cfg["video_language"] = "ko_en"
    res = auth_client.put("/api/settings/collect", json={"config": cfg})
    assert res.status_code == 200, res.get_json()
    out = res.get_json()
    assert out["config"]["amazon_queries"] == ["hangover", "hangover jelly"]
    assert out["config"]["tracked_asins"] == [{"asin": "B0ABCDEFGH", "label": "자사 젤리", "own": True}]
    assert out["config"]["min_views"] == 25000
    assert "amazon_queries" in out["changed"] and out["updatedBy"] == "테스터"
    with app.app_context():
        from tabs.settings.config import get_collect_config
        got = get_collect_config()
        assert got["video_language"] == "ko_en" and got["tracked_asins"][0]["own"] is True
    log = auth_client.get("/api/worklog?part=settings").get_json()["items"]
    assert any(i["action"] == "수집 설정 변경" for i in log)


@pytest.mark.parametrize("patch, word", [
    ({"tracked_asins": [{"asin": "B0SHORT"}]}, "ASIN"),
    ({"weee_urls": ["http://example.com/x"]}, "Weee"),
    ({"video_language": "fr"}, "영상 언어"),
    ({"min_views": "많이"}, "조회수"),
    ({"amazon_queries": ["q"] * 1 + [f"q{i}" for i in range(16)]}, "Amazon 검색어"),
])
def test_collect_validation(auth_client, patch, word):
    cfg = auth_client.get("/api/settings/collect").get_json()["config"]
    cfg.update(patch)
    res = auth_client.put("/api/settings/collect", json={"config": cfg})
    assert res.status_code == 400 and word in res.get_json()["error"]


# ───────── 연결 ─────────

def _item(client, name):
    return next(i for i in client.get("/api/settings/connections").get_json()["items"] if i["name"] == name)


def test_connections_list_never_returns_values(auth_client, monkeypatch):
    monkeypatch.setenv("SERPAPI_KEY", "serp-secret-value-9876")
    res = auth_client.get("/api/settings/connections")
    body = res.get_data(as_text=True)
    assert res.status_code == 200 and "serp-secret-value" not in body
    data = res.get_json()
    assert data["isAdmin"] is True and "redirectUri" in data["setup"]
    assert data["setup"]["redirectUri"].endswith("/api/cs/gmail/callback")
    serp = next(i for i in data["items"] if i["name"] == "SERPAPI_KEY")
    assert serp["connected"] and serp["source"] == "env" and serp["masked"] == "••••9876"
    assert serp["status"] == "확인 전" and serp["testable"] is True
    yt = next(i for i in data["items"] if i["name"] == "YOUTUBE_API_KEY")
    assert yt["status"] == "미연결"
    shopify = next(i for i in data["items"] if i["name"] == "SHOPIFY_TOKEN")
    assert shopify["testable"] is False


def test_key_entry_is_admin_only(client, login):
    login(email="member@reboot.test")
    res = client.put("/api/settings/connections/SERPAPI_KEY", json={"value": "member-tries-this-key"})
    assert res.status_code == 403
    data = client.get("/api/settings/connections").get_json()
    assert data["isAdmin"] is False


def test_key_entry_saves_masked_and_logs_without_value(auth_client, monkeypatch):
    monkeypatch.setenv("SERPAPI_KEY", "env-value-1111")
    res = auth_client.put("/api/settings/connections/SERPAPI_KEY", json={"value": "  screen-value-2222  "})
    assert res.status_code == 200
    item = res.get_json()
    assert item["source"] == "screen" and item["masked"] == "••••2222"
    assert "screen-value" not in res.get_data(as_text=True)
    logs = json.dumps(auth_client.get("/api/worklog?part=settings").get_json(), ensure_ascii=False)
    assert "screen-value-2222" not in logs and "연결 값 저장: SerpApi" in logs
    # 빈 값 → 화면 입력 삭제 → .env 값으로
    item = auth_client.put("/api/settings/connections/SERPAPI_KEY", json={"value": ""}).get_json()
    assert item["source"] == "env" and item["masked"] == "••••1111"


def test_key_entry_validation(auth_client):
    assert auth_client.put("/api/settings/connections/NOPE", json={"value": "abcdefghij"}).status_code == 404
    res = auth_client.put("/api/settings/connections/OPENAI_API_KEY", json={"value": "sk-abc def ghi jkl"})
    assert res.status_code == 400 and "공백" in res.get_json()["error"]
    assert auth_client.put("/api/settings/connections/OPENAI_API_KEY", json={"value": "short"}).status_code == 400
    # 설정값(비밀 아님)은 공백 허용 · 값이 그대로 보임
    item = auth_client.put("/api/settings/connections/GMAIL_QUERY", json={"value": "to:cs@example.com newer_than:30d"}).get_json()
    assert item["masked"] == "to:cs@example.com newer_than:30d"


def test_connection_test_serpapi_ok_then_stale_after_key_change(auth_client, monkeypatch, calls):
    monkeypatch.setenv("SERPAPI_KEY", "serp-test-key-1234")
    calls["routes"]["serpapi.com/account.json"] = FakeRes(200, {
        "plan_name": "Free Plan", "searches_per_month": 250, "total_searches_left": 231, "this_month_usage": 19})
    res = auth_client.post("/api/settings/connections/SERPAPI_KEY/test")
    out = res.get_json()
    assert res.status_code == 200 and out["status"] == "정상"
    assert "231" in out["message"] and out["detail"]["searchesLeft"] == 231
    assert calls["seen"][0]["params"]["api_key"] == "serp-test-key-1234"
    assert "serp-test-key" not in res.get_data(as_text=True)
    assert _item(auth_client, "SERPAPI_KEY")["status"] == "정상"
    # 키가 바뀌면 결과는 '확인 전'으로
    monkeypatch.setenv("SERPAPI_KEY", "serp-test-key-5678")
    assert _item(auth_client, "SERPAPI_KEY")["status"] == "확인 전"


def test_connection_test_serpapi_bad_key(auth_client, monkeypatch, calls):
    monkeypatch.setenv("SERPAPI_KEY", "serp-bad-key-0000")
    calls["routes"]["serpapi.com"] = FakeRes(401, {"error": "Invalid API key."})
    out = auth_client.post("/api/settings/connections/SERPAPI_KEY/test").get_json()
    assert out["status"] == "오류" and "거절" in out["message"]
    assert _item(auth_client, "SERPAPI_KEY")["status"] == "오류"
    log = auth_client.get("/api/worklog?part=settings").get_json()["items"][0]
    assert log["status"] == "실패" and "SerpApi" in log["action"]


def test_connection_test_openai_and_anthropic(auth_client, monkeypatch, calls):
    monkeypatch.setenv("OPENAI_API_KEY", "sk-openai-test-aaaa")
    monkeypatch.setenv("ANTHROPIC_API_KEY", "sk-ant-test-bbbb")
    calls["routes"]["api.openai.com/v1/models"] = FakeRes(200, {"data": [{"id": "gpt-4o-mini"}, {"id": "gpt-4o"}]})
    calls["routes"]["api.anthropic.com/v1/models"] = FakeRes(200, {"data": [{"id": "claude-x"}]})
    out = auth_client.post("/api/settings/connections/OPENAI_API_KEY/test").get_json()
    assert out["status"] == "정상" and out["detail"]["modelAvailable"] is True
    assert calls["seen"][-1]["headers"]["Authorization"] == "Bearer sk-openai-test-aaaa"
    out = auth_client.post("/api/settings/connections/ANTHROPIC_API_KEY/test").get_json()
    assert out["status"] == "정상"
    h = calls["seen"][-1]["headers"]
    assert h["x-api-key"] == "sk-ant-test-bbbb" and h["anthropic-version"] == "2023-06-01"
    calls["routes"]["api.openai.com/v1/models"] = FakeRes(401, {"error": {"message": "bad"}})
    out = auth_client.post("/api/settings/connections/OPENAI_API_KEY/test").get_json()
    assert out["status"] == "오류" and "401" in out["message"]


def test_connection_test_youtube_quota_and_ok(auth_client, monkeypatch, calls):
    monkeypatch.setenv("YOUTUBE_API_KEY", "AIza-test-youtube-cccc")
    calls["routes"]["youtube/v3/videos"] = FakeRes(403, {"error": {"errors": [{"reason": "quotaExceeded"}]}})
    out = auth_client.post("/api/settings/connections/YOUTUBE_API_KEY/test").get_json()
    assert out["status"] == "오류" and "사용량" in out["message"]
    calls["routes"]["youtube/v3/videos"] = FakeRes(200, {"items": [{"id": "abc"}]})
    out = auth_client.post("/api/settings/connections/YOUTUBE_API_KEY/test").get_json()
    assert out["status"] == "정상"
    p = calls["seen"][-1]["params"]
    assert p["chart"] == "mostPopular" and p["maxResults"] == 1 and p["part"] == "id"


def test_connection_test_firecrawl_and_kexim(auth_client, monkeypatch, calls):
    monkeypatch.setenv("FIRECRAWL_API_KEY", "fc-1234567890abcdefghij")
    calls["routes"]["firecrawl.dev"] = FakeRes(200, {"success": True, "data": {"remaining_credits": 812}})
    out = auth_client.post("/api/settings/connections/FIRECRAWL_API_KEY/test").get_json()
    assert out["status"] == "정상" and "812" in out["message"]
    calls["routes"]["firecrawl.dev"] = FakeRes(404, None, "not found")
    out = auth_client.post("/api/settings/connections/FIRECRAWL_API_KEY/test").get_json()
    assert out["status"] == "정상" and out["detail"]["formatOnly"] is True

    monkeypatch.setenv("KEXIM_API_KEY", "kexim-test-key-dddd")
    days = []

    def kexim(url, params, headers):
        days.append(params["searchdate"])
        if len(days) == 1:
            return FakeRes(200, [])  # 휴일 → 빈 목록 → 하루 전
        return FakeRes(200, [{"result": 1, "cur_unit": "USD", "deal_bas_r": "1,390.5"}])
    calls["routes"]["koreaexim.go.kr"] = kexim
    out = auth_client.post("/api/settings/connections/KEXIM_API_KEY/test").get_json()
    assert out["status"] == "정상" and "1,390.5" in out["message"] and len(days) == 2


def test_connection_test_missing_unsupported_and_network(auth_client, monkeypatch, calls):
    out = auth_client.post("/api/settings/connections/SERPAPI_KEY/test").get_json()
    assert out["status"] == "미연결" and calls["seen"] == []
    res = auth_client.post("/api/settings/connections/SHOPIFY_TOKEN/test")
    assert res.status_code == 400 and "지원하지 않" in res.get_json()["error"]
    assert auth_client.post("/api/settings/connections/NOPE/test").status_code == 404

    monkeypatch.setenv("SERPAPI_KEY", "serp-leaky-key-4321")

    def boom(url, params, headers):
        raise requests.ConnectionError(f"Max retries exceeded with url: /account.json?api_key={params['api_key']}")
    calls["routes"]["serpapi.com"] = boom
    res = auth_client.post("/api/settings/connections/SERPAPI_KEY/test")
    out = res.get_json()
    assert out["status"] == "오류" and "ConnectionError" in out["message"]
    assert "serp-leaky-key" not in res.get_data(as_text=True)
    logs = json.dumps(auth_client.get("/api/worklog").get_json(), ensure_ascii=False)
    assert "serp-leaky-key" not in logs


def test_connection_test_gmail_smtp(auth_client, monkeypatch):
    seen = []

    def fake_login(address, password):
        seen.append((address, password))
        if password != "abcdefghijklmnop":
            raise smtplib.SMTPAuthenticationError(535, b"bad")
    monkeypatch.setattr(connections, "_smtp_login", fake_login)
    out = auth_client.post("/api/settings/connections/GMAIL_APP_PASSWORD/test").get_json()
    assert out["status"] == "미연결" and seen == []
    monkeypatch.setenv("GMAIL_ADDRESS", "team@example.com")
    monkeypatch.setenv("GMAIL_APP_PASSWORD", "abcd efgh ijkl zzzz")
    out = auth_client.post("/api/settings/connections/GMAIL_APP_PASSWORD/test").get_json()
    assert out["status"] == "오류" and "앱 비밀번호" in out["message"]
    monkeypatch.setenv("GMAIL_APP_PASSWORD", "abcd efgh ijkl mnop")
    out = auth_client.post("/api/settings/connections/GMAIL_ADDRESS/test").get_json()
    assert out["status"] == "정상" and seen[-1] == ("team@example.com", "abcdefghijklmnop")
    # 두 항목 모두 결과가 붙음
    assert _item(auth_client, "GMAIL_ADDRESS")["status"] == "정상"
    assert _item(auth_client, "GMAIL_APP_PASSWORD")["status"] == "정상"


# ───────── 작업 로그 Excel ─────────

def test_worklog_xlsx_sheets_and_formula_safety(auth_client):
    from openpyxl import load_workbook
    auth_client.post("/api/worklog", json={"part": "insights", "kind": "수집", "action": "=HYPERLINK(\"x\")",
                                           "detail": {"count": 3}})
    auth_client.post("/api/worklog", json={"part": "content", "kind": "AI 프롬프트", "action": "카피 초안",
                                           "detail": {"model": "gpt-4o-mini", "tokens": {"total": 812},
                                                      "prompt": "[system] …", "response": "초안"}})
    res = auth_client.get("/api/settings/worklog.xlsx")
    assert res.status_code == 200
    assert "spreadsheetml" in res.headers["Content-Type"]
    assert "reboot-settings-worklog-" in res.headers["Content-Disposition"]
    wb = load_workbook(io.BytesIO(res.data))
    assert wb.sheetnames == ["전체", "수집", "AI 프롬프트", "승인 · 선정", "기타"]
    ws = wb["수집"]
    assert ws.cell(row=2, column=4).value == '=HYPERLINK("x")' and ws.cell(row=2, column=4).data_type == "s"
    ai = wb["AI 프롬프트"]
    assert ai.cell(row=2, column=5).value == "gpt-4o-mini" and ai.cell(row=2, column=6).value == 812
    only = load_workbook(io.BytesIO(auth_client.get("/api/settings/worklog.xlsx?part=content").data))
    assert only["전체"].max_row == 2
