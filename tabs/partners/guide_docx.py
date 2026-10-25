"""협업 가이드 Markdown → Word(.docx) — README 3-3 서식.

표지(제목·발행일) · 브랜드 색 머리행 표 · 한 줄 요약 콜아웃(> 인용) · 머리글 REBOOT · 바닥글 Confidential · 페이지 번호 · 맑은 고딕.
지원 Markdown: # 제목 · 표(| a | b |) · > 인용 · - 목록 · 1. 번호 · **굵게** · *기울임* · `코드` · [글](주소)
"""
import io
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FONT, FONT_KO = "Malgun Gothic", "맑은 고딕"
INK = RGBColor(0x19, 0x1C, 0x18)
MUTED = RGBColor(0x85, 0x89, 0x81)
DEEP = RGBColor(0x63, 0x7C, 0x13)
DARK, LIME, SOFT, LIME_LINE = "191C18", "D7FF3F", "F5FBDF", "BEDF39"
LIME_RGB = RGBColor(0xD7, 0xFF, 0x3F)

_INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\([^)\s]+\)|(?<![*\w])\*[^*\s][^*]*\*(?![*\w]))")
_SEP = re.compile(r"^:?-{2,}:?$")


def _set_font(rpr_owner):
    """rFonts를 맑은 고딕으로 (테마 글꼴 속성은 지움 — 남아 있으면 테마가 우선함)."""
    rpr = rpr_owner.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rfonts.get(qn(attr)) is not None:
            del rfonts.attrib[qn(attr)]
    rfonts.set(qn("w:ascii"), FONT)
    rfonts.set(qn("w:hAnsi"), FONT)
    rfonts.set(qn("w:cs"), FONT)
    rfonts.set(qn("w:eastAsia"), FONT_KO)


def _styles(doc):
    for name, size, bold, color in (("Normal", 10.5, None, INK), ("Heading 1", 18, True, INK), ("Heading 2", 14, True, INK),
                                    ("Heading 3", 12, True, DEEP), ("List Bullet", 10.5, None, INK)):
        try:
            st = doc.styles[name]
        except KeyError:
            continue
        _set_font(st.element)
        st.font.size = Pt(size)
        if bold is not None:
            st.font.bold = bold
        st.font.color.rgb = color
        if name.startswith("Heading"):
            st.font.italic = False
            st.paragraph_format.space_before = Pt(14 if name != "Heading 3" else 10)
            st.paragraph_format.space_after = Pt(6)
    doc.styles["Normal"].paragraph_format.space_after = Pt(4)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.25


def _shade(tc_pr, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def _borders(tc_pr, left=None):
    """셀 테두리: left=(굵기, 색) 이외는 없음 (콜아웃용)."""
    b = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{side}")
        if side == "left" and left:
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), str(left[0]))
            el.set(qn("w:color"), left[1])
        else:
            el.set(qn("w:val"), "nil")
        b.append(el)
    tc_pr.append(b)


def _inline(par, text, bold=False, color=None, size=None):
    for tok in _INLINE.split(str(text)):
        if not tok:
            continue
        run_bold, italic, mono = bold, False, False
        if tok.startswith("**") and tok.endswith("**") and len(tok) > 4:
            tok, run_bold = tok[2:-2], True
        elif tok.startswith("`") and tok.endswith("`") and len(tok) > 2:
            tok, mono = tok[1:-1], True
        elif tok.startswith("[") and "](" in tok and tok.endswith(")"):
            label, url = tok[1:-1].split("](", 1)
            tok = f"{label} ({url})"
        elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
            tok, italic = tok[1:-1], True
        run = par.add_run(tok)
        run.bold = run_bold or None
        run.italic = italic or None
        if mono:
            run.font.name = "Consolas"
        if color is not None:
            run.font.color.rgb = color
        if size:
            run.font.size = Pt(size)


def _page_field(par):
    """PAGE 필드 (Word가 열 때 페이지 번호로 바꿈)."""
    def run_with(el):
        r = par.add_run()
        r.font.size = Pt(8)
        r.font.color.rgb = MUTED
        r._r.append(el)
    b = OxmlElement("w:fldChar")
    b.set(qn("w:fldCharType"), "begin")
    run_with(b)
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    run_with(instr)
    s = OxmlElement("w:fldChar")
    s.set(qn("w:fldCharType"), "separate")
    run_with(s)
    t = OxmlElement("w:t")
    t.text = "1"
    run_with(t)
    e = OxmlElement("w:fldChar")
    e.set(qn("w:fldCharType"), "end")
    run_with(e)


def _header_footer(section):
    section.different_first_page_header_footer = True   # 표지에는 머리글·바닥글 없음
    hp = section.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = hp.add_run("REBOOT")
    r.bold = True
    r.font.size = Pt(9)
    r.font.color.rgb = INK
    r2 = hp.add_run("   Partner Collaboration Guide")
    r2.font.size = Pt(8)
    r2.font.color.rgb = MUTED
    fp = section.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run("Confidential · ")
    r.font.size = Pt(8)
    r.font.color.rgb = MUTED
    _page_field(fp)


def _cover(doc, title, subtitle, date_text):
    for _ in range(5):
        doc.add_paragraph()
    p = doc.add_paragraph()
    _inline(p, "REBOOT  ·  PARTNER COLLABORATION GUIDE", bold=True, color=MUTED, size=9)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    _inline(p, title, bold=True, color=INK, size=28)
    if subtitle:
        p = doc.add_paragraph()
        _inline(p, subtitle, color=MUTED, size=14)
    band = doc.add_table(rows=1, cols=1)
    cell = band.rows[0].cells[0]
    tc_pr = cell._tc.get_or_add_tcPr()
    _shade(tc_pr, DARK)
    cp = cell.paragraphs[0]
    cp.paragraph_format.space_before = Pt(6)
    cp.paragraph_format.space_after = Pt(6)
    _inline(cp, "  Restart Your Day", bold=True, color=LIME_RGB, size=12)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    _inline(p, f"발행일 {date_text}", color=INK, size=11)
    p = doc.add_paragraph()
    _inline(p, "Confidential — 협업 파트너 전용 문서입니다.", color=MUTED, size=9)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def _callout(doc, text):
    t = doc.add_table(rows=1, cols=1)
    cell = t.rows[0].cells[0]
    tc_pr = cell._tc.get_or_add_tcPr()
    _borders(tc_pr, left=(36, LIME_LINE))
    _shade(tc_pr, SOFT)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    _inline(p, text, bold=True, color=INK, size=11)
    doc.add_paragraph()


def _cells(line):
    s = line.strip()
    if s.startswith("|"):
        s = s[1:]
    if s.endswith("|"):
        s = s[:-1]
    return [c.strip() for c in re.split(r"(?<!\\)\|", s)]


def _table(doc, lines):
    rows = [_cells(x) for x in lines]
    rows = [r for r in rows if not all(_SEP.match(c.replace(" ", "")) for c in r if c) or not any(r)]
    rows = [r for r in rows if any(c for c in r)]
    if not rows:
        return
    ncols = max(len(r) for r in rows)
    t = doc.add_table(rows=len(rows), cols=ncols)
    t.style = doc.styles["Table Grid"]
    for i, r in enumerate(rows):
        for j in range(ncols):
            cell = t.rows[i].cells[j]
            par = cell.paragraphs[0]
            text = r[j].replace("\\|", "|") if j < len(r) else ""
            if i == 0:
                _shade(cell._tc.get_or_add_tcPr(), DARK)
                _inline(par, text.replace("**", ""), bold=True, color=LIME_RGB, size=10)
            else:
                for k, chunk in enumerate(re.split(r"<br\s*/?>", text)):
                    if k:
                        par = cell.add_paragraph()
                    _inline(par, chunk, size=10)
    doc.add_paragraph()


def _blocks(md):
    lines = str(md or "").replace("\r\n", "\n").split("\n")
    i = 0
    special = re.compile(r"^\s*(#{1,6}\s|\||>|[-*+]\s|\d+[.)]\s|(-{3,}|\*{3,}|_{3,})\s*$)")
    while i < len(lines):
        s = lines[i].strip()
        if not s:
            i += 1
            continue
        if s.startswith("|"):
            tbl = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                tbl.append(lines[i].strip())
                i += 1
            yield ("table", tbl)
            continue
        m = re.match(r"^(#{1,6})\s+(.*)$", s)
        if m:
            yield ("h", len(m.group(1)), m.group(2).strip())
            i += 1
            continue
        if s.startswith(">"):
            quote = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                quote.append(lines[i].strip().lstrip(">").strip())
                i += 1
            yield ("quote", " ".join(q for q in quote if q))
            continue
        if re.match(r"^(-{3,}|\*{3,}|_{3,})$", s):
            i += 1
            continue
        m = re.match(r"^[-*+]\s+(.*)$", s)
        if m:
            yield ("bullet", m.group(1))
            i += 1
            continue
        m = re.match(r"^(\d+)[.)]\s+(.*)$", s)
        if m:
            yield ("number", m.group(1), m.group(2))
            i += 1
            continue
        para = [s]
        i += 1
        while i < len(lines) and lines[i].strip() and not special.match(lines[i]):
            para.append(lines[i].strip())
            i += 1
        yield ("p", " ".join(para))


def first_heading(md):
    for b in _blocks(md):
        if b[0] == "h" and b[1] == 1:
            return b[2]
    return ""


def build(markdown, title="", subtitle="", date_text=""):
    """→ .docx 바이트"""
    doc = Document()
    _styles(doc)
    sec = doc.sections[0]
    sec.top_margin = sec.bottom_margin = Cm(2.2)
    sec.left_margin = sec.right_margin = Cm(2.2)
    title = (title or first_heading(markdown) or "REBOOT 협업 가이드").replace("#", "").strip()
    doc.core_properties.title = title
    doc.core_properties.author = "REBOOT Partnerships"
    doc.core_properties.subject = subtitle or ""
    _header_footer(sec)
    _cover(doc, title, subtitle, date_text)
    skipped_title = False
    for b in _blocks(markdown):
        kind = b[0]
        if kind == "h":
            level, text = b[1], b[2]
            if level == 1 and not skipped_title:
                skipped_title = True   # 표지에 이미 있음
                continue
            doc.add_heading(text.replace("**", ""), level=min(3, max(1, level - 1 if level > 1 else 1)))
        elif kind == "table":
            _table(doc, b[1])
        elif kind == "quote":
            _callout(doc, b[1])
        elif kind == "bullet":
            p = doc.add_paragraph(style="List Bullet")
            _inline(p, b[1])
        elif kind == "number":
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.7)
            p.paragraph_format.first_line_indent = Cm(-0.5)
            _inline(p, f"{b[1]}.  ", bold=True, color=DEEP)
            _inline(p, b[2])
        else:
            _inline(doc.add_paragraph(), b[1])
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
