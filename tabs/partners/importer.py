"""업체 파일 가져오기: 케이터링 CSV · 인플루언서 XLSX(6행 머리글) → partners_contacts.

데이터 매핑 (README 3-1)
  id        C + No (3자리) · I + 번호 (3자리) · 번호가 숫자가 아니면 CX01 · IX01 … (예시·테스트 행, 실적 더미에서 제외)
  name      Company Name          · 업체명
  email     Contact Email         · 공개 이메일            (첫 번째 이메일 주소만, 나머지 글자는 비고로)
  phone     Phone Number          · 공개 전화/문자
  url       Official Website URL  · 공식 사이트 URL
  location  Office Location       · 미국 시장 범위
  category  Category Focus        · 서비스 형태 · 서비스/검토 이유
  note      Why Relevant          · 주의/연락처 용도

세부 카테고리는 SUBCATEGORIES의 키워드로 자동 분류합니다 (한 업체에 태그 여러 개 가능).
분류 키워드는 이 파일의 SUBCATEGORIES에서 고칩니다. 영어는 대소문자 무시·단어 앞부분 일치(3글자 이하는 단어 전체 일치), 한국어는 포함 여부.
"""
import csv
import hashlib
import io
import json
import re
from pathlib import Path

from server.config import ROOT
from server.extensions import db
from .models import PartnerContact

CATERING, INFLUENCER = "케이터링", "인플루언서"
TYPES = (CATERING, INFLUENCER)
MAX_ROWS = 2000
MAX_FILE_BYTES = 5 * 1024 * 1024
HEADER_ROW_XLSX = 6   # 인플루언서 XLSX는 6행이 머리글

SUBCATEGORIES = {
    CATERING: [
        {"key": "festival", "emoji": "🎵", "label": "페스티벌·콘서트",
         "keywords": ["festival", "concert", "music", "edm", "live event", "amphitheater", "페스티벌", "콘서트", "공연", "음악"]},
        {"key": "nightlife", "emoji": "🌙", "label": "나이트라이프·클럽",
         "keywords": ["nightlife", "night life", "nightclub", "night club", "club", "lounge", "after party", "afterparty", "dj",
                      "나이트", "클럽", "라운지"]},
        {"key": "campus", "emoji": "🎓", "label": "대학·테일게이트",
         "keywords": ["college", "university", "campus", "tailgate", "tailgating", "game day", "gameday", "fraternity",
                      "sorority", "greek life", "student", "alumni", "대학", "캠퍼스", "테일게이트"]},
        {"key": "vip", "emoji": "💎", "label": "VIP·럭셔리·갈라",
         "keywords": ["vip", "luxury", "gala", "upscale", "high end", "premium", "black tie", "exclusive", "celebrity",
                      "럭셔리", "갈라", "고급"]},
        {"key": "mobile_bar", "emoji": "🍸", "label": "모바일 바·믹솔로지",
         "keywords": ["mobile bar", "bartend", "mixolog", "cocktail", "bar service", "beverage service", "drink service",
                      "bar cart", "칵테일", "바텐", "믹솔로지", "모바일 바"]},
        {"key": "brand_popup", "emoji": "🚀", "label": "브랜드 팝업·런칭",
         "keywords": ["brand activation", "activation", "pop up", "popup", "launch", "experiential", "sampling",
                      "product launch", "brand event", "팝업", "런칭", "론칭", "샘플링"]},
        {"key": "pool_beach", "emoji": "🏖️", "label": "풀·비치·루프탑",
         "keywords": ["pool", "beach", "rooftop", "roof top", "yacht", "boat", "resort", "cabana",
                      "풀파티", "수영장", "비치", "해변", "루프탑"]},
    ],
    INFLUENCER: [
        {"key": "agency", "emoji": "🏢", "label": "대행사·매니지먼트",
         "keywords": ["agency", "agencies", "management", "talent", "representation", "roster", "mcn",
                      "대행", "에이전시", "매니지먼트", "소속사"]},
        {"key": "platform", "emoji": "🧩", "label": "플랫폼",
         "keywords": ["platform", "marketplace", "software", "saas", "app", "database", "network",
                      "플랫폼", "마켓플레이스", "솔루션"]},
        {"key": "ugc", "emoji": "🎬", "label": "UGC·광고 소재",
         "keywords": ["ugc", "user generated", "creative", "ad creative", "content production", "content creation",
                      "video ad", "광고 소재", "소재", "콘텐츠 제작", "영상 제작"]},
        {"key": "pr_social", "emoji": "📣", "label": "PR·소셜",
         "keywords": ["pr", "public relations", "social media", "social", "communications", "publicity",
                      "홍보", "소셜", "언론"]},
        {"key": "micro_seeding", "emoji": "🌱", "label": "마이크로·제품 시딩",
         "keywords": ["micro", "nano", "seeding", "gifting", "product seeding", "sampling",
                      "마이크로", "나노", "시딩", "협찬", "제품 제공"]},
        {"key": "food_bev", "emoji": "🍔", "label": "식음료·외식",
         "keywords": ["food", "beverage", "f&b", "drink", "restaurant", "foodie", "culinary", "cpg", "snack",
                      "식음료", "외식", "음료", "푸드", "맛집"]},
        {"key": "tiktok_dtc", "emoji": "🛒", "label": "TikTok Shop·DTC",
         "keywords": ["tiktok shop", "dtc", "d2c", "direct to consumer", "e commerce", "ecommerce", "shopify",
                      "affiliate", "live shopping", "틱톡샵", "틱톡 샵", "커머스", "라이브 커머스"]},
    ],
}

# 원본 열 이름 (소문자·공백 정리 후 비교). 앞에 있는 이름이 우선
CATERING_COLUMNS = {
    "no": ("no", "no.", "#", "번호"),
    "name": ("company name", "company", "업체명", "name"),
    "email": ("contact email", "email", "e-mail", "이메일"),
    "phone": ("phone number", "phone", "전화번호"),
    "url": ("official website url", "website url", "website", "url"),
    "location": ("office location", "location", "지역"),
    "category": ("category focus", "category", "카테고리"),
    "note": ("why relevant", "notes", "note", "비고"),
}
INFLUENCER_COLUMNS = {
    "no": ("번호", "no", "no."),
    "name": ("업체명", "company name", "name"),
    "email": ("공개 이메일", "이메일", "email"),
    "phone": ("공개 전화/문자", "공개 전화", "전화", "phone"),
    "url": ("공식 사이트 url", "공식 사이트", "url", "website"),
    "location": ("미국 시장 범위", "시장 범위", "location"),
    "category": ("서비스 형태",),
    "category2": ("서비스/검토 이유", "서비스 / 검토 이유", "검토 이유"),
    "note": ("주의/연락처 용도", "주의 / 연락처 용도", "비고", "note"),
}

EMAIL_RE = re.compile(r"[A-Za-z0-9._%+'-]+@[A-Za-z0-9-]+(?:\.[A-Za-z0-9-]+)*\.[A-Za-z]{2,}")
_DOMAIN_RE = re.compile(r"^(?:www\.)?[A-Za-z0-9-]+(?:\.[A-Za-z0-9-]+)*\.[A-Za-z]{2,}(?:[/?#]\S*)?$")


class ImportFailed(ValueError):
    """화면에 그대로 보여 줄 가져오기 실패 이유."""


# ───────── 세부 카테고리 ─────────

def _pattern(keyword):
    kw = keyword.lower().strip()
    if re.fullmatch(r"[a-z0-9&+.\s-]+", kw):
        body = r"[\s-]*".join(re.escape(p) for p in re.split(r"[\s-]+", kw) if p)
        tail = r"(?![a-z0-9])" if len(kw) <= 3 else ""   # pr · app · dj 는 단어 전체만
        return re.compile(r"(?<![a-z0-9])" + body + tail)
    return re.compile(re.escape(kw))


_COMPILED = {}


def _compiled(ptype):
    key = (ptype, json.dumps(SUBCATEGORIES.get(ptype, []), ensure_ascii=False))
    if key not in _COMPILED:
        _COMPILED[key] = [(sc["key"], [_pattern(k) for k in sc["keywords"]]) for sc in SUBCATEGORIES.get(ptype, [])]
    return _COMPILED[key]


def tag(ptype, text):
    src = str(text or "").lower()
    return [key for key, pats in _compiled(ptype) if any(p.search(src) for p in pats)]


def subcategory_meta():
    return {t: [{"key": s["key"], "label": s["label"], "emoji": s["emoji"]} for s in SUBCATEGORIES[t]] for t in TYPES}


# ───────── 파일 읽기 ─────────

def _decode(raw):
    for enc in ("utf-8-sig", "cp949"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("latin-1")


def read_csv_rows(raw):
    text = _decode(raw)
    first = text.split("\n", 1)[0]
    delim = "\t" if first.count("\t") > first.count(",") else ","
    rows = []
    for i, row in enumerate(csv.reader(io.StringIO(text), delimiter=delim)):
        if i > MAX_ROWS + 50:
            raise ImportFailed(f"행이 너무 많습니다 ({MAX_ROWS}행까지).")
        rows.append(row)
    return rows


def read_xlsx_rows(raw):
    from openpyxl import load_workbook
    try:
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception:
        raise ImportFailed("XLSX 파일을 읽지 못했습니다. 엑셀에서 .xlsx 형식으로 다시 저장해 보세요.")
    try:
        sheets = []
        for ws in wb.worksheets:
            rows = []
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i > MAX_ROWS + 50:
                    raise ImportFailed(f"행이 너무 많습니다 ({MAX_ROWS}행까지).")
                rows.append(list(row))
            sheets.append(rows)
        return sheets
    finally:
        wb.close()


def _norm(h):
    return re.sub(r"\s+", " ", str(h if h is not None else "").replace("﻿", "")).strip().lower()


def _index(header, columns):
    norm = [_norm(h) for h in header]
    idx = {}
    for field, names in columns.items():
        for n in names:
            if n in norm:
                idx[field] = norm.index(n)
                break
    return idx


def _header_ok(idx):
    return "name" in idx and len(idx) >= 3


def _find_header(rows, columns, prefer=None):
    """머리글 행 찾기: prefer(0부터) 행이 맞으면 그 행, 아니면 앞 30행에서 처음 맞는 행."""
    if prefer is not None and prefer < len(rows):
        idx = _index(rows[prefer], columns)
        if _header_ok(idx):
            return prefer, idx
    for i, row in enumerate(rows[:30]):
        idx = _index(row, columns)
        if _header_ok(idx):
            return i, idx
    return None


def _cell(row, idx, field):
    i = idx.get(field)
    if i is None or i >= len(row):
        return ""
    v = row[i]
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def _num(s):
    s = str(s or "").strip()
    return int(float(s)) if re.fullmatch(r"\d+(?:\.0+)?", s) else None


def _url(s):
    s = (s or "").strip()
    if re.match(r"^https?://", s, re.I):
        return s[:500]
    if _DOMAIN_RE.match(s):
        return ("https://" + s)[:500]
    return s[:500]


def _record(row, idx, ptype):
    name = _cell(row, idx, "name")
    if not name:
        return None
    category = _cell(row, idx, "category")
    if ptype == INFLUENCER:
        category = " · ".join(x for x in (category, _cell(row, idx, "category2")) if x)
    note = _cell(row, idx, "note")
    email_raw = _cell(row, idx, "email")
    emails = EMAIL_RE.findall(email_raw)
    rest = email_raw
    for e in emails[:1]:
        rest = rest.replace(e, "", 1)
    rest = re.sub(r"^[\s,;/|·-]+|[\s,;/|·-]+$", "", rest.replace("mailto:", ""))
    if rest:  # 이메일 칸의 다른 글자(다른 주소·'문의 양식' 등)는 비고로
        note = (note + " · " if note else "") + f"이메일 칸: {rest}"
    return {"no": _cell(row, idx, "no"), "name": name[:200], "email": emails[0][:255] if emails else "",
            "phone": _cell(row, idx, "phone")[:100], "url": _url(_cell(row, idx, "url")),
            "location": _cell(row, idx, "location")[:200], "category": category[:2000], "note": note[:2000]}


def _assign_ids(recs, prefix):
    used, x = set(), 0
    for r in recs:
        n = _num(r["no"])
        if n is None:
            x += 1
            base = f"{prefix}X{x:02d}"
        else:
            base = f"{prefix}{n:03d}"
        rid, k = base, 1
        while rid in used:
            k += 1
            rid = f"{base}-{k}"
        used.add(rid)
        r["id"] = rid


def parse_rows(ptype, sheets, prefer=None):
    """sheets: [[행…]…] (시트별 행 목록) → (업체 목록, 안내 문구 목록)"""
    columns = CATERING_COLUMNS if ptype == CATERING else INFLUENCER_COLUMNS
    found = None
    for rows in sheets:
        h = _find_header(rows, columns, prefer)
        if h is not None:
            found = (rows, h)
            break
    if found is None:
        need = "No · Company Name · Contact Email …" if ptype == CATERING else "번호 · 업체명 · 공개 이메일 … (6행)"
        raise ImportFailed(f"{ptype} 파일에서 머리글({need})을 찾지 못했습니다.")
    rows, (hi, idx) = found
    warnings = []
    if prefer is not None and hi != prefer:
        warnings.append(f"머리글을 {prefer + 1}행이 아니라 {hi + 1}행에서 찾았습니다.")
    recs = []
    for row in rows[hi + 1:]:
        if not row or all(v in (None, "") or str(v).strip() == "" for v in row):
            continue
        r = _record(row, idx, ptype)
        if r is None:
            continue
        recs.append(r)
        if len(recs) > MAX_ROWS:
            raise ImportFailed(f"행이 너무 많습니다 ({MAX_ROWS}행까지).")
    if not recs:
        raise ImportFailed(f"{ptype} 파일에 가져올 업체 행이 없습니다.")
    _assign_ids(recs, "C" if ptype == CATERING else "I")
    for r in recs:
        r["tags"] = tag(ptype, r["category"] + " " + r["name"])
    missing = [label for field, label in (("email", "이메일"), ("location", "지역"), ("category", "카테고리")) if field not in idx]
    if missing:
        warnings.append("찾지 못한 열: " + ", ".join(missing))
    return recs, warnings


def parse_file(ptype, filename, raw):
    ext = Path(str(filename or "")).suffix.lower()
    if len(raw) > MAX_FILE_BYTES:
        raise ImportFailed("파일이 너무 큽니다 (5MB까지).")
    if ext == ".xlsx":
        return parse_rows(ptype, read_xlsx_rows(raw), prefer=HEADER_ROW_XLSX - 1 if ptype == INFLUENCER else None)
    if ext in (".csv", ".txt", ".tsv"):
        return parse_rows(ptype, [read_csv_rows(raw)], prefer=0 if ptype == CATERING else None)
    raise ImportFailed("CSV 또는 XLSX 파일만 가져올 수 있습니다.")


def detect_type(filename, raw):
    """구분 없이 올린 파일: .xlsx → 인플루언서, CSV는 머리글로 판단."""
    ext = Path(str(filename or "")).suffix.lower()
    if ext == ".xlsx":
        return INFLUENCER
    if ext not in (".csv", ".txt", ".tsv"):
        raise ImportFailed("CSV 또는 XLSX 파일만 가져올 수 있습니다.")
    for row in read_csv_rows(raw)[:30]:
        norm = {_norm(c) for c in row}
        if "업체명" in norm:
            return INFLUENCER
        if "company name" in norm:
            return CATERING
    raise ImportFailed("파일 종류를 알 수 없습니다. 케이터링(Company Name) 또는 인플루언서(업체명) 머리글이 있어야 합니다.")


# ───────── 기본 업체 파일 (자동 불러오기) ─────────

BUNDLED_DIR = ROOT / "tabs" / "partners" / "data" / "contacts"
BUNDLED_PREFIX = "기본 파일 · "


def _bundled_label(name, raw):
    return f"{BUNDLED_PREFIX}{name} · {hashlib.sha1(raw).hexdigest()[:8]}"


def bundled_files():
    """BUNDLED_DIR의 CSV·XLSX → {구분: (파일 이름, 내용)} (구분마다 이름순 첫 파일)"""
    out = {}
    if not BUNDLED_DIR.is_dir():
        return out
    for f in sorted(BUNDLED_DIR.iterdir()):
        if f.suffix.lower() not in (".csv", ".xlsx") or not f.is_file():
            continue
        raw = f.read_bytes()
        try:
            t = detect_type(f.name, raw)
        except ImportFailed:
            continue
        out.setdefault(t, (f.name, raw))
    return out


def sync_bundled():
    """기본 파일을 DB에 반영: 구분이 비었거나, 기본 파일에서 가져온 목록인데 파일 내용이 바뀌었을 때만.
    화면에서 직접 올린 파일(출처가 기본 파일이 아님)은 덮어쓰지 않습니다. 반영한 구분 목록을 돌려줌."""
    loaded = []
    for t, (name, raw) in bundled_files().items():
        label = _bundled_label(name, raw)
        cur = PartnerContact.query.filter_by(ptype=t).first()
        if cur is not None and (cur.source_file == label or not cur.source_file.startswith(BUNDLED_PREFIX)):
            continue
        try:
            recs, _ = parse_file(t, name, raw)
        except ImportFailed:
            continue
        replace_contacts(t, recs, label, "자동")
        loaded.append(t)
    return loaded


def replace_contacts(ptype, recs, source_file, user=""):
    """같은 구분의 업체를 새 목록으로 바꿈 (협업 이력은 업체 id·이름을 따로 저장하므로 그대로 남음)."""
    PartnerContact.query.filter_by(ptype=ptype).delete()
    for i, r in enumerate(recs):
        db.session.add(PartnerContact(
            id=r["id"], ptype=ptype, name=r["name"], email=r["email"], phone=r["phone"], url=r["url"],
            location=r["location"], category=r["category"], note=r["note"],
            tags=json.dumps(r["tags"], ensure_ascii=False), sort=i, source_file=str(source_file or "")[:200],
            imported_by=str(user or "")[:100]))
    db.session.commit()
    return len(recs)


# ───────── 예시 · 양식 ─────────

def examples():
    """tabs/partners/data/examples.json — 화면 확인용 가상 업체 2곳씩 (example.com)."""
    path = ROOT / "tabs" / "partners" / "data" / "examples.json"
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return {}


def example_records(ptype):
    ex = examples().get("catering" if ptype == CATERING else "influencer") or {}
    if not ex.get("header") or not ex.get("rows"):
        raise ImportFailed("예시 데이터 파일(tabs/partners/data/examples.json)을 읽지 못했습니다.")
    return parse_rows(ptype, [[ex["header"]] + ex["rows"]])[0]


def template_csv():
    ex = examples().get("catering") or {}
    out = io.StringIO()
    w = csv.writer(out)
    w.writerow(ex.get("header") or list(n[0].title() for n in CATERING_COLUMNS.values()))
    for row in ex.get("rows") or []:
        w.writerow(row)
    return "﻿" + out.getvalue()


def template_xlsx():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    ex = examples().get("influencer") or {}
    wb = Workbook()
    ws = wb.active
    ws.title = "Contacts"
    for i, line in enumerate((ex.get("title") or [])[:HEADER_ROW_XLSX - 1], start=1):
        ws.cell(row=i, column=1, value=line)
    header = ex.get("header") or [n[0] for n in INFLUENCER_COLUMNS.values()]
    for j, h in enumerate(header, start=1):
        c = ws.cell(row=HEADER_ROW_XLSX, column=j, value=h)
        c.font = Font(bold=True, color="D7FF3F")
        c.fill = PatternFill("solid", fgColor="191C18")
        ws.column_dimensions[c.column_letter].width = 22
    for i, row in enumerate(ex.get("rows") or [], start=HEADER_ROW_XLSX + 1):
        for j, v in enumerate(row, start=1):
            ws.cell(row=i, column=j, value=v)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
