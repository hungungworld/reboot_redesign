"""인플루언서 · 파트너 API — 인터넷·키 없이 (OpenAI · SMTP · IMAP은 가짜로 바꿔 끼움)."""
import io
import json
import shutil
import smtplib
import subprocess
from datetime import date, datetime, timedelta, timezone
from email.message import EmailMessage
from email.utils import format_datetime, make_msgid
from pathlib import Path

import pytest

from tabs.partners import importer, mailer, sales

ROOT = Path(__file__).resolve().parents[3]  # 저장소 맨 위


@pytest.fixture(autouse=True)
def no_bundled_contacts(tmp_path, monkeypatch):
    """기본 업체 파일(tabs/partners/data/contacts/) 대신 빈 폴더 — 각 테스트는 빈 목록에서 시작."""
    d = tmp_path / "bundled-contacts"
    d.mkdir()
    monkeypatch.setattr(importer, "BUNDLED_DIR", d)
    return d

CATERING_CSV = (
    "No,Company Name,Contact Email,Phone Number,Official Website URL,Office Location,Category Focus,Why Relevant\n"
    "1,Festival Bar Co,mailto:info@festbar.example.com / backup@festbar.example.com,555-0100,festbar.example.com,"
    "\"Austin, TX\",Music festival & concert mobile bar service,Large summer festivals\n"
    "2,Campus Tailgate Catering,events@campus.example.com,555-0101,https://campus.example.com,\"Columbus, OH\","
    "College tailgate and game day catering,Big Ten game days\n"
    "x,Test Row Caterer,,,,,Product launch pop-up events,test row\n"
)


def make_xlsx(rows, header_row=6):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.cell(row=1, column=1, value="US Influencer Public Contacts (테스트)")
    header = ["번호", "업체명", "공개 이메일", "공개 전화/문자", "공식 사이트 URL", "미국 시장 범위", "서비스 형태", "서비스/검토 이유", "주의/연락처 용도"]
    for j, h in enumerate(header, start=1):
        ws.cell(row=header_row, column=j, value=h)
    for i, row in enumerate(rows, start=header_row + 1):
        for j, v in enumerate(row, start=1):
            ws.cell(row=i, column=j, value=v)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


INFLUENCER_ROWS = [
    [1, "Alpha Talent Agency", "hello@alpha.example.com", "", "https://alpha.example.com", "미국 전역",
     "인플루언서 대행사", "식음료 브랜드 캠페인 경험, TikTok Shop", "공개 문의 메일"],
    [2.0, "Beta UGC Platform", "team@beta.example.com", "", "beta.example.com", "LA · NY",
     "UGC 광고 소재 플랫폼", "크리에이터 매칭", ""],
    ["예시)", "예시) Test Creator", "test@example.com", "", "", "", "PR agency", "테스트 행", ""],
]


def upload(client, **files):
    data = {k: (io.BytesIO(v[1]), v[0]) for k, v in files.items()}
    return client.post("/api/partners/import", data=data, content_type="multipart/form-data")


def import_both(client):
    res = upload(client, catering=("top50catering.csv", CATERING_CSV.encode("utf-8")),
                 influencer=("US_Influencer_Public_Contacts_72.xlsx", make_xlsx(INFLUENCER_ROWS)))
    assert res.status_code == 200, res.get_json()
    return res.get_json()


class FakeAI:
    """server.core.ai.openai_json 대신 — 받은 프롬프트를 기록하고 정해진 답을 돌려줌."""

    def __init__(self, reply):
        self.reply = reply
        self.calls = []

    def __call__(self, system, user, **kw):
        self.calls.append({"system": system, "user": user if isinstance(user, str) else json.dumps(user, ensure_ascii=False), **kw})
        return dict(self.reply)


@pytest.fixture()
def ai(monkeypatch):
    monkeypatch.setenv("OPENAI_API_KEY", "sk-test-not-real")

    def use(reply):
        fake = FakeAI(reply)
        monkeypatch.setattr("server.core.ai.openai_json", fake)
        return fake
    return use


class FakeSMTP:
    instances = []

    def __init__(self, fail=None):
        self.fail, self.sent, self.login_args = fail, [], None
        FakeSMTP.instances.append(self)

    def login(self, user, password):
        self.login_args = (user, password)
        if self.fail == "auth":
            raise smtplib.SMTPAuthenticationError(535, b"5.7.8 Username and Password not accepted")

    def send_message(self, msg):
        self.sent.append(msg)
        return {}

    def quit(self):
        pass


@pytest.fixture()
def gmail(monkeypatch):
    monkeypatch.setenv("GMAIL_ADDRESS", "team@reboot-test.example.com")
    monkeypatch.setenv("GMAIL_APP_PASSWORD", "abcd efgh ijkl mnop")
    FakeSMTP.instances = []

    def use(fail=None):
        monkeypatch.setattr(mailer, "connect_smtp", lambda: FakeSMTP(fail))
        return FakeSMTP.instances
    use()
    return use


def send(client, **over):
    body = {"to": "events@campus.example.com", "subject": "REBOOT x Campus Tailgate Catering",
            "body": "Hello team,\n\nWe would love to collaborate.\n\nREBOOT Partnerships Team", "owner": "혁주",
            "confirm": True, "partner_id": "C002", "partner_name": "Campus Tailgate Catering", "partner_type": "케이터링"}
    body.update(over)
    return client.post("/api/partners/email/send", json=body)


# ───────── 로그인 · 상태 ─────────

def test_requires_login(client):
    assert client.get("/api/partners/contacts").status_code == 401
    assert client.post("/api/partners/email/send", json={"confirm": True}).status_code == 401
    assert client.get("/api/partners/sales").status_code == 401


def test_status_and_config(auth_client, monkeypatch):
    for name in ("REGULATION_DB_URL", "PRODUCT_DESC"):
        monkeypatch.delenv(name, raising=False)
    st = auth_client.get("/api/partners/status").get_json()
    assert set(st) == {"regulationDb", "openai", "model", "smtp", "sender"}
    assert (st["regulationDb"], st["openai"], st["smtp"], st["sender"]) == (False, False, False, "")
    monkeypatch.setenv("REGULATION_DB_URL", "postgresql://example.invalid/regulation")
    assert auth_client.get("/api/partners/status").get_json()["regulationDb"] is True
    cfg = auth_client.get("/api/partners/config").get_json()
    assert [len(cfg["subcategories"][t]) for t in ("케이터링", "인플루언서")] == [7, 7]
    assert [len(cfg["guidePurposes"][t]) for t in ("케이터링", "인플루언서")] == [6, 6]
    assert all("{partner}" in p and "{ideal}" in p and "{product}" in p for p in cfg["prompts"].values())
    assert len(cfg["guideSections"]) == 8 and "hangover" not in cfg["productDesc"].lower()


def test_partner_page_loads_scripts_in_order(auth_client):
    html = auth_client.get("/dashboard/partners").get_data(as_text=True)
    order = [html.index(f"/parts/partners/{f}") for f in
             ("data-outreach.js", "data-records.js", "outreach.js", "records.js", "partners.js")]
    assert order == sorted(order) and "/parts/partners/partners.css" in html


# ───────── 업체 가져오기 · 검색 ─────────

def test_empty_state(auth_client):
    data = auth_client.get("/api/partners/contacts").get_json()
    assert data["items"] == [] and data["total"] == 0
    ov = auth_client.get("/api/partners/overview").get_json()
    assert (ov["contacted"], ov["active"], ov["replied"], ov["awaiting"]) == (0, 0, 0, 0)


def test_import_catering_csv_and_influencer_xlsx(auth_client):
    out = import_both(auth_client)
    by_type = {r["type"]: r for r in out["results"]}
    assert by_type["케이터링"]["count"] == 3 and by_type["인플루언서"]["count"] == 3
    items = {c["id"]: c for c in auth_client.get("/api/partners/contacts").get_json()["items"]}
    assert set(items) == {"C001", "C002", "CX01", "I001", "I002", "IX01"}
    c1 = items["C001"]
    assert c1["email"] == "info@festbar.example.com" and "backup@festbar.example.com" in c1["note"]
    assert c1["url"] == "https://festbar.example.com" and c1["location"] == "Austin, TX"
    assert {"festival", "mobile_bar"} <= set(c1["tags"]) and "campus" in items["C002"]["tags"]
    i1 = items["I001"]
    assert i1["type"] == "인플루언서" and i1["category"] == "인플루언서 대행사 · 식음료 브랜드 캠페인 경험, TikTok Shop"
    assert {"agency", "food_bev", "tiktok_dtc"} <= set(i1["tags"])
    assert {"ugc", "platform"} <= set(items["I002"]["tags"]) and items["I002"]["url"] == "https://beta.example.com"
    assert items["IX01"]["example"] is True and items["CX01"]["example"] is True and items["I001"]["example"] is False


def test_bundled_files_load_automatically(auth_client, no_bundled_contacts):
    d = no_bundled_contacts
    (d / "catering.csv").write_text(CATERING_CSV, encoding="utf-8")
    data = auth_client.get("/api/partners/contacts").get_json()
    assert data["total"] == 3 and data["imports"]["케이터링"]["file"].startswith("기본 파일 · catering.csv")
    assert data["imports"]["인플루언서"]["count"] == 0
    # 파일 내용이 바뀌면 다시 불러옴 (협업 이력은 그대로)
    (d / "catering.csv").write_text("".join(CATERING_CSV.splitlines(keepends=True)[:3]), encoding="utf-8")  # 테스트 행 뺌
    assert auth_client.get("/api/partners/contacts").get_json()["total"] == 2
    # 화면에서 직접 올린 파일은 기본 파일로 덮어쓰지 않음
    upload(auth_client, catering=("mine.csv", CATERING_CSV.encode("utf-8")))
    data = auth_client.get("/api/partners/contacts").get_json()
    assert data["total"] == 3 and data["imports"]["케이터링"]["file"] == "mine.csv"
    # 종류를 알 수 없는 파일은 건너뜀
    (d / "notes.csv").write_text("foo,bar\n1,2\n", encoding="utf-8")
    assert auth_client.get("/api/partners/overview").status_code == 200


def test_repo_bundled_contacts_parse():
    """저장소에 넣어 둔 기본 업체 파일이 머리글을 찾아 읽히는지 (내용 값은 보지 않음)."""
    d = ROOT / "tabs" / "partners" / "data" / "contacts"
    for f in sorted(d.glob("*.csv")):
        raw = f.read_bytes()
        recs, _ = importer.parse_file(importer.detect_type(f.name, raw), f.name, raw)
        assert recs and all(r["name"] for r in recs)


def test_xlsx_header_must_be_found(auth_client):
    res = upload(auth_client, influencer=("x.xlsx", make_xlsx(INFLUENCER_ROWS, header_row=9)))
    assert res.status_code == 200 and "9행" in res.get_json()["results"][0]["warnings"][0]
    bad = upload(auth_client, catering=("c.csv", b"foo,bar\n1,2\n"))
    assert bad.status_code == 400 and "머리글" in bad.get_json()["error"]
    assert upload(auth_client, catering=("c.pdf", b"%PDF")).status_code == 400
    assert auth_client.post("/api/partners/import", data={}, content_type="multipart/form-data").status_code == 400
    # 구분 없이 올린 파일은 머리글로 판단
    auto = upload(auth_client, file=("list.csv", CATERING_CSV.encode("utf-8")))
    assert auto.status_code == 200 and auto.get_json()["results"][0]["type"] == "케이터링"


def test_reimport_replaces_only_that_type_and_keeps_history(auth_client):
    import_both(auth_client)
    res = auth_client.post("/api/partners/history", json={"partner_id": "C002", "activity": "전화", "progress": "협의 중",
                                                          "owner": "효선", "memo": "통화함"})
    assert res.status_code == 201
    one = CATERING_CSV.split("\n")[0] + "\n7,New Rooftop Bar,hi@roof.example.com,,,,Rooftop pool parties,\n"
    assert upload(auth_client, catering=("top50catering.csv", one.encode("cp949"))).status_code == 200
    items = auth_client.get("/api/partners/contacts").get_json()["items"]
    assert [c["id"] for c in items if c["type"] == "케이터링"] == ["C007"]
    assert len([c for c in items if c["type"] == "인플루언서"]) == 3
    hist = auth_client.get("/api/partners/history?partner_id=C002").get_json()["items"]
    assert len(hist) == 1 and hist[0]["partnerName"] == "Campus Tailgate Catering"


def test_contact_filters(auth_client):
    import_both(auth_client)

    def ids(qs):
        return sorted(c["id"] for c in auth_client.get("/api/partners/contacts" + qs).get_json()["items"])
    assert ids("?type=케이터링") == ["C001", "C002", "CX01"]
    assert ids("?type=influencer") == ["I001", "I002", "IX01"]
    assert ids("?tags=festival,campus") == ["C001", "C002"]          # 칩은 OR
    assert ids("?q=tailgate") == ["C002"]
    assert ids("?q=alpha.example") == ["I001"]                      # 이메일 검색
    assert ids("?type=인플루언서&tags=ugc") == ["I002"]
    assert ids("?q=columbus") == ["C002"]                           # 지역
    data = auth_client.get("/api/partners?type=케이터링").get_json()  # README 주소
    assert data["counts"] == {"total": 3, "withEmail": 2}
    assert auth_client.get("/api/partners/contacts?type=etc").status_code == 400


def test_subcategory_keywords():
    assert {"pool_beach", "vip"} <= set(importer.tag("케이터링", "Rooftop pool parties & VIP gala nights"))
    assert "pr_social" not in importer.tag("인플루언서", "Product reviews and premium content")   # pr은 단어 전체만
    assert {"pr_social", "agency"} <= set(importer.tag("인플루언서", "PR agency for lifestyle brands"))
    assert importer.tag("케이터링", "") == []


# ───────── 메일 생성 · 번역 ─────────

def test_email_generate_uses_prompt_and_rules(auth_client, ai):
    import_both(auth_client)
    fake = ai({"subject": "REBOOT x Festival Bar Co", "body": "Hello " * 160})
    res = auth_client.post("/api/partners/email/generate", json={
        "type": "케이터링", "partner": {"id": "C001", "name": "Festival Bar Co"}, "ideal": "페스티벌 현장 바",
        "prompt": "Write to {partner}. Product: {product}. Ideal: {ideal}.", "purpose": "현장 샘플링",
        "period": {"start": "2026-10-23", "end": "2026-11-22"}, "proposal": "1) 샘플링\n2) 메뉴 페어링", "requests": "일정 공유"})
    assert res.status_code == 200, res.get_json()
    out = res.get_json()
    assert out["subject"] == "REBOOT x Festival Bar Co" and out["ai"] is True and out["words"] == 160
    call = fake.calls[0]
    assert "Write to Festival Bar Co." in call["user"] and "Ideal: 페스티벌 현장 바." in call["user"]
    assert "{partner}" not in call["user"] and "REBOOT" in call["user"] and "Music festival" in call["user"]
    assert "150-200 words" in call["system"] and "excessive" in call["system"] and "2026-11-22" in call["user"]
    assert call["schema"]["required"] == ["subject", "body"] and call["part"] == "partners"


def test_email_generate_without_key(auth_client):
    res = auth_client.post("/api/partners/email/generate", json={"type": "인플루언서", "partner": {"name": "X"}})
    assert res.status_code == 503 and "OPENAI_API_KEY" in res.get_json()["error"]


def test_translate(auth_client, ai):
    fake = ai({"text": "안녕하세요 REBOOT #ad"})
    res = auth_client.post("/api/partners/translate", json={"text": "Hello REBOOT #ad", "to": "ko"})
    assert res.get_json()["text"] == "안녕하세요 REBOOT #ad"
    assert "hashtags" in fake.calls[0]["system"] and "Korean" in fake.calls[0]["system"]
    assert auth_client.post("/api/partners/translate", json={"text": "x", "to": "jp"}).status_code == 400
    assert auth_client.post("/api/partners/translate", json={"text": " ", "to": "en"}).status_code == 400


# ───────── 첨부파일 ─────────

def upload_att(client, name, content):
    return client.post("/api/partners/attachments", data={"file": (io.BytesIO(content), name)},
                       content_type="multipart/form-data")


def test_attachments_upload_list_download_delete(app, auth_client):
    res = upload_att(auth_client, "협업 브리프.pdf", b"%PDF-1.4 test")
    assert res.status_code == 201
    att = res.get_json()["items"][0]
    assert att["name"] == "협업 브리프.pdf" and att["size"] == 13
    folder = Path(app.config["DATA_DIR"]) / "uploads" / "partners" / "attachments"
    stored = list(folder.iterdir())
    assert len(stored) == 1 and stored[0].name.startswith(att["id"]) and stored[0].suffix == ".pdf"
    assert upload_att(auth_client, "../../evil.exe", b"MZ").status_code == 400
    assert upload_att(auth_client, "empty.txt", b"").status_code == 400
    listing = auth_client.get("/api/partners/attachments").get_json()
    assert [a["id"] for a in listing["items"]] == [att["id"]] and listing["maxTotalMB"] == 25
    dl = auth_client.get(f"/api/partners/attachments/{att['id']}")
    assert dl.status_code == 200 and dl.data == b"%PDF-1.4 test"
    for bad in ("..%2F..%2Freboot.db", "abc", "0" * 31 + "g"):
        assert auth_client.get(f"/api/partners/attachments/{bad}").status_code == 404
        assert auth_client.delete(f"/api/partners/attachments/{bad}").status_code == 404
    assert auth_client.delete(f"/api/partners/attachments/{att['id']}").status_code == 200
    assert list(folder.iterdir()) == []
    assert auth_client.delete(f"/api/partners/attachments/{att['id']}").status_code == 404


# ───────── 발송 ─────────

def test_send_guards(auth_client, monkeypatch):
    assert send(auth_client, confirm=False).status_code == 400
    assert send(auth_client, confirm="true").status_code == 400        # 문자열 'true'는 확인으로 보지 않음
    assert send(auth_client, owner=" ").status_code == 400
    assert send(auth_client, to="not-an-email").status_code == 400
    assert send(auth_client, to="a@x.example.com\r\nBcc: evil@x.example.com").status_code == 400
    res = send(auth_client)                                            # Gmail 미연결
    assert res.status_code == 503 and "설정 › 연결" in res.get_json()["error"]
    assert auth_client.get("/api/partners/history").get_json()["items"] == []
    assert auth_client.get("/api/health").get_json()["smtp"] is False


def test_send_success_saves_history_with_attachment(auth_client, gmail):
    import_both(auth_client)
    att = upload_att(auth_client, "brief.csv", b"a,b\n1,2\n").get_json()["items"][0]
    res = send(auth_client, attachment_ids=[att["id"]], subject="Hello\r\nBcc: x@y.z")
    assert res.status_code == 200, res.get_json()
    out = res.get_json()
    h = out["history"]
    assert out["ok"] is True and h["mailStatus"] == "발송 성공" and h["progress"] == "컨택 중" and h["kind"] == "자동"
    assert h["attachments"] == ["brief.csv"] and h["owner"] == "혁주" and h["partnerId"] == "C002"
    assert h["messageId"].startswith("<") and h["messageId"].endswith("@reboot-test.example.com>")
    smtp = FakeSMTP.instances[-1]
    assert smtp.login_args == ("team@reboot-test.example.com", "abcdefghijklmnop")
    msg = smtp.sent[0]
    assert msg["Message-ID"] == h["messageId"] and msg["To"] == "events@campus.example.com"
    assert "\n" not in msg["Subject"] and msg["Subject"].startswith("Hello")
    assert [p.get_filename() for p in msg.iter_attachments()] == ["brief.csv"]
    assert auth_client.get("/api/partners/overview").get_json()["awaiting"] == 1
    logs = auth_client.get("/api/worklog?part=partners").get_json()["items"]
    assert any("협업 메일 발송" in x["action"] and x["approval"] == "혁주" for x in logs)
    assert "events@campus.example.com" not in json.dumps(logs, ensure_ascii=False)   # 이메일은 가려서 기록


def test_send_failure_is_saved(auth_client, gmail):
    gmail(fail="auth")
    res = send(auth_client)
    assert res.status_code == 200
    out = res.get_json()
    assert out["ok"] is False and "로그인" in out["error"] and "abcd" not in out["error"]
    h = out["history"]
    assert h["mailStatus"] == "발송 실패" and h["messageId"] == "" and "발송 실패" in h["memo"]
    assert auth_client.get("/api/partners/overview").get_json()["failed"] == 1


def test_send_rejects_oversized_attachments(auth_client, gmail, monkeypatch):
    from tabs.partners import routes
    att = upload_att(auth_client, "big.zip", b"x" * 2048).get_json()["items"][0]
    monkeypatch.setattr(routes, "MAX_ATTACH_TOTAL", 1024)
    res = send(auth_client, attachment_ids=[att["id"]])
    assert res.status_code == 400 and "25MB" in res.get_json()["error"]
    assert send(auth_client, attachment_ids=["../../etc"]).status_code == 400


# ───────── 협업 가이드 ─────────

GUIDE_MD = """# REBOOT 인플루언서 협업 가이드

## 1. 한 줄 요약
> 21세 이상 팔로워에게 REBOOT 애프터파티 리추얼을 소개합니다.

## 2. 협업 개요
| 항목 | 내용 |
|---|---|
| 브랜드 | **REBOOT** |
| 기간 | [확인 필요: 기간] |

## 7. FTC 광고 표기
- 첫 줄에 #ad 표기
1. 초안 제출
"""


def test_guide_docx_attach_and_download(app, auth_client):
    from docx import Document
    res = auth_client.post("/api/partners/guide/docx", json={"markdown": GUIDE_MD, "type": "인플루언서",
                                                            "purpose": "제품 시딩·리뷰", "attach": True})
    assert res.status_code == 201, res.get_json()
    att = res.get_json()["attachment"]
    assert att["name"].startswith("REBOOT_인플루언서_협업가이드_") and att["name"].endswith(".docx") and att["source"] == "guide"
    folder = Path(app.config["DATA_DIR"]) / "uploads" / "partners" / "attachments"
    doc = Document(str(next(folder.iterdir())))
    sec = doc.sections[0]
    assert "REBOOT" in sec.header.paragraphs[0].text and "Confidential" in sec.footer.paragraphs[0].text
    assert "PAGE" in sec.footer._element.xml
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "REBOOT 인플루언서 협업 가이드" in text and "#ad" in text
    grid = [t for t in doc.tables if len(t.rows) == 3][0]
    assert grid.rows[0].cells[0].text == "항목" and 'w:fill="191C18"' in grid.rows[0].cells[0]._tc.xml
    assert grid.rows[1].cells[1].text == "REBOOT"                       # **굵게** 표시는 글자만
    assert any('w:fill="F5FBDF"' in t._tbl.xml for t in doc.tables)    # 한 줄 요약 콜아웃
    assert any(a["id"] == att["id"] for a in auth_client.get("/api/partners/attachments").get_json()["items"])
    again = auth_client.post("/api/partners/guide/docx", json={"markdown": GUIDE_MD, "type": "인플루언서", "attach": True}).get_json()
    assert again["replaced"] == [att["id"]]                              # 같은 날 같은 유형 가이드는 바꿔 넣음
    names = [a["name"] for a in auth_client.get("/api/partners/attachments").get_json()["items"]]
    assert names.count(att["name"]) == 1 and len(list(folder.iterdir())) == 1
    dl = auth_client.post("/api/partners/guide/docx", json={"markdown": GUIDE_MD, "type": "인플루언서", "attach": False})
    assert dl.status_code == 200 and dl.data[:2] == b"PK" and "wordprocessingml" in dl.content_type
    assert auth_client.post("/api/partners/guide/docx", json={"markdown": "  "}).status_code == 400


def test_guide_generate(auth_client, ai):
    fake = ai({"title": "REBOOT 케이터링 협업 가이드", "markdown": "## 1. 한 줄 요약\n> 요약"})
    res = auth_client.post("/api/partners/guide/generate", json={
        "type": "케이터링", "purpose": "현장 샘플링", "schedule": "11/1 킥오프", "must_include": "21세 이상",
        "deliverables": "현장 사진 10장", "banned": [{"term": "hangover cure", "alt": "after-party ritual"}]})
    out = res.get_json()
    assert res.status_code == 200 and out["markdown"].startswith("# REBOOT 케이터링 협업 가이드")
    call = fake.calls[0]
    for title in ("한 줄 요약", "협업 개요", "일정", "핵심 메시지", "결과물 및 제출 기준", "Do & Don't", "FTC 광고 표기", "승인 절차"):
        assert title in call["system"]
    assert "hangover cure" in call["user"] and "현장 샘플링" in call["user"]


def test_guide_generate_selected_products_only(auth_client, ai):
    cfg = auth_client.get("/api/partners/config").get_json()
    ids = [p["id"] for p in cfg["products"]]
    assert len(ids) >= 2 and all(p["name"] for p in cfg["products"])
    fake = ai({"title": "가이드", "markdown": "## 1. 한 줄 요약\n> 요약"})
    auth_client.post("/api/partners/guide/generate", json={"type": "인플루언서", "purpose": "제품 시딩·리뷰",
                                                           "product_ids": [ids[1]]})
    sent = json.loads(fake.calls[0]["user"])["brand_and_products"]["products"]
    assert [p["internal_name"] for p in sent] == [cfg["products"][1]["name"]]
    assert "Enjoy a brighter day with REBOOT" in fake.calls[0]["system"]
    # 고르지 않으면 전체
    auth_client.post("/api/partners/guide/generate", json={"type": "인플루언서", "purpose": "제품 시딩·리뷰"})
    assert len(json.loads(fake.calls[1]["user"])["brand_and_products"]["products"]) == len(ids)


# ───────── 이력 ─────────

def test_history_manual_edit_delete(auth_client):
    import_both(auth_client)
    res = auth_client.post("/api/partners/history", json={
        "partner_id": "C001", "activity": "전화", "progress": "협의 중", "owner": "효선", "memo": "단가 문의",
        "at": "2026-10-20T10:00:00+09:00"})
    assert res.status_code == 201
    h = res.get_json()
    assert h["kind"] == "수기" and h["partnerName"] == "Festival Bar Co" and h["partnerType"] == "케이터링"
    assert h["at"] == "2026-10-20T01:00:00Z" and h["mailStatus"] == ""
    assert auth_client.post("/api/partners/history", json={"partner_id": "C001", "progress": "끝"}).status_code == 400
    assert auth_client.post("/api/partners/history", json={"memo": "업체 없음"}).status_code == 400
    ed = auth_client.patch(f"/api/partners/history/{h['id']}", json={"progress": "협업 중", "mail_status": "회신 대기",
                                                                    "memo": "계약 진행", "subject": "a\nb"})
    assert ed.status_code == 200 and ed.get_json()["progress"] == "협업 중" and ed.get_json()["subject"] == "a b"
    assert auth_client.patch(f"/api/partners/history/{h['id']}", json={"mail_status": "읽음"}).status_code == 400
    assert auth_client.patch("/api/partners/history/9999", json={"memo": "x"}).status_code == 404
    ov = auth_client.get("/api/partners/overview").get_json()
    assert ov["contacted"] == 1 and ov["active"] == 1 and ov["activeBreakdown"]["협업 중"] == 1 and ov["awaiting"] == 1
    assert [x["id"] for x in auth_client.get("/api/partners/history?kind=수기").get_json()["items"]] == [h["id"]]
    assert auth_client.post("/api/partners/history/delete", json={"ids": []}).status_code == 400
    assert auth_client.post("/api/partners/history/delete", json={"ids": [h["id"]]}).get_json()["deleted"] == 1
    assert auth_client.get("/api/partners/history").get_json()["items"] == []


def test_history_restore_after_save_failure(auth_client):
    res = auth_client.post("/api/partners/history", json={"restore": {
        "partnerId": "C001", "partnerName": "Festival Bar Co", "activity": "메일 발송", "mailStatus": "발송 성공",
        "progress": "컨택 중", "owner": "혁주", "at": "2026-10-23T01:00:00Z", "to": "info@festbar.example.com",
        "subject": "Hi", "body": "Body", "messageId": "<x@reboot>", "attachments": ["a.pdf"]}})
    assert res.status_code == 201 and res.get_json()["kind"] == "자동" and res.get_json()["messageId"] == "<x@reboot>"


# ───────── 답장 확인 (IMAP) ─────────

def raw_msg(frm, to, subject, body, msgid=None, in_reply_to=None, headers=None, when=None):
    m = EmailMessage()
    m["From"], m["To"], m["Subject"] = frm, to, subject
    m["Message-ID"] = msgid or make_msgid(domain="partner.example.com")
    m["Date"] = format_datetime(when or datetime.now(timezone.utc))
    if in_reply_to:
        m["In-Reply-To"] = in_reply_to
        m["References"] = in_reply_to
    for k, v in (headers or {}).items():
        m[k] = v
    m.set_content(body)
    return m.as_bytes()


class FakeIMAP:
    def __init__(self, messages):
        from email import message_from_bytes
        self.raw = {str(i + 1).encode(): m for i, m in enumerate(messages)}
        self.parsed = {k: message_from_bytes(v) for k, v in self.raw.items()}
        self.calls, self.selected = [], None

    def login(self, user, password):
        self.user = user

    def list(self):
        return "OK", [b'(\\HasNoChildren) "/" "INBOX"', b'(\\All \\HasNoChildren) "/" "[Gmail]/&yATMtLz0rQDVaA-"']

    def select(self, box, readonly=False):
        self.selected = (box, readonly)
        return "OK", [b"2"]

    def uid(self, cmd, *args):
        self.calls.append((cmd,) + args)
        if cmd == "SEARCH":
            hits = []
            for k, m in self.parsed.items():
                ok, i = True, 0
                while i < len(args):
                    a = args[i]
                    if a == "HEADER":
                        ok &= args[i + 2].strip('"') in str(m.get(args[i + 1], ""))
                        i += 3
                    elif a == "FROM":
                        ok &= args[i + 1].strip('"').lower() in str(m.get("From", "")).lower()
                        i += 2
                    else:
                        i += 2   # SINCE 날짜는 무시
                if ok:
                    hits.append(k)
            return "OK", [b" ".join(hits)]
        if cmd == "FETCH":
            assert "BODY.PEEK" in args[1], "읽음 표시를 바꾸지 않게 BODY.PEEK로 읽어야 함"
            key = args[0] if isinstance(args[0], bytes) else str(args[0]).encode()
            return "OK", [(key + b" (UID BODY[] {1}", self.raw[key]), b")"]
        raise AssertionError(cmd)

    def logout(self):
        pass


def test_replies_check_with_fake_imap(auth_client, gmail, monkeypatch):
    import_both(auth_client)
    a = send(auth_client).get_json()["history"]
    b = send(auth_client, to="nobody@bounce.example.org", partner_id="C001", partner_name="Festival Bar Co").get_json()["history"]
    c = send(auth_client, to="quiet@alpha.example.com", partner_id="I001", partner_name="Alpha Talent Agency").get_json()["history"]
    later = datetime.now(timezone.utc) + timedelta(minutes=5)
    reply_id = "<reply-1@campus.example.com>"
    reply = raw_msg("Campus Team <events@campus.example.com>", "team@reboot-test.example.com", "Re: " + a["subject"],
                    "Thanks for reaching out! We are interested in November.\nCould you share samples?\n\n"
                    "On Thu, Oct 23, 2026 at 3:00 PM REBOOT <team@reboot-test.example.com> wrote:\n> Hello team,",
                    msgid=reply_id, in_reply_to=a["messageId"], when=later)
    bounce = raw_msg("Mail Delivery Subsystem <mailer-daemon@googlemail.com>", "team@reboot-test.example.com",
                     "Delivery Status Notification (Failure)", "Address not found: nobody@bounce.example.org",
                     headers={"X-Failed-Recipients": "nobody@bounce.example.org"}, when=later)
    own = raw_msg("REBOOT <team@reboot-test.example.com>", "events@campus.example.com", "Re: x", "our own follow-up",
                  in_reply_to=a["messageId"], when=later + timedelta(minutes=1))
    fake = FakeIMAP([own, reply, bounce])
    monkeypatch.setattr(mailer, "connect_imap", lambda: fake)
    res = auth_client.post("/api/partners/replies/check")
    assert res.status_code == 200, res.get_json()
    out = res.get_json()
    assert (out["checked"], out["replies"], out["bounces"]) == (3, 1, 1)
    assert fake.selected == ('"[Gmail]/&yATMtLz0rQDVaA-"', True)       # 전체보관함 · 읽기 전용
    items = {h["id"]: h for h in auth_client.get("/api/partners/history").get_json()["items"]}
    ha, hb, hc = items[a["id"]], items[b["id"]], items[c["id"]]
    assert ha["mailStatus"] == "회신 받음" and "[회신" in ha["memo"] and "Thanks for reaching out!" in ha["memo"]
    assert ha["reply"]["messageId"] == reply_id and "> Hello team" not in ha["reply"]["body"]
    assert "Could you share samples?" in ha["reply"]["body"] and ha["reply"]["estimated"] is False
    assert hb["mailStatus"] == "반송" and "nobody@bounce.example.org" in hb["memo"]
    assert hc["mailStatus"] == "발송 성공"
    assert auth_client.get("/api/partners/overview").get_json()["awaiting"] == 1

    # 회신 요약 (캐시) · AI 답장 초안
    monkeypatch.setenv("OPENAI_API_KEY", "sk-test-not-real")
    fake_ai = FakeAI({"core": "11월 샘플링에 관심", "requests": "샘플 요청", "mood": "긍정적", "next_action": "샘플 일정 회신"})
    monkeypatch.setattr("server.core.ai.openai_json", fake_ai)
    s1 = auth_client.post(f"/api/partners/history/{a['id']}/reply-summary").get_json()
    s2 = auth_client.post(f"/api/partners/history/{a['id']}/reply-summary").get_json()
    assert s1["summary"]["mood"] == "긍정적" and s1["cached"] is False and s2["cached"] is True and len(fake_ai.calls) == 1
    fake_ai.reply = {"subject": "ignored", "body": "Thank you for your interest."}
    d = auth_client.post(f"/api/partners/history/{a['id']}/reply-draft", json={"points": "샘플 20개 보낼 수 있음"}).get_json()
    assert d["subject"] == "Re: " + a["subject"] and d["body"] == "Thank you for your interest."
    assert "샘플 20개" in fake_ai.calls[-1]["user"] and "Could you share samples?" in fake_ai.calls[-1]["user"]

    # 같은 대화로 답장 → In-Reply-To · References, 진행 상태 '협의 중'
    r = send(auth_client, subject=d["subject"], body="Thank you for your interest.", reply_to=a["id"], partner_id="")
    assert r.status_code == 200 and r.get_json()["ok"] is True
    msg = FakeSMTP.instances[-1].sent[-1]
    assert msg["In-Reply-To"] == reply_id and a["messageId"] in msg["References"] and reply_id in msg["References"]
    nh = r.get_json()["history"]
    assert nh["activity"] == "답장 발송" and nh["progress"] == "협의 중" and nh["partnerId"] == "C002" and nh["replyOf"] == a["id"]
    orig =[h for h in auth_client.get("/api/partners/history").get_json()["items"] if h["id"] == a["id"]][0]
    assert orig["progress"] == "협의 중"


def test_reply_to_self_is_found(auth_client, gmail, monkeypatch):
    """내 Gmail → 내 Gmail로 보낸 테스트: 같은 주소에서 온 답장도 회신으로 잡고, 대시보드에서 보낸 우리 메일은 빼기."""
    me = "team@reboot-test.example.com"
    a = send(auth_client, to=me, partner_id="", partner_name="Self Test").get_json()["history"]
    later = datetime.now(timezone.utc) + timedelta(minutes=5)
    sent_copy = raw_msg(f"REBOOT <{me}>", me, a["subject"], "Hello team", msgid=a["messageId"], when=later)
    reply = raw_msg(f"HJ <{me}>", me, "Re: " + a["subject"], "Got it, thanks!", msgid="<self-reply@reboot-test.example.com>",
                    in_reply_to=a["messageId"], when=later + timedelta(minutes=1))
    monkeypatch.setattr(mailer, "connect_imap", lambda: FakeIMAP([sent_copy, reply]))
    out = auth_client.post("/api/partners/replies/check").get_json()
    assert out["replies"] == 1
    h = [x for x in auth_client.get("/api/partners/history").get_json()["items"] if x["id"] == a["id"]][0]
    assert h["mailStatus"] == "회신 받음" and "Got it, thanks!" in h["reply"]["body"]

    # 다른 업체에 보낸 메일에는 여전히 '내가 보낸 메일'을 회신으로 잡지 않음
    b = send(auth_client).get_json()["history"]
    own = raw_msg(f"REBOOT <{me}>", "events@campus.example.com", "Re: x", "our follow-up", in_reply_to=b["messageId"],
                  when=later + timedelta(minutes=2))
    monkeypatch.setattr(mailer, "connect_imap", lambda: FakeIMAP([own]))
    assert auth_client.post("/api/partners/replies/check").get_json()["replies"] == 0


def test_replies_check_needs_gmail(auth_client):
    assert auth_client.post("/api/partners/replies/check").status_code == 503


def test_reply_estimated_without_message_id(auth_client, gmail, monkeypatch):
    rec = auth_client.post("/api/partners/history", json={"restore": {
        "partnerName": "Old Partner", "activity": "메일 발송", "mailStatus": "발송 성공", "owner": "혁주",
        "at": (datetime.now(timezone.utc) - timedelta(days=2)).isoformat(), "to": "old@partner.example.com",
        "subject": "Collaboration with REBOOT", "body": "hi"}}).get_json()
    reply = raw_msg("old@partner.example.com", "team@reboot-test.example.com", "RE: Collaboration with REBOOT",
                    "Sounds good. Let's talk next week.")
    monkeypatch.setattr(mailer, "connect_imap", lambda: FakeIMAP([reply]))
    out = auth_client.post("/api/partners/replies/check").get_json()
    assert out["replies"] == 1
    h = auth_client.get("/api/partners/history").get_json()["items"][0]
    assert h["id"] == rec["id"] and h["mailStatus"] == "회신 받음" and "추정" in h["memo"] and h["reply"]["estimated"] is True


# ───────── 판매 실적 (더미) ─────────

def test_sales_placeholder_and_deterministic(auth_client):
    a = auth_client.get("/api/partners/sales").get_json()
    b = auth_client.get("/api/partners/sales").get_json()
    assert a == b and a["dummy"] is True and a["source"] == "placeholder"
    assert [c["name"] for c in sorted(a["creators"], key=lambda c: c["name"])][:2] == ["Creator A", "Creator B"]
    assert len(a["creators"]) == 12 and len(a["daily"]) == 30
    t = a["totals"]
    assert t["revenue"] == round(sum(c["revenue"] for c in a["creators"]), 2)
    assert t["orders"] == sum(d["orders"] for d in a["daily"]) and t["refunds"] > 0
    assert all(c["revenue"] == round(c["units"] * sales.UNIT_PRICE, 2) for c in a["creators"])
    one = auth_client.get(f"/api/partners/sales?start={a['end']}&end={a['end']}").get_json()
    assert len(one["daily"]) == 1 and one["start"] == one["end"] == a["end"]
    assert auth_client.get("/api/partners/sales?start=2026-13-01").status_code == 400


def test_sales_uses_imported_influencers_without_examples(auth_client):
    import_both(auth_client)
    data = auth_client.get("/api/partners/sales?start=2000-01-01&end=2100-01-01").get_json()
    assert data["source"] == "imported" and len(data["daily"]) == sales.DAYS
    assert sorted(c["name"] for c in data["creators"]) == ["Alpha Talent Agency", "Beta UGC Platform"]
    assert {c["code"] for c in data["creators"]} == {"ALPHATALEN10", "BETAUGCPLA10"}


def test_sales_values_fixed_per_day():
    r1 = sales.report(["Creator X"], date(2026, 10, 23), date(2026, 10, 1), date(2026, 10, 10))
    r2 = sales.report(["Creator X"], date(2026, 10, 30), date(2026, 10, 1), date(2026, 10, 10))
    assert r1["daily"] == r2["daily"] and r1["totals"] == r2["totals"]


# ───────── 예시 데이터 · 양식 ─────────

def test_example_data_and_templates(auth_client):
    out = auth_client.post("/api/partners/example").get_json()
    assert out["loaded"] == ["케이터링", "인플루언서"] and out["samples"] == 2
    items = auth_client.get("/api/partners/contacts").get_json()["items"]
    assert all(c["example"] and "example.com" in c["email"] and c["name"].startswith("예시)") for c in items)
    ov = auth_client.get("/api/partners/overview").get_json()
    assert ov["contacted"] == 0 and ov["awaiting"] == 0 and len(ov["partners"]) == 2 and all(p["sample"] for p in ov["partners"])
    assert auth_client.post("/api/partners/example").get_json()["loaded"] == []        # 두 번째는 그대로
    assert auth_client.get("/api/partners/sales").get_json()["source"] == "placeholder"  # 예시(IX)는 실적에서 제외
    csv_text = auth_client.get("/api/partners/template/catering.csv").get_data(as_text=True)
    assert "Company Name" in csv_text and "example.com" in csv_text
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(auth_client.get("/api/partners/template/influencer.xlsx").data))
    assert wb.active.cell(row=6, column=2).value == "업체명"
    assert auth_client.get("/api/partners/template/x.csv").status_code == 404


# ───────── 화면 JS 문법 ─────────

@pytest.mark.skipif(shutil.which("node") is None, reason="node 없음")
def test_partner_js_syntax():
    for f in sorted((ROOT / "tabs" / "partners" / "static").glob("*.js")):
        res = subprocess.run(["node", "--check", str(f)], capture_output=True, text=True, timeout=60)
        assert res.returncode == 0, f"{f.name}: {res.stderr}"


def test_guide_prompt_fixed_summary_and_no_age_key_message():
    from tabs.partners import prompts
    system, _ = prompts.guide_messages("인플루언서", "제품 시딩·리뷰", "", "", "", [], {}, "", "2026-10-23")
    assert f"> {prompts.GUIDE_SUMMARY}" in system and "{summary}" not in system
    assert "never copy the format notes" in system and "must not include an adults-21+" in system
    assert "hangover" not in prompts.GUIDE_SUMMARY.lower() and "숙취" not in prompts.GUIDE_SUMMARY
