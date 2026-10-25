"""FDA 경고장 수집 · AI 페르소나 — 인터넷·키 없이 확인 (외부 요청은 모두 가짜로 바꿈)."""
import io
import json
from datetime import date, timedelta
from pathlib import Path
from types import SimpleNamespace

from openpyxl import Workbook

from tabs.compliance import ai_review, refresh, warning_letters

LETTER = """WARNING LETTER  MARCS-CMS 612345  Product: Party Gummies
Your website makes claims such as: “Party Gummies prevent hangovers so you can drink all night” and
“Protects your liver from alcohol damage”. See section 201(g)(1) of the Act, 21 U.S.C. 321(g)(1)."""


def _xlsx():
    wb = Workbook()
    ws = wb.active
    ws.append(["Posted Date", "Letter Issue Date", "Company Name", "Issuing Office", "Subject"])
    recent = (date.today() - timedelta(days=10)).strftime("%m/%d/%Y")
    old = (date.today() - timedelta(days=900)).strftime("%m/%d/%Y")
    ws.append([recent, recent, "Party Gummies LLC", "Human Foods Program", "Unapproved New Drugs/Misbranded"])
    ws["C2"].hyperlink = "/inspections-compliance-enforcement-and-criminal-investigations/warning-letters/party-gummies-llc-612345-10012026"
    ws.append([recent, recent, "Pill Maker Inc", "Center for Drug Evaluation and Research", "CGMP/Finished Pharmaceuticals"])
    ws.append([old, old, "Old Juice Co", "Human Foods Program", "Juice HACCP/Adulterated"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _offline(monkeypatch):
    """자동 경로는 fda.gov 봇 차단처럼, 다른 출처는 0건으로."""
    def blocked(*a, **k):
        raise warning_letters.Blocked("fda.gov가 자동 요청을 막았습니다 (봇 차단).")
    monkeypatch.setattr(warning_letters, "_fetch", blocked)
    monkeypatch.setattr(refresh, "SOURCES", [refresh.SOURCES[0]])
    for k in ("FDA_DASHBOARD_USER", "FDA_DASHBOARD_KEY"):   # .env의 실제 키로 FDA를 부르지 않게
        monkeypatch.delenv(k, raising=False)


def test_extract_claims_skips_legal_text():
    quotes = warning_letters.extract_claims(LETTER)
    assert quotes == ["Party Gummies prevent hangovers so you can drink all night", "Protects your liver from alcohol damage"]


def test_xlsx_upload_keeps_recent_food_letters(auth_client):
    res = auth_client.post("/api/compliance/letters/xlsx", data={"file": (io.BytesIO(_xlsx()), "letters.xlsx")},
                           content_type="multipart/form-data")
    assert res.status_code == 200, res.get_json()
    out = res.get_json()
    assert out["letters"] == {"rows": 3, "food": 1, "added": 1}
    src = next(s for s in out["sources"] if s["id"] == "fda-warning-letters")
    assert [i["id"] for i in src["items"]] == ["612345"]
    assert src["items"][0]["url"].startswith("https://www.fda.gov/")


def test_xlsx_upload_rejects_other_files(auth_client):
    res = auth_client.post("/api/compliance/letters/xlsx", data={"file": (io.BytesIO(b"hello"), "a.xlsx")},
                           content_type="multipart/form-data")
    assert res.status_code == 400
    assert auth_client.post("/api/compliance/letters/xlsx", data={}, content_type="multipart/form-data").status_code == 400


def test_text_upload_becomes_casebook_candidates(auth_client):
    res = auth_client.post("/api/compliance/letters/text", json={
        "text": LETTER, "url": "https://www.fda.gov/x/party-gummies-llc-612345-10012026", "firm": "Party Gummies LLC", "issued": "2026-10-01"})
    assert res.status_code == 200, res.get_json()
    item = next(s for s in res.get_json()["sources"] if s["id"] == "fda-warning-letters")["items"][0]
    assert item["id"] == "612345" and len(item["quotes"]) == 2 and item["date"] == "2026-10-01"


def test_text_upload_input_checks(auth_client):
    assert auth_client.post("/api/compliance/letters/text", json={"text": "short"}).status_code == 400
    assert auth_client.post("/api/compliance/letters/text", json={"text": LETTER, "url": "https://evil.example/x"}).status_code == 400
    assert auth_client.post("/api/compliance/letters/text", json={"text": "x" * 100}).status_code == 400   # 인용 문구 없음


def test_refresh_reports_block_but_keeps_uploaded(auth_client, monkeypatch):
    _offline(monkeypatch)
    auth_client.post("/api/compliance/letters/text", json={"text": LETTER, "firm": "Party Gummies LLC"})
    out = auth_client.post("/api/compliance/monitor/refresh", json={}).get_json()
    src = next(s for s in out["sources"] if s["id"] == "fda-warning-letters")
    assert src["items"] and src["items"][0]["quotes"]
    errs = " ".join(e["error"] for e in out["errors"])
    assert "봇 차단" in errs and "Data Dashboard 키 없음" in errs


def test_ai_review_sends_persona(app, monkeypatch):
    sent = {}

    class Fake:
        def __init__(self, **kw):
            self.beta = SimpleNamespace(messages=SimpleNamespace(create=self.create))

        def create(self, **kw):
            sent.update(kw)
            body = {"verdict": "BLOCK", "confidence": "high", "summary": "s", "issues": [], "rewrite": "r"}
            return SimpleNamespace(stop_reason="end_turn", model="claude-test", content=[SimpleNamespace(type="text", text=json.dumps(body))])

    monkeypatch.setenv("ANTHROPIC_API_KEY", "test-key")
    monkeypatch.setattr(ai_review.anthropic, "Anthropic", Fake)
    with app.app_context():
        ai_review.review("Prevents hangovers.", persona="미국 광고 전문가 " + "가" * 700)
        texts = [b["text"] for b in sent["system"]]
        assert texts[0] == ai_review.SYSTEM and "never makes a verdict less strict" in texts[1]
        assert len(texts[1]) == len(ai_review.PERSONA_INTRO) + ai_review.MAX_PERSONA   # 600자에서 자름
        assert "cache_control" in sent["system"][-1]
        sent.clear()
        ai_review.review("Prevents hangovers.")
        assert len(sent["system"]) == 2   # 페르소나가 없으면 블록도 없음


def test_translation_cache_reuses_results(app, monkeypatch):
    """같은 문장은 두 번째부터 AI를 부르지 않음 (일괄 번역은 새 줄만 보냄)."""
    from tabs.compliance import translate
    calls = []

    def fake_call(system, content, schema, what):
        calls.append(content)
        if "items" in schema.get("properties", {}):
            lines = [json.loads(l) for l in content.split("\n")[1:]]
            return {"items": [{"i": x["i"], "text": "EN:" + x["ko"], "note_ko": ""} for x in lines]}, "claude-test"
        return {"text": "번역", "note_ko": ""}, "claude-test"

    monkeypatch.setattr(translate, "_call", fake_call)
    monkeypatch.setattr(translate, "available", lambda: True)
    with app.app_context():
        assert translate.to_korean("Say goodbye.")["text"] == "번역"
        again = translate.to_korean("Say goodbye.")
        assert again["cached"] is True and len(calls) == 1
        first = translate.to_english_batch(["가", "나"])
        assert [x["text"] for x in first["items"]] == ["EN:가", "EN:나"] and len(calls) == 2
        second = translate.to_english_batch(["가", "나", "다"])
        assert second["cached"] == 2 and len(calls) == 3 and '"다"' in calls[-1] and '"가"' not in calls[-1]
        assert translate.to_english("가")["cached"] is True and len(calls) == 3


def test_only_fda_urls_are_kept_or_fetched(app):
    """업로드 파일 속 임의 주소로 서버가 내부망을 읽지 않게 (SSRF 방지)."""
    ok = "https://www.fda.gov/inspections-compliance-enforcement-and-criminal-investigations/warning-letters/a-612345-10012026"
    assert warning_letters.safe_url(ok) == ok
    assert warning_letters.safe_url("https://datadashboard.fda.gov/oii/cd/firm.htm")
    for bad in ("https://127.0.0.1/internal", "http://www.fda.gov/x", "https://www.fda.gov.evil.example/x",
                "https://user@www.fda.gov/x", "https://www.fda.gov:8443/x", "file:///etc/passwd", "https://[::1/"):
        assert warning_letters.safe_url(bad) == "", bad
    assert warning_letters.safe_url("https://datadashboard.fda.gov/x", {warning_letters.LETTER_HOST}) == ""
    wb = Workbook()
    ws = wb.active
    ws.append(["Posted Date", "Letter Issue Date", "Company Name", "Issuing Office", "Subject", "URL"])
    recent = (date.today() - timedelta(days=5)).strftime("%m/%d/%Y")
    ws.append([recent, recent, "Evil Foods", "Human Foods Program", "Misbranded", "https://127.0.0.1/internal"])
    buf = io.BytesIO()
    wb.save(buf)
    rows, _ = warning_letters.parse_xlsx(buf.getvalue())
    assert rows and rows[0]["url"] == ""
    with app.app_context():
        try:
            warning_letters._fetch("https://127.0.0.1/internal")
            raise AssertionError("내부 주소를 읽으면 안 됨")
        except warning_letters.LetterError:
            pass


def test_failed_source_keeps_previous_items(app, monkeypatch):
    """한 출처가 실패해도 지난번 정상 자료를 지우지 않고 stale로 표시 · 출처별 갱신 시각."""
    def ok():
        return [{"id": "r1", "date": "2026-10-01", "title": "Jelly recall", "summary": ""}]

    def boom():
        raise RuntimeError("일시 오류")
    with app.app_context():
        monkeypatch.setattr(refresh, "SOURCES", [("openfda-food-recalls", "리콜", ok)])
        first = refresh.collect()
        refresh.write(first)
        assert first["sources"][0]["updatedAt"]
        monkeypatch.setattr(refresh, "SOURCES", [("openfda-food-recalls", "리콜", boom)])
        second = refresh.collect()
        src = second["sources"][0]
        assert src["stale"] is True and [i["id"] for i in src["items"]] == ["r1"]
        assert src["updatedAt"] == first["sources"][0]["updatedAt"]
        assert second["errors"] and "일시 오류" in second["errors"][0]["error"]
        refresh.write(first)
        before = refresh._current()["generatedAt"]
        after = refresh.update_source({"id": "fda-warning-letters", "label": "경고장", "items": []})
        assert after["generatedAt"] == before   # 출처 하나만 바꾸면 전체 수집 시각은 그대로
        assert next(s for s in after["sources"] if s["id"] == "fda-warning-letters")["updatedAt"]


def test_ai_usage_records_cost_and_budget_stops(app, monkeypatch):
    """AI 호출마다 토큰·예상 비용 기록 · 이번 달 한도를 넘으면 AI 기능만 멈춤 (GPT 검토 ⑭)."""
    from tabs.compliance import ai_usage
    res = SimpleNamespace(model="claude-opus-5-5", usage=SimpleNamespace(
        input_tokens=1_000_000, output_tokens=100_000, cache_creation_input_tokens=0, cache_read_input_tokens=0))
    with app.app_context():
        monkeypatch.setenv("AI_MONTHLY_BUDGET_USD", "10")
        assert ai_usage.over_budget() == ""
        ai_usage.record("compliance", "AI 2차 판정", res)          # 4 + 2 = 6달러
        s = ai_usage.summary()
        assert s["calls"] == 1 and abs(s["usd"] - 6.0) < 1e-6 and s["features"][0]["feature"] == "AI 2차 판정"
        assert ai_usage.over_budget() == ""
        ai_usage.record("cs", "CS 답장 초안", res)                 # 12달러 > 10달러
        assert "한도" in ai_usage.over_budget()
        monkeypatch.setenv("ANTHROPIC_API_KEY", "test-key")
        try:
            ai_review.review("Tastes like lime.")
            raise AssertionError("한도를 넘으면 AI를 부르지 않아야 함")
        except ai_review.ReviewError as e:
            assert e.status == 429
        monkeypatch.setenv("AI_MONTHLY_BUDGET_USD", "0")         # 0 = 한도 없음
        assert ai_usage.over_budget() == ""


def test_ai_review_has_few_shot_without_eval_sentences():
    """few-shot 예시는 시스템 프롬프트에 있고, '처음 보는 문장 시험' 문장은 예시에 넣지 않음 (시험이 외운 결과가 되지 않게)."""
    import re
    src = (Path(__file__).parent.parent / "static" / "data.js").read_text(encoding="utf-8")
    block = src[src.index("evalSet: ["):src.index("],", src.index("evalSet: ["))]
    evals = re.findall(r"text: '([^']+)'", block)
    assert len(evals) >= 25
    assert ai_review.FEW_SHOT in ai_review.SYSTEM and ai_review.SYSTEM.startswith(ai_review.SYSTEM_BASE)
    assert not [t for t in evals if t in ai_review.FEW_SHOT]
    assert {"checks", "uncertain"} <= set(ai_review.SCHEMA["required"])
