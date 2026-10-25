"""FDA 경고장 수집 → 규제 모니터링의 'FDA 경고장' 출처 + 사례집 후보(quotes).

fda.gov는 자동 요청을 봇 차단(abuse-detection)으로 막습니다 (2026-10-24 확인: 경고장 목록 XLSX·목록 페이지·본문 모두
차단 안내 페이지로 이동하거나 404). 차단을 우회하지 않고 되는 길만 씁니다.
  1) FDA Data Dashboard API (공식 · 무료 키 FDA_DASHBOARD_USER · FDA_DASHBOARD_KEY): 식품 분야 경고장 목록 (회사·날짜)
     → 본문 문구는 주지 않으므로 사례집 후보는 생기지 않음
  2) fda.gov 경고장 목록 XLSX 자동 받기 (정직한 User-Agent · 1회). 막히면 이유만 기록.
     하루 1회 작업에서는 새 식품 경고장 본문도 CRAWL_DELAY_SECONDS(기본 30초) 간격으로 최대 REFRESH_MAX_LETTERS통 읽음
  3) 사람이 올리기: 브라우저에서 받은 경고장 목록 XLSX 업로드 · 경고장 본문 붙여넣기 → 따옴표 안 문구 → 사례집 후보
수집한 문구는 사례집·규칙에 자동으로 들어가지 않습니다 ('사례집 후보' → 사람이 승인).

저장: DATA_DIR/parts_override/compliance/warning-letters.json (수집할 때마다 monitor-data.js의 'fda-warning-letters'로 합쳐짐)
원본: FDA 광고 콘솔 crawler.py · refresh.py (README 부록 B 4단계)
"""
import hashlib
import io
import json
import re
import time
from datetime import date, datetime, timedelta
from html.parser import HTMLParser
from pathlib import Path
from urllib.parse import urlparse

import requests
from flask import current_app

from server.envkeys import get as env

SOURCE_ID = "fda-warning-letters"
LABEL = "FDA 경고장 (식품 분야 전체 · 🍺 표시만 숙취·음주 관련)"
BASE = "https://www.fda.gov"
WL_XLSX = (BASE + "/inspections-compliance-enforcement-and-criminal-investigations/"
           "compliance-actions-and-activities/warning-letters/datatables-data?page&_format=xlsx")
DASHBOARD = "https://api-datadashboard.fda.gov/v1/compliance_actions"
UA = "REBOOT-compliance-monitor/0.2 (internal compliance research)"
MAX_XLSX = 10 * 1024 * 1024
MAX_TEXT = 200_000
MAX_ITEMS = 80   # 문구(quotes)가 있는 항목은 개수와 상관없이 남김

# 일반식품 기준: 예전 콘솔은 'Dietary Supplement|Unapproved New Drug'만 봤지만, 일반식품 경고장 제목
# (Food/Adulterated, Misbranded, Juice, Seafood 등)과 식품 담당 부서 이름도 함께 봄 (README 부록 B 4단계)
FOOD_SUBJECT = re.compile(r"dietary supplement|unapproved new drug|misbrand|food|juice|seafood|beverage|nutrition|labeling", re.I)
FOOD_OFFICE = re.compile(r"food safety|applied nutrition|cfsan|human foods|nutrition", re.I)
TOPIC = re.compile(r"hangover|intoxicat|alcohol|dihydromyricetin|\bDHM\b|hovenia|liver detox|after[\s-]party", re.I)
CLAIM_HINT = re.compile(r"hangover|alcohol|drink|sober|intoxicat|liver|detox|toxin|acetaldehyde|party", re.I)
QUOTE = re.compile(r"[“\"]([^”\"]{12,400})[”\"]")
LEGAL_NOISE = re.compile(r"U\.S\.C\.|C\.F\.R\.|\bCFR\b|section \d|Fed\. Reg\.|^new drug$|^dietary supplement$", re.I)


class Blocked(Exception):
    """fda.gov가 자동 요청을 막음 (봇 차단 안내 페이지 · 404 · XLSX가 아닌 응답)."""


class LetterError(Exception):
    """화면에 그대로 보여 줄 수 있는 입력 오류."""


# ───────── 저장 ─────────

def _path():
    return Path(current_app.config["DATA_DIR"]) / "parts_override" / "compliance" / "warning-letters.json"


def load():
    try:
        data = json.loads(_path().read_text(encoding="utf-8"))
        return data if isinstance(data.get("letters"), dict) else {"letters": {}}
    except (OSError, ValueError):
        return {"letters": {}}


def _save(data):
    p = _path()
    p.parent.mkdir(parents=True, exist_ok=True)
    tmp = p.with_suffix(".json.tmp")
    tmp.write_text(json.dumps(data, ensure_ascii=False, indent=1), encoding="utf-8")
    tmp.replace(p)


def _merge(data, item):
    """같은 id면 빈 칸만 채우고, 문구(quotes)는 새로 읽은 쪽을 씀."""
    cur = data["letters"].get(item["id"], {})
    merged = {**item, **{k: v for k, v in cur.items() if v and not item.get(k)}}
    if item.get("quotes"):
        merged["quotes"] = item["quotes"]
    data["letters"][item["id"]] = merged


def items(data=None):
    data = data or load()
    rows = sorted(data["letters"].values(), key=lambda x: x.get("date") or "", reverse=True)
    with_quotes = [r for r in rows if r.get("quotes")]
    rest = [r for r in rows if not r.get("quotes")][:max(0, MAX_ITEMS - len(with_quotes))]
    return sorted(with_quotes + rest, key=lambda x: x.get("date") or "", reverse=True)


def as_source():
    return {"id": SOURCE_ID, "label": LABEL, "items": items()}


# ───────── 경고장 목록 XLSX (자동 · 업로드 공통) ─────────

def _iso(v):
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    s = str(v or "").strip()
    for fmt in ("%m/%d/%Y", "%Y-%m-%d", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(s[:19], fmt).date().isoformat()
        except ValueError:
            continue
    return s[:10]


LETTER_HOST = "www.fda.gov"                              # 서버가 본문을 읽는 곳은 여기뿐
LINK_HOSTS = {LETTER_HOST, "datadashboard.fda.gov"}      # 화면 링크로 저장해도 되는 곳


def safe_url(url, hosts=LINK_HOSTS):
    """FDA 주소만 통과 (업로드한 XLSX 속 임의 주소로 서버가 내부망을 읽는 SSRF 방지 · 2026-10-24).
    https · 허용 호스트 · 기본 포트 · 사용자 정보 없음이 아니면 빈 문자열."""
    url = str(url or "").strip()
    try:
        p = urlparse(url)
        port = p.port
    except ValueError:
        return ""
    ok = p.scheme == "https" and (p.hostname or "").lower() in hosts and port in (None, 443) and not p.username and not p.password
    return url if ok else ""


def marcs_from_url(url):
    m = re.search(r"-(\d{5,7})-\d{8}/?$", url or "")
    return m.group(1) if m else None


def _lookback():
    try:
        days = int(env("REFRESH_LOOKBACK_DAYS", "365"))
    except ValueError:
        days = 365
    return (date.today() - timedelta(days=max(1, days))).isoformat()


def _summary(subject, office, company):
    hot = TOPIC.search(f"{subject} {company}")
    return " · ".join(x for x in [subject, office, "🍺 숙취·음주 관련 단어" if hot else ""] if x)[:400]


def parse_xlsx(content):
    """경고장 목록 XLSX → 식품 관련 경고장 dict 목록. 링크 칸이 없으면 셀 하이퍼링크에서 찾음."""
    from openpyxl import load_workbook
    if not content.startswith(b"PK"):
        raise LetterError("XLSX 파일이 아닙니다. FDA 경고장 페이지의 Excel 내보내기 파일을 올려 주세요.")
    try:
        ws = load_workbook(io.BytesIO(content), read_only=False).active
    except Exception:
        raise LetterError("XLSX 파일을 읽지 못했습니다.")
    header = [str(c.value or "").strip() for c in ws[1]]
    find = lambda p: next((i for i, h in enumerate(header) if re.search(p, h, re.I)), None)
    col = {k: find(p) for k, p in {"posted": r"posted", "issued": r"issue", "company": r"company|firm",
                                    "office": r"office", "subject": r"subject", "url": r"url|link"}.items()}
    if col["company"] is None or col["subject"] is None:
        raise LetterError(f"경고장 목록 형식이 아닙니다 (칸 이름: {', '.join(h for h in header if h)[:200]}).")
    since = _lookback()
    out, total = [], 0
    for cells in ws.iter_rows(min_row=2):
        def val(k):
            i = col[k]
            return cells[i].value if i is not None and i < len(cells) else None
        company = str(val("company") or "").strip()
        if not company:
            continue
        total += 1
        url = str(val("url") or "").strip() or next((c.hyperlink.target for c in cells if c.hyperlink and c.hyperlink.target), "")
        if url.startswith("/"):
            url = BASE + url
        subject, office = str(val("subject") or "").strip(), str(val("office") or "").strip()
        posted, issued = _iso(val("posted")), _iso(val("issued"))
        if not (FOOD_SUBJECT.search(subject) or FOOD_OFFICE.search(office)):
            continue
        if (posted or issued) and (posted or issued) < since:
            continue
        url = safe_url(url)
        marcs = marcs_from_url(url)
        out.append({"id": marcs or "x-" + hashlib.sha1(f"{company}|{issued}".encode()).hexdigest()[:10],
                    "date": issued or posted, "title": subject or "Warning Letter", "firm": company,
                    "url": url, "summary": _summary(subject, office, company),
                    "via": "xlsx"})
    return out, total


def import_xlsx(content):
    if len(content) > MAX_XLSX:
        raise LetterError("파일이 너무 큽니다 (10MB 이하).")
    rows, total = parse_xlsx(content)
    data = load()
    before = set(data["letters"])
    for r in rows:
        _merge(data, r)
    data["updatedAt"] = datetime.now().isoformat(timespec="seconds")
    _save(data)
    return {"rows": total, "food": len(rows), "added": len(set(data["letters"]) - before)}


def _fetch(url, timeout=60):
    """fda.gov 요청 1회. 봇 차단(abuse-detection 이동 · 짧은 404)은 Blocked로."""
    if not safe_url(url, {LETTER_HOST}):
        raise LetterError("fda.gov 주소만 읽습니다.")
    r = requests.get(url, headers={"User-Agent": UA}, timeout=timeout, allow_redirects=False)
    if r.is_redirect or (r.status_code in (403, 404) and len(r.content) < 200):
        # 2026-10-24: 302 → /apology_objects/abuse-detection-apology.html, 목록 주소는 10바이트 'Not found'
        raise Blocked("fda.gov가 자동 요청을 막았습니다 (봇 차단). 브라우저에서 경고장 목록 XLSX를 받아 [XLSX 올리기]로 올려 주세요.")
    r.raise_for_status()
    return r


def auto_xlsx():
    r = _fetch(WL_XLSX, timeout=120)
    if not r.content.startswith(b"PK"):
        raise Blocked("fda.gov가 XLSX 대신 다른 페이지를 돌려줬습니다 (봇 차단 가능성).")
    return import_xlsx(r.content)


# ───────── 경고장 본문 → 인용 문구 ─────────

def extract_claims(text):
    """본문에서 따옴표로 인용된 문구 (법 조항·짧은 낱말 제외). 숙취·음주 관련 문구가 있으면 그것만."""
    out = []
    for m in QUOTE.finditer(text):
        t = re.sub(r"\s+", " ", m.group(1)).strip()
        if LEGAL_NOISE.search(t) or t in out:
            continue
        out.append(t)
    related = [c for c in out if CLAIM_HINT.search(c)]
    return (related or out)[:30]


def add_text(text, url="", firm="", issued="", subject=""):
    """사람이 붙여 넣은 경고장 본문 → 인용 문구를 뽑아 그 경고장 항목의 quotes로 저장."""
    text = str(text or "")
    if len(text.strip()) < 40:
        raise LetterError("경고장 본문을 붙여 넣어 주세요 (브라우저에서 경고장을 열고 Ctrl+A → Ctrl+C).")
    if len(text) > MAX_TEXT:
        raise LetterError("본문이 너무 깁니다 (20만 자 이하).")
    url = url.strip()
    if url and not safe_url(url, {LETTER_HOST}):
        raise LetterError("경고장 주소는 https://www.fda.gov/ 로 시작해야 합니다.")
    quotes = extract_claims(text)
    if not quotes:
        raise LetterError("본문에서 따옴표로 인용된 광고 문구를 찾지 못했습니다.")
    m = re.search(r"Product:\s*([^\n]{3,120})", text)
    marcs = marcs_from_url(url) or (re.search(r"MARCS-CMS\s*(\d{5,7})", text) or [None, None])[1]
    data = load()
    lid = marcs or "t-" + hashlib.sha1(text[:2000].encode()).hexdigest()[:10]
    cur = data["letters"].get(lid, {})
    item = {"id": lid, "date": _iso(issued) or cur.get("date", ""), "title": subject.strip() or cur.get("title") or (m.group(1).strip() if m else "Warning Letter"),
            "firm": firm.strip() or cur.get("firm", ""), "url": url or cur.get("url", ""),
            "summary": cur.get("summary") or _summary(subject, "", firm), "quotes": quotes, "via": "text"}
    _merge(data, item)
    data["updatedAt"] = datetime.now().isoformat(timespec="seconds")
    _save(data)
    return {"id": lid, "quotes": len(quotes)}


class _Text(HTMLParser):
    def __init__(self):
        super().__init__()
        self.parts = []

    def handle_data(self, d):
        self.parts.append(d)


def _html_text(html):
    p = _Text()
    p.feed(html)
    return " ".join(p.parts)


def read_bodies(limit):
    """하루 1회 작업용: 문구가 아직 없는 식품 경고장 본문을 30초 간격으로 읽음. 차단되면 바로 멈춤."""
    try:
        delay = max(30.0, float(env("CRAWL_DELAY_SECONDS", "30")))   # fda.gov robots.txt Crawl-Delay 30
    except ValueError:
        delay = 30.0
    data = load()
    todo = [x for x in items(data) if safe_url(x.get("url"), {LETTER_HOST}) and not x.get("quotes") and not x.get("checked")][:limit]
    read = 0
    try:
        for i, x in enumerate(todo):
            if i:
                time.sleep(delay)
            body = _html_text(_fetch(x["url"]).text)   # 막히면 Blocked → 여기서 멈춤 (읽은 것까지는 저장)
            x["checked"] = date.today().isoformat()
            if TOPIC.search(body):
                x["quotes"] = extract_claims(body)
            _merge(data, x)
            read += 1
    finally:
        _save(data)
    return read


# ───────── FDA Data Dashboard API (공식 · 키 필요) ─────────

def dashboard():
    user, key = env("FDA_DASHBOARD_USER"), env("FDA_DASHBOARD_KEY")
    if not (user and key):
        return None
    body = {"start": 1, "rows": 200, "sort": "ActionTakenDate", "sortorder": "DESC",
            "filters": {"ProductType": ["Food/Cosmetics"], "ActionType": ["Warning Letter"], "ActionTakenDateFrom": [_lookback()]},
            "columns": ["LegalName", "ActionTakenDate", "CaseInjunctionID", "Center", "CountryName", "FirmProfile"]}
    r = requests.post(DASHBOARD, json=body, timeout=60, headers={
        "User-Agent": UA, "Content-Type": "application/json", "Authorization-User": user, "Authorization-Key": key})
    r.raise_for_status()
    j = r.json()
    if str(j.get("statuscode")) not in ("400", "200"):   # 이 API는 성공을 400으로 돌려줌 (공식 문서)
        raise RuntimeError(f"Data Dashboard 응답 {j.get('statuscode')}: {str(j.get('message'))[:80]}")
    data = load()
    for x in j.get("result") or []:
        cid = str(x.get("CaseInjunctionID") or "").strip()
        if not cid:
            continue
        _merge(data, {"id": "ca-" + cid, "date": _iso(x.get("ActionTakenDate")), "title": "Warning Letter · Food/Cosmetics",
                      "firm": str(x.get("LegalName") or "")[:160],
                      "url": safe_url(x.get("FirmProfile")),
                      "summary": " · ".join(str(v) for v in (x.get("Center"), x.get("CountryName"),
                                                              "🍺 숙취·음주 관련 단어" if TOPIC.search(str(x.get("LegalName") or "")) else "",
                                                              "제목·본문은 FDA 경고장 목록에서 회사 이름으로 확인") if v)[:400], "via": "dashboard"})
    data["updatedAt"] = datetime.now().isoformat(timespec="seconds")
    _save(data)
    return len(j.get("result") or [])


# ───────── 규제 모니터링 수집에서 부르는 곳 ─────────

def collect(bodies=False):
    """자동 경로를 모두 시도 → (항목, 실패 목록, 참고 안내). 자동 경로가 다 막혀도 올려 둔 항목은 그대로 돌려줌.
    목록을 Data Dashboard로 받았으면 fda.gov 차단은 실패가 아니라 '참고 안내'(화면에 출처 아래 회색 글씨)."""
    warns, note = [], ""
    n = None
    try:
        n = dashboard()
        if n is None:
            warns.append("Data Dashboard 키 없음 (FDA_DASHBOARD_USER · FDA_DASHBOARD_KEY)")
    except Exception as e:
        warns.append(f"Data Dashboard 실패: {type(e).__name__}: {str(e)[:100]}")
    try:
        auto_xlsx()
        if bodies:
            try:
                limit = max(1, int(env("REFRESH_MAX_LETTERS", "10")))
            except ValueError:
                limit = 10
            read_bodies(limit)
    except Blocked as e:
        # 목록은 Data Dashboard로 받았으면 실패가 아님 → 본문 문구만 직접 올리기로 넣으라는 안내
        if n:
            note = "경고장 목록은 FDA Data Dashboard로 정상 수집 · 본문 문구는 fda.gov 봇 차단으로 자동 수집 불가 → 필요한 경고장만 아래 '경고장 직접 올리기 ②'로"
        else:
            warns.append(str(e))
    except Exception as e:
        warns.append(f"fda.gov 경고장 목록 실패: {type(e).__name__}: {str(e)[:100]}")
    return items(), warns, note
