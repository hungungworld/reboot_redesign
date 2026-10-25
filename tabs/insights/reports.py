"""내려받기 — 주간 리포트(Word 1장) · 월간 리포트(Word) · 데이터(Excel). README '내려받기' 표 기준."""
import io
import json

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from server.models import WorkLog
from tabs.settings import xlsx
from .models import Comment, Review, ShelfItem, Video
from .service import EXPORT_HEAD, SOURCE_LABEL, competitor_export, competitors_view, latest_trends, monthly_view, shelf_dict, shelf_maps, weekly_view
from .util import BUCKET_LABEL, week_label

GRADE_COLOR = {"긴급": RGBColor(0xB7, 0x5B, 0x48), "기회": RGBColor(0x63, 0x7C, 0x13), "참고": RGBColor(0x62, 0x7A, 0xA5)}
FONT = "맑은 고딕"


def _doc():
    d = Document()
    for name in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3"):
        try:
            st = d.styles[name]
        except KeyError:
            continue
        st.font.name = "Malgun Gothic"
        rpr = st.element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = rpr.makeelement(qn("w:rFonts"), {})
            rpr.append(rfonts)
        rfonts.set(qn("w:eastAsia"), FONT)
    d.styles["Normal"].font.size = Pt(9.5)
    for s in d.sections:
        s.top_margin = s.bottom_margin = Cm(1.6)
        s.left_margin = s.right_margin = Cm(1.8)
    return d


def _p(d, text, bold=False, size=None, color=None, italic=False):
    p = d.add_paragraph()
    run = p.add_run(text)
    run.bold, run.italic = bold, italic
    if size:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    p.paragraph_format.space_after = Pt(3)
    return p


def _table(d, header, rows, widths=None):
    t = d.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    for i, h in enumerate(header):
        cell = t.rows[0].cells[i]
        cell.text = ""
        r = cell.paragraphs[0].add_run(str(h))
        r.bold = True
        r.font.size = Pt(8.5)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            r = cells[i].paragraphs[0].add_run("" if v is None else str(v))
            r.font.size = Pt(8.5)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    return t


def _usd(v):
    return f"${v:,.2f}" if v is not None else "—"


def _stock(v):
    return "품절" if v is False else "재고 있음" if v is True else "모름"


def _save(d):
    buf = io.BytesIO()
    d.save(buf)
    buf.seek(0)
    return buf


def weekly_docx(t=None):
    v = weekly_view(t)
    d = _doc()
    d.add_heading(f"REBOOT 주간 리포트 · {v['week']} ({v['weekLabel']})", level=1)
    _p(d, f"만든 날 {v['today']} · 시장 · 고객 인사이트 · 판단은 규칙, AI 문장은 참고용", size=8.5, italic=True,
       color=RGBColor(0x85, 0x89, 0x81))
    d.add_heading("이번 주 액션", level=2)
    if not v["actions"]:
        _p(d, "특이사항 없음 — 규칙에 걸린 신호가 없어요." if v["hasData"] else "수집된 데이터가 없어요. 데이터 업데이트를 먼저 실행하세요.")
    for a in v["actions"][:5]:
        p = d.add_paragraph()
        r = p.add_run(f"[{a['grade']}] ")
        r.bold, r.font.color.rgb = True, GRADE_COLOR[a["grade"]]
        p.add_run(a["fact"]).bold = True
        p.paragraph_format.space_after = Pt(1)
        _p(d, f"영향: {a['why']}")
        _p(d, f"할 일: {a['todo']}" + ("  (AI 문장 · 참고용)" if a.get("ai") else ""))
    s = v["shelf"]
    d.add_heading("진열대 요약", level=2)
    if s["total"]:
        cc = s.get("compChanges")
        cmp_txt = (f"지난 수집 대비 경쟁사 변화 {cc['total']}건 (가격 변동 {cc['price']} · 새 상품 {cc['new']})" if s["compare"] and cc
                   else "첫 수집이라 변화 비교는 다음 수집부터예요")
        _p(d, f"상품 {s['total']}개 · 경쟁 상품 품절 {s['oos']}개 · 자사 품절 {s['ownOos']}개 · {cmp_txt}.")
        if s.get("ownWeee"):
            o = s["ownWeee"]
            _p(d, f"Weee!에서 자사 1회분 단가 {_usd(o['perServing'])} — {o['total']}개 중 {o['rank']}번째로 저렴 ({o['bucketLabel']}).")
        rows = ShelfItem.query.order_by(ShelfItem.week.desc(), ShelfItem.source, ShelfItem.position).limit(200).all()
        latest = {}
        for r in rows:
            latest.setdefault(r.source, r.week)
        maps = shelf_maps()
        items = [shelf_dict(r, *maps) for r in rows if latest.get(r.source) == r.week][:10]
        _table(d, ["판매처", "상품", "가격", "1회분", "재고"],
               [[x["sourceLabel"], ("[자사] " if x["own"] else "") + x["title"][:60], _usd(x["price"]), _usd(x["perServing"]),
                 _stock(x["in_stock"])] for x in items], widths=[2, 9.5, 2, 2, 2])
    else:
        _p(d, "진열대 데이터가 없어요.")
    d.add_heading("새 리뷰", level=2)
    rv = v["reviews"]
    praise = ", ".join(f"{x['topic']} {x['count']}" for x in rv["praise"]) or "없음"
    comp = ", ".join(f"{x['topic']} {x['count']}" for x in rv["complaint"]) or "없음"
    _p(d, f"지난 7일 새로 수집한 리뷰 {rv['count']}개 (자사 {rv['own']} · 경쟁사 {rv['other']}) · 칭찬 상위: {praise} · 불만 상위: {comp}")
    d.add_heading("다가오는 시즌", level=2)
    if v["seasons"]:
        for se in v["seasons"]:
            ly = se.get("lastYear")
            extra = f" · 작년 같은 시기 '{ly['keyword']}' 검색 평균 대비 {ly['pct']:+.0f}%" if ly else ""
            _p(d, f"{se['name']} ({se['start']}, D-{se['dday']}){extra}")
    else:
        _p(d, "모니터링 중인 다가오는 시즌이 없어요 (설정 › 수집 설정).")
    d.add_heading("출처", level=2)
    srcs = [f"{SOURCE_LABEL.get(src, src)} {w['now']}" for src, w in s["weeks"].items()]
    _p(d, "진열대: " + (" · ".join(srcs) if srcs else "없음") + " · Amazon은 SerpApi(검색 상위 · 광고 제외), Weee!는 Firecrawl")
    _p(d, "리뷰: Amazon 대표 리뷰(SerpApi) · Weee!(공개 리뷰 페이지) · 검색 수요: Google Trends(SerpApi, 미국 12개월)")
    return _save(d), f"reboot-insights-weekly-{v['week']}.docx"


def monthly_docx(month=None, t=None):
    v = monthly_view(month, t)
    d = _doc()
    d.add_heading(f"REBOOT 월간 리포트 · {v['label']}", level=1)
    _p(d, "시장 · 고객 인사이트 · AI 결과는 참고용 (근거 데이터와 함께 확인)", size=8.5, italic=True, color=RGBColor(0x85, 0x89, 0x81))
    d.add_heading("이달의 요약", level=2)
    if v["summary"]:
        for line in v["summary"]:
            d.add_paragraph(line, style="List Bullet")
        _p(d, f"AI 요약 · {v['model']}", size=8, italic=True)
    else:
        _p(d, "AI 요약이 없어요. 월간 분석 실행(요약 · 전략)을 돌리면 채워져요 (OpenAI 키 필요).")
    d.add_heading("지표", level=2)
    m = v["metrics"]
    se = m.get("search")
    if not se:
        search_txt = "데이터 없음"
    elif se["changePct"] is None:
        search_txt = f"{se['value']}"
    else:
        search_txt = f"{se['value']} (지난달 {se['prev']} → {se['changePct']:+.0f}%)"
    _table(d, ["지표", "값"], [
        ["자사 품절 주수", "데이터 없음" if m["ownOosWeeks"] is None else f"{m['ownOosWeeks']}주"],
        ["경쟁사 가격 변동 횟수 (±10%)", "2주차부터" if m["priceChanges"] is None else f"{m['priceChanges']}회"],
        ["새 리뷰 수", f"자사 {m['newReviews']['own']} · 경쟁사 {m['newReviews']['other']}"],
        ["검색 지수 (지난달 대비)", search_txt],
    ], widths=[7, 9])
    d.add_heading("고객 인사이트", level=2)
    if v["insights"]:
        for i in v["insights"]:
            _p(d, i["title"], bold=True)
            if i.get("detail"):
                _p(d, i["detail"])
            if i.get("faq"):
                _p(d, "FAQ: " + " / ".join(i["faq"]))
            if i.get("expressions"):
                _p(d, "고객 표현: " + " / ".join(i["expressions"]))
    else:
        _p(d, "아직 고객 데이터가 없어요 (유튜브 수집 · 리뷰).")
    d.add_heading("경쟁사 동향", level=2)
    for line in v["competitorTrends"] or ["경쟁사 데이터가 없어요."]:
        d.add_paragraph(line, style="List Bullet")
    d.add_heading("전략 업데이트", level=2)
    st = v["strategy"]
    if st:
        if st.get("changes"):
            _p(d, "지난달 대비 바뀐 점", bold=True)
            for c in st["changes"]:
                d.add_paragraph(c, style="List Bullet")
        if st.get("personas"):
            _p(d, "페르소나", bold=True)
            _table(d, ["이름", "한 줄", "핵심 니즈", "쓸 메시지"],
                   [[p["name"], p["oneLine"], p["needs"], p["message"]] for p in st["personas"]], widths=[3, 4.5, 4.5, 4.5])
        sw = st.get("swot") or {}
        _p(d, "SWOT", bold=True)
        _table(d, ["강점 S", "약점 W", "기회 O", "위협 T"],
               [["\n".join(sw.get(k) or []) for k in ("S", "W", "O", "T")]], widths=[4.2, 4.2, 4.2, 4.2])
        _p(d, f"타깃: {st.get('target') or '—'}")
        _p(d, f"포지셔닝: {st.get('positioning') or '—'}")
    else:
        _p(d, "전략 업데이트가 없어요. 월간 분석 실행(요약 · 전략)을 돌리면 페르소나 · SWOT · 타깃 · 포지셔닝이 채워져요.")
    return _save(d), f"reboot-insights-monthly-{v['month']}.docx"


def data_xlsx(t=None):
    wb = xlsx.new_workbook()
    maps = shelf_maps()
    shelf = [shelf_dict(r, *maps) for r in ShelfItem.query.order_by(ShelfItem.week.desc(), ShelfItem.source, ShelfItem.position).all()]
    xlsx.add_sheet(wb, "진열대(주차별)", ["주차", "기간", "판매처", "순위", "상품", "브랜드", "가격($)", "회분", "회분 근거", "1회분($)", "구간",
                                     "별점", "리뷰 수", "재고", "자사", "추적", "주소"],
                   [[x["week"], week_label(x["week"]), x["sourceLabel"], x["position"], x["title"], x["brand"], x["price"],
                     x["servings"], {"manual": "직접 수정", "parsed": x["servingsRule"], "default": "확인 필요(1회분으로 계산)"}[x["servingsSource"]],
                     x["perServing"], BUCKET_LABEL.get(x["bucket"], ""), x["rating"], x["reviews"], _stock(x["in_stock"]),
                     x["own"], x["tracked"], x["link"]] for x in shelf], widths={5: 50, 17: 40})
    reviews = Review.query.order_by(Review.collected_at.desc()).all()
    xlsx.add_sheet(wb, "구매자 리뷰", ["출처", "상품", "브랜드", "자사", "별점", "제목", "내용", "주제", "감정", "한국어 요약", "분류", "날짜"],
                   [[r.source, r.product_title, r.brand, r.own, r.rating, r.title, r.text, ", ".join(r.topic_list()),
                     {"+": "칭찬", "-": "불만", "0": "중립"}.get(r.sentiment, r.sentiment), r.summary_ko,
                     "AI" if r.classified_by == "ai" else "규칙", r.date_text] for r in reviews], widths={2: 40, 7: 70, 10: 50})
    vids = {v.id: v for v in Video.query.all()}
    comments = Comment.query.order_by(Comment.month.desc(), Comment.likes.desc()).all()
    voc = [["유튜브", c.month, vids[c.video_id].title if c.video_id in vids else "", c.text, c.likes, "예" if c.relevant else "아니요",
            "예" if c.is_question else "", c.topic, "AI" if c.classified_by == "ai" else "규칙"] for c in comments]
    voc += [["리뷰", r.week, r.product_title, r.text, None, "예", "", ", ".join(r.topic_list()),
             "AI" if r.classified_by == "ai" else "규칙"] for r in reviews]
    xlsx.add_sheet(wb, "VOC(유튜브·리뷰)", ["출처", "월/주차", "영상/상품", "내용", "좋아요", "관련", "질문", "주제", "분류"], voc,
                   widths={3: 40, 4: 70})
    series, meta = latest_trends()
    xlsx.add_sheet(wb, "검색 트렌드", ["키워드", "날짜(주 시작)", "지수 (0~100)", "출처", "수집 주차"],
                   [[k, p["date"], p["value"], meta[k]["source"], meta[k]["week"]] for k, pts in series.items() for p in pts])
    comp = competitors_view()
    xlsx.add_sheet(wb, "경쟁사 후보·점수", ["구분", "브랜드", "브랜드 추정", "신호 점수", "같은 판매처", "고객 언급", "가격대 비슷", "주시", "추천", "근거", "상품"],
                   [[x["tierLabel"], x["brand"], x["brandGuessed"], x["score"], x["signals"]["sameSeller"], x["signals"]["mention"],
                     x["signals"]["similarPrice"], x["watched"], x["recommended"], x["reason"],
                     " / ".join(p["title"][:60] for p in x["products"])] for col in comp["columns"].values() for x in col],
                   widths={10: 50, 11: 60})
    xlsx.add_sheet(wb, "경쟁사 분석", [label for _, label in EXPORT_HEAD],
                   [[r[k] for k, _ in EXPORT_HEAD] for r in competitor_export()], widths={4: 40, 11: 40, 12: 30, 13: 40, 14: 40})
    wk = weekly_view(t)
    xlsx.add_sheet(wb, "이번 주 액션", ["주차", "등급", "규칙", "사실", "영향", "할 일", "AI 문장", "규칙에 걸린 신호"],
                   [[wk["week"], a["grade"], a["ruleLabel"], a["fact"], a["why"], a["todo"], a["ai"],
                     json.dumps(a["facts"], ensure_ascii=False)[:2000]] for a in wk["actions"]], widths={4: 50, 5: 50, 6: 50, 8: 60})
    logs = WorkLog.query.filter_by(part="insights").order_by(WorkLog.id.desc()).limit(500).all()
    xlsx.add_sheet(wb, "작업 로그", ["일시(UTC)", "구분", "작업", "상태", "승인", "실행", "상세"],
                   [[r.at.strftime("%Y-%m-%d %H:%M") if r.at else "", r.kind, r.action, r.status, r.approval, r.user_name,
                     (r.detail or "")[:2000]] for r in logs], widths={7: 70})
    return xlsx.to_bytes(wb), f"reboot-insights-data-{wk['today']}.xlsx"
