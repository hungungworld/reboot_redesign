"""시장 · 고객 인사이트 — 규칙 엔진 · 1회분 단가 · 수집기(가짜 응답) · 데이터 업데이트 · 화면 API · 내려받기 (인터넷 · 실제 키 없이)."""
import io
import json
from datetime import date, timedelta

import pytest

from tabs.insights import http, rules, service, update
from tabs.insights.util import bucket, parse_servings, review_month, week_key

TODAY = date(2026, 10, 23)


class FakeRes:
    def __init__(self, status=200, data=None):
        self.status_code = status
        self._data = data

    def json(self):
        if self._data is None:
            raise ValueError("no json")
        return self._data


# ───────── 가짜 외부 응답 ─────────

def _weekly_points(spike=True):
    """2025-10-26부터 52주, 평균 40 근처 · 작년 핼러윈 다음 주(2025-11-02)에 90."""
    start = date(2025, 10, 26)
    out = []
    for i in range(52):
        d = start + timedelta(weeks=i)
        v = 40 + (i % 5)
        if spike and d == date(2025, 11, 2):
            v = 90
        out.append({"date": d.strftime("%b %d, %Y"), "timestamp": str(int((d - date(1970, 1, 1)).total_seconds())),
                    "values": [{"query": "hangover", "value": str(v), "extracted_value": v}]})
    return {"interest_over_time": {"timeline_data": out}}


SEARCH = {
    "hangover": {"organic_results": [
        {"asin": "B0SPONSOR1", "title": "Sponsored hangover thing", "sponsored": True, "extracted_price": 9.99},
        {"asin": "B0COMP0001", "title": "Flyby Hangover Recovery Aid, 30 Count", "price": "$29.99", "extracted_price": 29.99,
         "rating": 4.3, "reviews": 5400, "link": "https://www.amazon.com/dp/B0COMP0001"},
        {"asin": "B0COMP0002", "title": "Korean Hangover Jelly Stick 5 pack*4 boxes", "extracted_price": 40.0, "rating": 4.1,
         "reviews": 120},
        {"asin": "B0COMP0003", "title": "Haejang soup instant kit", "extracted_price": 12.0, "reviews": 30},
    ]},
    "korean hangover": {"organic_results": [
        {"asin": "B0COMP0002", "title": "Korean Hangover Jelly Stick 5 pack*4 boxes", "extracted_price": 40.0, "reviews": 120},
        {"asin": "B0COMP0004", "title": "Liquid I.V. Hydration Multiplier - 16 Count", "extracted_price": 25.0, "reviews": 90000},
    ]},
}
PRODUCTS = {
    "B0COMP0001": {"product_results": {"title": "Flyby Hangover Recovery Aid, 30 Count", "brand": "Visit the Flyby Store",
                                       "extracted_price": 29.99, "stock": "In Stock", "rating": 4.3, "reviews": 5400},
                   "about_item": ["Supports recovery", "30 capsules"],
                   "reviews_information": {"summary": {"text": "Customers like the convenience."},
                                           "authors_reviews": [{"title": "Works", "text": "It worked great after a party, easy to carry.",
                                                                "rating": 5, "date": "October 3, 2026"}]}},
    "B0COMP0004": {"product_results": {"title": "Liquid I.V. Hydration Multiplier - 16 Count", "brand": "Liquid I.V.",
                                       "extracted_price": 25.0, "stock": "In Stock"}},
    "B0COMP0002": {"product_results": {"title": "Korean Hangover Jelly Stick 5 pack*4 boxes", "extracted_price": 40.0}},
    "B0OWNASIN1": {"product_results": {"title": "Aldicom Hangover Jelly 1 pack*5 pack", "extracted_price": 15.0,
                                       "stock": "Currently unavailable."},
                   "reviews_information": {"authors_reviews": [
                       {"title": "Bad taste", "text": "The taste is too bitter and the flavor is gross.", "rating": 2},
                       {"title": "Meh", "text": "Bitter aftertaste, did not like the taste.", "rating": 1}]}},
}
WEEE_MD = """# Hangover

[![Condition Hangover Drink](https://img.weeecdn.com/a.jpg)](https://www.sayweee.com/en/product/Condition-Hangover-Drink/12345)

[Condition Hangover Drink 100ml x 10](https://www.sayweee.com/en/product/Condition-Hangover-Drink/12345)

$19.99$24.99

[![Aldicom Jelly](https://img.weeecdn.com/b.jpg)](https://www.sayweee.com/en/product/Aldicom-Jelly/67890)

[Aldicom Hangover Jelly 1 pack*5 pack](https://www.sayweee.com/en/product/Aldicom-Jelly/67890)

$12.99

Sold out

[Haejang Soup](https://www.sayweee.com/en/product/Haejang-Soup/11111)

$8.99
"""
YT_SEARCH = {"items": [{"id": {"videoId": "vid00000001"}, "snippet": {"title": "Korean hangover drink review", "channelTitle": "K-Snack Lab",
                                                                    "channelId": "UC123", "publishedAt": "2026-09-01T00:00:00Z"}},
                       {"id": {"videoId": "vid00000002"}, "snippet": {"title": "Small video", "channelTitle": "Tiny",
                                                                    "channelId": "UC999", "publishedAt": "2026-09-02T00:00:00Z"}}]}
YT_STATS = {"items": [{"id": "vid00000001", "statistics": {"viewCount": "50000", "likeCount": "900", "commentCount": "4"}},
                      {"id": "vid00000002", "statistics": {"viewCount": "500", "commentCount": "1"}}]}
YT_COMMENTS = {"items": [
    {"snippet": {"topLevelComment": {"id": f"c{i}", "snippet": {"textOriginal": t, "likeCount": 10 - i,
                                                                 "publishedAt": "2026-09-03T00:00:00Z"}}}}
    for i, t in enumerate(["Where can I buy this in the US?", "Where to buy the korean hangover drink online?",
                           "Does it actually work for a bad hangover?", "great video bro",
                           "Is this available at Costco or Walmart?", "Can I order this jelly on Amazon?",
                           "Where do you find this in stores?", "anyone else extremely hungover right now?"])]}


@pytest.fixture()
def fake_net(monkeypatch):
    calls = []

    def fake_get(url, params=None, headers=None, timeout=25):
        p = params or {}
        calls.append(("GET", url, dict(p)))
        if "serpapi.com" in url:
            if p.get("api_key") == "bad-serp-key-000":
                return FakeRes(401, {"error": "Invalid API key. Your API key should be here: https://serpapi.com/manage-api-key"})
            eng = p.get("engine")
            if eng == "amazon":
                return FakeRes(200, SEARCH.get(p["k"], {"organic_results": []}))
            if eng == "amazon_product":
                return FakeRes(200, PRODUCTS.get(p["asin"], {"error": "Google hasn't returned any results for this query."}))
            if eng == "google_trends":
                return FakeRes(200, _weekly_points(spike=p["q"] == "hangover"))
        if "googleapis.com/youtube" in url:
            if url.endswith("/search"):
                return FakeRes(200, YT_SEARCH)
            if url.endswith("/videos"):
                return FakeRes(200, YT_STATS)
            if url.endswith("/commentThreads"):
                if p.get("videoId") == "vid00000001":
                    return FakeRes(200, YT_COMMENTS)
                return FakeRes(403, {"error": {"errors": [{"reason": "commentsDisabled"}]}})
        raise AssertionError(f"예상하지 못한 호출: {url}")

    def fake_post(url, json=None, headers=None, timeout=90):
        calls.append(("POST", url, dict(json or {})))
        if "firecrawl.dev" in url:
            assert headers["Authorization"].startswith("Bearer ")
            page = (json or {}).get("url", "")
            if "/social/review/more/12345" in page:  # Condition만 리뷰 1개 (Aldicom은 0개)
                if "json" in json["formats"]:
                    return FakeRes(200, {"success": True, "data": {"json": {"reviews": [
                        {"rating": 5, "date": "Oct 17, 2026", "text": "Works great, no headache next morning"}]}}})
                return FakeRes(200, {"success": True, "data": {"markdown": "# Reviews (1)\n\nKimi\nVerified Buyer\nWorks great"}})
            if "/social/review/more/" in page:
                return FakeRes(200, {"success": True, "data": {"markdown": "# Reviews (0)\n\nPost a review"}})
            return FakeRes(200, {"success": True, "data": {"markdown": WEEE_MD}})
        raise AssertionError(f"예상하지 못한 호출: {url}")

    monkeypatch.setattr(http, "get", fake_get)
    monkeypatch.setattr(http, "post", fake_post)
    return calls


@pytest.fixture()
def fixed_today(monkeypatch):
    monkeypatch.setattr(service, "today", lambda: TODAY)
    return TODAY


@pytest.fixture()
def sync_jobs(monkeypatch):
    """데이터 업데이트를 뒤가 아니라 바로 실행 (테스트용).
    jobs.run_background는 시작하면 Thread, 이미 실행 중이면 None을 돌려주므로 None이 아닌 값을 돌려줌."""
    def run_now(app, name, fn, *a, **k):
        fn(*a, **k)
        return "started"
    monkeypatch.setattr(update.jobs, "run_background", run_now)


def _all_keys(monkeypatch):
    monkeypatch.setenv("SERPAPI_KEY", "serp-test-key-1234")
    monkeypatch.setenv("FIRECRAWL_API_KEY", "fc-test-key-1234567890")
    monkeypatch.setenv("YOUTUBE_API_KEY", "AIza-test-key-5678")


def _own_asin_config(client):
    cfg = client.get("/api/settings/collect").get_json()["config"]
    cfg["tracked_asins"] = [{"asin": "B0OWNASIN1", "label": "알디콤 젤리", "own": True}]
    cfg["trend_keywords"] = ["hangover", "hangover drink"]
    cfg["search_sets"] = {"A": ["korean hangover drink review"], "B": [], "K": []}
    assert client.put("/api/settings/collect", json={"config": cfg}).status_code == 200


# ───────── 1회분 단가 ─────────

@pytest.mark.parametrize("text, n", [
    ("5 pack*4 boxes", 20), ("Pack of 6", 6), ("Hangover jelly 10 sticks (min 2)", 10), ("1 pack*5 pack", 5),
    ("21 ml × 4 Sachets", 4), ("숙취해소 젤리 20팩", 20), ("12 Count (Pack of 2)", 24), ("Morning Recovery 3.4 fl oz, 6 Pack", 6),
    ("(min 2) hangover drink", None), ("DHM 1000 mg", None),
])
def test_parse_servings(text, n):
    assert parse_servings(text)[0] == n


def test_buckets_and_month_rules():
    assert [bucket(2.99), bucket(3.0), bucket(5.99), bucket(6.0), bucket(None)] == ["low", "mid", "mid", "high", None]
    assert review_month(date(2026, 11, 5)) == "2026-10" and review_month(date(2026, 11, 8)) == "2026-11"
    assert week_key(TODAY) == "2026-W43"


# ───────── 규칙 엔진 (화면과 같은 모양의 고정 데이터) ─────────

def _item(key, title, price, source="amazon", own=False, in_stock=True):
    return {"key": key, "title": title, "price": price, "source": source, "own": own, "in_stock": in_stock}


def _ctx(**kw):
    ctx = {"today": TODAY, "shelf_now": [], "shelf_prev": None, "trends": {}, "questions": [], "own_complaints": [],
           "seasons": [{"name": "Halloween", "start": "2026-10-31", "monitor": True},
                       {"name": "Thanksgiving Eve", "start": "2026-11-25", "monitor": True},
                       {"name": "Super Bowl", "start": "2027-02-14", "monitor": True}]}
    ctx.update(kw)
    return ctx


def test_rules_no_data_means_no_actions():
    assert rules.weekly_actions(_ctx()) == []


def test_rules_own_out_of_stock_needs_season_within_45_days():
    now = [_item("amazon:OWN", "Aldicom Jelly", 15, own=True, in_stock=False)]
    acts = rules.weekly_actions(_ctx(shelf_now=now))
    assert [a["rule"] for a in acts] == ["own_oos_season"] and acts[0]["grade"] == "긴급"
    assert "Halloween" in acts[0]["why"] and acts[0]["facts"]["남은 날"] == 8
    far = _ctx(shelf_now=now, seasons=[{"name": "Super Bowl", "start": "2027-02-14", "monitor": True}])
    assert rules.weekly_actions(far) == []
    off = _ctx(shelf_now=now, seasons=[{"name": "Halloween", "start": "2026-10-31", "monitor": False}])
    assert rules.weekly_actions(off) == []


def test_rules_price_changes_and_new_entries():
    prev = [_item("amazon:A", "Flyby 30 Count", 30.0), _item("amazon:B", "Liquid IV 16 Count", 25.0),
            _item("amazon:C", "Jelly 10 sticks", 20.0), _item("amazon:OWN", "Aldicom", 15.0, own=True)]
    now = [_item("amazon:A", "Flyby 30 Count", 24.0),            # -20% → 인하 긴급
           _item("amazon:B", "Liquid IV 16 Count", 28.0),         # +12% → 인상 참고
           _item("amazon:C", "Jelly 10 sticks", 21.0),            # +5% → 없음
           _item("amazon:OWN", "Aldicom", 9.0, own=True),         # 자사 가격 변동은 경쟁사 규칙에서 제외
           _item("amazon:D", "New Hangover Patch", 12.0),         # 신규
           _item("weee:1", "Weee item", 10.0, source="weee")]     # Weee!는 지난 묶음이 없어 비교 안 함
    acts = {a["rule"]: a for a in rules.weekly_actions(_ctx(shelf_now=now, shelf_prev=prev))}
    assert set(acts) == {"price_cut", "price_up", "new_entry"}
    assert acts["price_cut"]["grade"] == "긴급" and "-20%" in acts["price_cut"]["fact"]
    assert acts["price_up"]["grade"] == "참고" and "+12%" in acts["price_up"]["fact"]
    assert acts["new_entry"]["grade"] == "참고" and "1개" in acts["new_entry"]["fact"]
    order = [a["grade"] for a in rules.weekly_actions(_ctx(shelf_now=now, shelf_prev=prev))]
    assert order == sorted(order, key=lambda g: rules.GRADE_ORDER[g])
    # 첫 주(지난 묶음 없음)는 비교하지 않음
    assert rules.weekly_actions(_ctx(shelf_now=now, shelf_prev=None)) == []


def _lined(key, line, source, in_stock=True, left=None, own=False):
    return dict(_item(key, f"{line} 10 sticks", 20.0, source=source, own=own, in_stock=in_stock),
                line={"id": line, "label": line}, stockLeft=left)


def test_rules_competitor_stock_signals():
    now = [_lined("amazon:A", "상쾌환 · 환", "amazon", in_stock=False), _lined("weee:A", "상쾌환 · 환", "weee"),     # Weee!엔 있음 → 한 채널
           _lined("amazon:B", "컨디션 · 젤리", "amazon", in_stock=False), _lined("weee:B", "컨디션 · 젤리", "weee", in_stock=False),
           _lined("amazon:C", "Flyby · 캡슐", "amazon", left=8), _lined("amazon:D", "Cheers · 캡슐", "amazon", left=40),
           _lined("amazon:OWN", "알디콤 · 젤리", "amazon", left=3, own=True)]                                     # 자사는 경쟁 규칙 제외
    acts = {a["rule"]: a for a in rules.weekly_actions(_ctx(shelf_now=now))}
    assert set(acts) == {"comp_oos", "comp_oos_multi", "low_stock"}
    assert acts["comp_oos"]["grade"] == "기회" and "상쾌환 · 환 — Amazon에서 품절" in acts["comp_oos"]["fact"]
    assert "입찰" in acts["comp_oos"]["todo"] and rules.COMPLIANCE in acts["comp_oos"]["todo"]
    assert acts["comp_oos_multi"]["grade"] == "긴급" and "Amazon · Weee!에서 모두 품절" in acts["comp_oos_multi"]["fact"]
    assert acts["low_stock"]["grade"] == "참고" and "8개 남았어요" in acts["low_stock"]["fact"]
    assert acts["low_stock"]["todo"].startswith("경쟁 제품 품절에 대비하세요")
    # 지난 수집에도 품절이면 'N주째' · 지난 품절 → 이번 재고 있음 = 재입고
    prev = [_lined("amazon:B", "컨디션 · 젤리", "amazon", in_stock=False), _lined("weee:B", "컨디션 · 젤리", "weee", in_stock=False),
            _lined("amazon:C", "Flyby · 캡슐", "amazon", in_stock=False)]
    acts = {a["rule"]: a for a in rules.weekly_actions(_ctx(shelf_now=now, shelf_prev=prev))}
    assert "계속 품절" in acts["comp_oos_multi"]["fact"] and "성과분석" in acts["comp_oos_multi"]["todo"]
    assert acts["restock"]["grade"] == "참고" and "Flyby" in acts["restock"]["fact"] and "입찰가" in acts["restock"]["todo"]
    assert rules.COMPLIANCE not in acts["restock"]["todo"]


def test_comment_rules_drop_chatter_and_unknown_questions():
    from tabs.insights import classify
    noise = classify.classify_comment("anyone else extremely hungover right now?")
    assert noise == {"relevant": False, "is_question": False, "topic": ""}
    assert classify.classify_comment("Where can I buy this jelly in the US?")["topic"] == "buy"
    assert classify.classify_comment("Is it better than Liquid IV for hangovers?")["topic"] == "compare"
    odd = classify.classify_comment("Why is the hangover cure in the thumbnail blue?")
    assert odd["relevant"] and not odd["is_question"] and odd["topic"] == ""  # 주제를 모르면 질문 아님 ('기타' 없음)
    assert not classify.is_relevant("I went drinking last night and feel bad")  # 숙취 잡담
    assert classify.video_on_topic("Korean hangover drink review") and not classify.video_on_topic("My crazy weekend vlog")
    assert classify.video_on_topic("Aldicom unboxing", ["aldicom"])
    assert "other" not in classify.QUESTION_LABEL


def test_rules_top_question_needs_volume_and_skips_other():
    few = _ctx(questions=[{"key": "buy", "label": "어디서 사나요 (구매처)", "count": 4, "examples": []}])
    assert "top_question" not in {a["rule"] for a in rules.weekly_actions(few)}
    thin = _ctx(questions=[{"key": "buy", "label": "어디서 사나요 (구매처)", "count": 5, "examples": []}]
                + [{"key": k, "label": k, "count": 5, "examples": []} for k in ("works", "when", "safety", "taste", "price", "ingredients", "compare")])
    assert "top_question" not in {a["rule"] for a in rules.weekly_actions(thin)}  # 5/40 = 12.5% < 15%
    other = _ctx(questions=[{"key": "other", "label": "기타 질문", "count": 68, "examples": []},
                            {"key": "taste", "label": "맛이 어떤가요", "count": 6, "examples": []}])
    top = next(a for a in rules.weekly_actions(other) if a["rule"] == "top_question")
    assert top["facts"]["질문 주제"] == "맛이 어떤가요"


def test_rules_season_search_lift_questions_and_complaints():
    pts = [{"date": (date(2025, 10, 26) + timedelta(weeks=i)).isoformat(), "value": 90 if i == 1 else 40} for i in range(52)]
    lift = rules.season_lift(pts, "2026-10-31")
    assert lift and lift["date"] == "2025-11-02" and lift["pct"] > 100
    flat = [dict(p, value=40) for p in pts]
    assert rules.season_lift(flat, "2026-10-31")["pct"] == 0
    ctx = _ctx(trends={"hangover": pts, "flat": flat},
               questions=[{"label": "어디서 사나요 (구매처)", "count": 7, "examples": ["Where can I buy this?"]}],
               own_complaints=[{"topic": "맛·향", "count": 3, "examples": ["too bitter"]}])
    acts = rules.weekly_actions(ctx)
    by = {a["rule"]: a for a in acts}
    assert by["season_search"]["grade"] == "기회" and "'hangover'" in by["season_search"]["fact"]
    assert by["top_question"]["grade"] == "기회" and "7건" in by["top_question"]["fact"]
    assert by["own_complaint"]["grade"] == "긴급" and "맛·향" in by["own_complaint"]["fact"]
    assert acts[0]["rule"] == "own_complaint"  # 긴급이 먼저
    # Thanksgiving Eve(33일 뒤)는 30일 밖이라 기회 아님
    assert all("Thanksgiving" not in a["fact"] for a in acts)


# ───────── 화면 API (빈 상태) ─────────

def test_insights_requires_login(client):
    assert client.get("/api/insights/weekly").status_code == 401


def test_insights_page_loads_part_files(auth_client):
    html = auth_client.get("/dashboard/insights").get_data(as_text=True)
    assert "/parts/insights/insights.js" in html and "/parts/insights/insights.css" in html


def test_empty_states(auth_client, fixed_today):
    w = auth_client.get("/api/insights/weekly").get_json()
    assert w["hasData"] is False and w["actions"] == [] and w["week"] == "2026-W43"
    assert w["seasons"][0]["name"] == "Halloween" and w["seasons"][0]["dday"] == 8
    m = auth_client.get("/api/insights/monthly").get_json()
    assert m["month"] == "2026-10" and m["summary"] == [] and m["strategy"] is None
    assert m["metrics"]["ownOosWeeks"] is None and m["metrics"]["priceChanges"] is None
    for tab in ("shelf", "competitors", "reviews", "audience", "trends"):
        assert auth_client.get(f"/api/insights/data/{tab}").status_code == 200, tab
    assert auth_client.get("/api/insights/data/nope").status_code == 404
    assert auth_client.get("/api/insights/trends").get_json() == {"keywords": [], "updatedAt": None}
    assert auth_client.get("/api/insights/personas").get_json() == []
    assert auth_client.get("/api/insights/monthly?month=2026-13").status_code == 400


def test_update_plan_reports_missing_keys(auth_client):
    plan = auth_client.get("/api/insights/update/plan").get_json()
    by = {i["id"]: i for i in plan["items"]}
    assert by["amazon"]["ready"] is False and by["amazon"]["missing"] == ["SERPAPI_KEY"]
    assert by["reviews"]["ready"] is True and "규칙" in by["reviews"]["note"]
    assert by["monthly"]["ready"] is False
    assert by["amazon"]["cost"]["serpapi"] == 2 + 6  # 검색어 2 + 상세 6
    assert plan["limits"]["serpapi"] == 250
    res = auth_client.post("/api/insights/update", json={"items": ["amazon"]})
    assert res.status_code == 400 and "SerpApi" not in res.get_json()["error"] and "Amazon" in res.get_json()["error"]
    assert auth_client.post("/api/insights/update", json={"items": []}).status_code == 400
    assert auth_client.post("/api/insights/update", json={"items": ["hack"]}).status_code == 400


# ───────── 데이터 업데이트 (가짜 응답) ─────────

def test_update_job_collects_everything(auth_client, monkeypatch, fake_net, fixed_today, sync_jobs):
    _all_keys(monkeypatch)
    _own_asin_config(auth_client)
    items = ["weee", "amazon", "trends", "youtube", "reviews", "voc", "actions"]
    res = auth_client.post("/api/insights/update", json={"items": items})
    assert res.status_code == 202, res.get_json()
    st = auth_client.get("/api/insights/update/status").get_json()
    run = st["run"]
    assert run["status"] == "done", run["results"]
    assert set(run["results"]) == set(items)
    assert run["usage"]["serpapi"] >= 2 + 2 and run["usage"]["firecrawl"] == 1 + 2 + 5 and run["usage"]["youtube"] == 100 + 1 + 1
    # 키가 결과 · 로그에 남지 않음
    logs = auth_client.get("/api/worklog?part=insights").get_json()["items"]
    dumped = json.dumps(logs, ensure_ascii=False) + json.dumps(st, ensure_ascii=False)
    assert "serp-test-key" not in dumped and "fc-test-key" not in dumped and "AIza-test-key" not in dumped
    kinds = {(i["kind"], i["action"]) for i in logs}
    assert ("수집", "Amazon 진열대 · 리뷰 수집") in kinds and ("수집", "Weee! 진열대 · 리뷰 수집") in kinds
    assert ("승인 · 선정", "이번 주 액션 선정 (규칙)") in kinds
    # 광고 제외 · soup 제외 · 같은 ASIN은 한 번
    amazon_calls = [c for c in fake_net if c[0] == "GET" and c[2].get("engine") == "amazon"]
    assert [c[2]["k"] for c in amazon_calls] == ["hangover", "korean hangover"]
    trend_calls = [c[2]["q"] for c in fake_net if c[0] == "GET" and c[2].get("engine") == "google_trends"]
    assert trend_calls == ["hangover", "hangover drink"]  # 키워드마다 따로

    shelf = auth_client.get("/api/insights/data/shelf").get_json()
    titles = [x["title"] for x in shelf["items"]]
    assert not any("Sponsored" in t or "soup" in t.lower() for t in titles)
    keys = [x["key"] for x in shelf["items"]]
    assert keys.count("amazon:B0COMP0002") == 1 and "amazon:B0OWNASIN1" in keys
    jelly = next(x for x in shelf["items"] if x["key"] == "amazon:B0COMP0002")
    assert jelly["servings"] == 20 and jelly["perServing"] == 2.0 and jelly["bucket"] == "low"
    own = next(x for x in shelf["items"] if x["key"] == "amazon:B0OWNASIN1")
    assert own["own"] and own["in_stock"] is False and own["tracked"] and own["servings"] == 5
    weee = [x for x in shelf["items"] if x["source"] == "weee"]
    assert {x["title"] for x in weee} == {"Condition Hangover Drink 100ml x 10", "Aldicom Hangover Jelly 1 pack*5 pack"}
    wo = next(x for x in weee if x["own"])
    assert wo["price"] == 12.99 and wo["in_stock"] is False
    # Weee! 리뷰: 리뷰 페이지마다 수를 보고, 있는 상품만 JSON 추출 · 작성자 이름은 저장 안 함
    rv = auth_client.get("/api/insights/data/reviews").get_json()
    wr = [r for r in rv["items"] if r["source"] == "weee"]
    assert len(wr) == 1 and wr[0]["sourceLabel"] == "Weee!" and wr[0]["rating"] == 5 and "Kimi" not in json.dumps(wr)
    review_pages = [c[2] for c in fake_net if c[0] == "POST" and "/social/review/more/" in c[2].get("url", "")]
    assert sorted(p["formats"][0] for p in review_pages) == ["json", "markdown", "markdown"]
    flyby = next(x for x in shelf["items"] if x["key"] == "amazon:B0COMP0001")
    assert flyby["brand"] == "Flyby" and flyby["detailed"] and flyby["in_stock"] is True

    week = auth_client.get("/api/insights/weekly").get_json()
    rules_hit = {a["rule"]: a for a in week["actions"]}
    assert {"own_oos_season", "own_complaint", "season_search", "top_question"} <= set(rules_hit)
    assert "price_cut" not in rules_hit and week["shelf"]["compare"] is False  # 첫 주는 비교 안 함
    assert rules_hit["top_question"]["facts"]["질문 주제"] == "어디서 사나요 (구매처)"
    assert all(a["ai"] is False for a in week["actions"]) and week["aiReady"] is False
    assert week["shelf"]["ownWeee"]["perServing"] == 2.6

    aud = auth_client.get("/api/insights/data/audience").get_json()
    assert aud["process"]["found"] == 2 and aud["process"]["kept"] == 1
    assert aud["questions"][0]["key"] == "buy" and aud["questions"][0]["count"] == 5
    assert aud["creators"][0]["channel"] == "K-Snack Lab"
    assert all("great video" not in c["text"] and "anyone else" not in c["text"] for c in aud["comments"])

    rv = auth_client.get("/api/insights/data/reviews").get_json()
    own_row = next(b for b in rv["table"] if b["own"])
    assert own_row["topics"]["맛·향"]["neg"] == 2
    assert rv["summaries"][0]["text"] == "Customers like the convenience."

    tr = auth_client.get("/api/insights/trends").get_json()
    assert [k["keyword"] for k in tr["keywords"]] == ["hangover", "hangover drink"]
    assert len(tr["keywords"][0]["points"]) == 52 and tr["updatedAt"]
    nxt = auth_client.get("/api/insights/data/trends").get_json()["next"][0]
    assert nxt["name"] == "Halloween" and nxt["opportunity"] is True

    comp = auth_client.get("/api/insights/data/competitors").get_json()
    brands = {x["brand"]: x for col in comp["columns"].values() for x in col}
    assert "Flyby" in brands and brands["Flyby"]["tier"] == "category"
    assert comp["ownFound"] is True


def test_update_bad_key_fails_without_leaking(auth_client, monkeypatch, fake_net, fixed_today, sync_jobs):
    monkeypatch.setenv("SERPAPI_KEY", "bad-serp-key-000")
    assert auth_client.post("/api/insights/update", json={"items": ["amazon", "trends"]}).status_code == 202
    run = auth_client.get("/api/insights/update/status").get_json()["run"]
    assert run["status"] == "failed"
    assert "키를 거절" in run["results"]["amazon"]["message"]
    logs = auth_client.get("/api/worklog?part=insights").get_json()["items"]
    assert logs[0]["status"] == "실패"
    assert "bad-serp-key" not in json.dumps(logs, ensure_ascii=False) + json.dumps(run, ensure_ascii=False)


def test_update_refuses_while_running(auth_client, monkeypatch):
    monkeypatch.setattr(update.jobs, "is_running", lambda name: True)
    res = auth_client.post("/api/insights/update", json={"items": ["reviews"]})
    assert res.status_code == 409


# ───────── 2주차: 가격 변동 · 신규 · 수정 ─────────

def _seed_two_weeks(app):
    from server.extensions import db
    from tabs.insights.models import ShelfItem
    with app.app_context():
        for week, price in (("2026-W42", 30.0), ("2026-W43", 24.0)):
            db.session.add(ShelfItem(week=week, source="amazon", key="amazon:B0COMP0001", position=1,
                                     title="Flyby Hangover Recovery Aid, 30 Count", price=price, servings=30, in_stock=True))
        db.session.add(ShelfItem(week="2026-W43", source="amazon", key="amazon:B0NEW00001", position=2,
                                 title="Brand New Hangover Patch (Pack of 6)", price=18.0, servings=6, in_stock=True))
        db.session.add(ShelfItem(week="2026-W43", source="amazon", key="amazon:B0NOSERVE1", position=3,
                                 title="Mystery hangover powder", price=9.0, in_stock=None))
        db.session.commit()


def test_second_week_changes_overrides_and_edits(app, auth_client, fixed_today):
    _seed_two_weeks(app)
    week = auth_client.get("/api/insights/weekly").get_json()
    by = {a["rule"]: a for a in week["actions"]}
    assert by["price_cut"]["grade"] == "긴급" and "$30.00 → $24.00" in by["price_cut"]["fact"]
    assert by["new_entry"]["grade"] == "참고" and week["shelf"]["compare"] is True
    cc = week["shelf"]["compChanges"]
    assert cc["price"] == 1 and cc["new"] == 2 and cc["threat"] == 3 and cc["demo"] is False
    comp = auth_client.get("/api/insights/data/competitors").get_json()
    assert cc["total"] == len(comp["changes"])  # 주간 요약과 경쟁사 탭이 같은 숫자
    shelf = auth_client.get("/api/insights/data/shelf").get_json()
    mystery = next(x for x in shelf["items"] if x["key"] == "amazon:B0NOSERVE1")
    assert mystery["servingsSource"] == "default" and mystery["perServing"] == 9.0  # 못 읽으면 1회분으로 (확인 필요)
    flyby = next(x for x in shelf["items"] if x["key"] == "amazon:B0COMP0001")
    assert flyby["change"]["pricePct"] == -20.0
    # 회분 수 직접 수정 → 모든 주차에 적용
    res = auth_client.put("/api/insights/servings", json={"key": "amazon:B0NOSERVE1", "servings": 3})
    assert res.status_code == 200
    mystery = next(x for x in auth_client.get("/api/insights/data/shelf").get_json()["items"] if x["key"] == "amazon:B0NOSERVE1")
    assert mystery["servingsSource"] == "manual" and mystery["perServing"] == 3.0 and mystery["bucket"] == "mid"
    assert auth_client.put("/api/insights/servings", json={"key": "amazon:B0NOSERVE1", "servings": 0}).status_code == 400
    assert auth_client.put("/api/insights/servings", json={"key": "nope", "servings": 2}).status_code == 404
    assert auth_client.put("/api/insights/servings", json={"key": "amazon:B0NOSERVE1", "servings": None}).get_json()["servings"] is None
    # 한 줄 고치기: 가격 · 재고 · 자사 표시
    res = auth_client.patch(f"/api/insights/shelf/{mystery['id']}", json={"price": "10.5", "in_stock": False, "own": True})
    out = res.get_json()
    assert res.status_code == 200 and out["price"] == 10.5 and out["in_stock"] is False and out["own"] is True
    assert out["editedBy"] == "테스터"
    assert auth_client.patch(f"/api/insights/shelf/{mystery['id']}", json={"price": "abc"}).status_code == 400
    assert auth_client.patch(f"/api/insights/shelf/{mystery['id']}", json={}).status_code == 400
    m = auth_client.get("/api/insights/monthly").get_json()
    assert m["metrics"]["priceChanges"] == 1 and m["metrics"]["ownOosWeeks"] == 1


def test_line_edit_confirm_and_weekly_shelf_fields(app, auth_client, fixed_today):
    _seed_two_weeks(app)
    shelf = auth_client.get("/api/insights/data/shelf").get_json()
    patch = next(x for x in shelf["items"] if x["key"] == "amazon:B0NEW00001")
    assert patch["line"]["confirmed"] is False and shelf["lines"] and shelf["forms"]
    assert auth_client.put("/api/insights/lines", json={"key": "amazon:B0NEW00001", "brand": "PatchCo", "form": "other"}).status_code == 200
    assert auth_client.put("/api/insights/lines", json={"key": "amazon:B0NEW00001", "form": "rock"}).status_code == 400
    assert auth_client.put("/api/insights/lines", json={"key": "nope", "brand": "x"}).status_code == 404
    patch = next(x for x in auth_client.get("/api/insights/data/shelf").get_json()["items"] if x["key"] == "amazon:B0NEW00001")
    assert patch["line"]["label"] == "PatchCo · 기타" and patch["line"]["confirmed"] and patch["brand"] == "PatchCo"
    assert auth_client.post("/api/insights/lines/confirm").get_json()["confirmed"] == 3
    s = auth_client.get("/api/insights/weekly").get_json()["shelf"]
    assert s["unconfirmed"] == 0 and s["unknownStock"] == 1 and s["lowLimit"] == 20
    assert set(s["lastCollect"]) == set()  # 실행 기록이 없으면 빈 값


def test_watch_toggle_marks_competitor(app, auth_client, fixed_today):
    _seed_two_weeks(app)
    assert auth_client.put("/api/insights/watch", json={"brand": "Flyby", "watch": True}).get_json()["watch"] == ["Flyby"]
    comp = auth_client.get("/api/insights/data/competitors").get_json()
    flyby = next(x for col in comp["columns"].values() for x in col if x["brand"] == "Flyby")
    assert flyby["watched"] is True
    log = auth_client.get("/api/worklog?part=insights").get_json()["items"][0]
    assert log["kind"] == "승인 · 선정" and log["approval"] == "테스터"
    assert auth_client.put("/api/insights/watch", json={"brand": "Flyby", "watch": False}).get_json()["watch"] == []
    assert auth_client.put("/api/insights/watch", json={"brand": ""}).status_code == 400


# ───────── 구매자 리뷰 표 · CSV 올리기 ─────────

def test_review_table_rule_topics(app, auth_client):
    from tabs.insights import store
    rows = [{"source": "weee", "product_key": "weee:1", "product_title": "Aldicom Jelly", "brand": "Aldicom", "rating": 2.0,
             "date_text": "Oct 1, 2026", "title": "", "text": "The taste is too sweet", "own": True},
            {"source": "weee", "product_key": "weee:2", "product_title": "Condition Drink", "brand": "Condition", "rating": 5.0,
             "date_text": "Oct 2, 2026", "title": "Great", "text": "Works great and easy to carry", "own": False}]
    with app.app_context():
        assert store.save_reviews(rows, "2026-W43") == 2
        assert store.save_reviews(rows, "2026-W43") == 0  # 같은 리뷰는 한 번
    rv = auth_client.get("/api/insights/data/reviews").get_json()
    own = next(b for b in rv["table"] if b["own"])
    assert own["topics"]["맛·향"]["neg"] == 1
    cond = next(b for b in rv["table"] if b["brand"] == "Condition")
    assert cond["topics"]["효과"]["pos"] == 1 and cond["topics"]["복용 편의"]["pos"] == 1
    # Weee! 리뷰 CSV 올리기는 없앰 (Weee! 리뷰는 진열대 수집 때 자동으로)
    assert auth_client.post("/api/insights/reviews/upload", json={"csv": "text\nx\n"}).status_code in (404, 405)
    assert auth_client.get("/api/insights/templates/weee-reviews.csv").status_code == 404


def test_review_slack_share(app, auth_client, monkeypatch):
    """미리 보기는 보내지 않음 · 키 없음/모양 틀림은 안내 · 보내면 기록 · 웹훅 주소는 응답 · 로그에 안 남음."""
    from tabs.insights import store
    from tabs.insights.models import Review
    with app.app_context():
        store.save_reviews([{"source": "amazon", "product_key": "amazon:A1", "product_title": "ALDICOM A Liquid 5-Pack", "rating": 2.0,
                             "date_text": "Oct 1, 2026", "text": "Too pricey <b>& sweet</b>", "own": True}], "2026-W43")
        rid = Review.query.first().id
    sent = []

    class Res:
        status_code = 200
    monkeypatch.setattr(http, "post", lambda url, json=None, headers=None, timeout=90: sent.append((url, json)) or Res())
    url = f"/api/insights/reviews/{rid}/slack"
    pv = auth_client.post(url, json={"preview": True}).get_json()
    assert pv["ready"] is False and "*[알디콤 A · Amazon]* ★★☆☆☆ 불만" in pv["text"] and "&lt;b&gt;&amp; sweet" in pv["text"]
    assert "대응: 대응 필요" in pv["text"] and sent == []
    res = auth_client.post(url, json={"note": "논의 필요"})
    assert res.status_code == 400 and "SLACK_WEBHOOK_URL" in res.get_json()["error"]
    monkeypatch.setenv("SLACK_WEBHOOK_URL", "https://example.com/hook")
    assert "모양" in auth_client.post(url, json={}).get_json()["error"] and sent == []
    hook = "https://hooks.slack.com/services/T000/B000/secretTEST1234"
    monkeypatch.setenv("SLACK_WEBHOOK_URL", hook)
    res = auth_client.post(url, json={"note": "논의 필요"})
    assert res.status_code == 200 and res.get_json()["slackBy"]
    assert sent[0][0] == hook and "💬" in sent[0][1]["text"] and "논의 필요" in sent[0][1]["text"]
    rv = auth_client.get("/api/insights/data/reviews").get_json()
    assert rv["slackReady"] and rv["cs"][0]["slackBy"] and rv["items"][0]["slackAt"]
    logs = json.dumps(auth_client.get("/api/worklog?part=insights").get_json(), ensure_ascii=False)
    assert "리뷰 슬랙 보내기" in logs and "secretTEST1234" not in logs and "secretTEST1234" not in res.get_data(as_text=True)
    Res.status_code = 404
    assert "거절" in auth_client.post(url, json={}).get_json()["error"]


def test_weee_review_count():
    from tabs.insights.weee import review_count
    assert review_count("# Reviews (1,506)\n\nBased on") == 1506
    assert review_count("[Back](x)\n\nReviews (0)") == 0 and review_count("no header") is None


def test_trends_csv_upload(auth_client):
    csv_text = "Category: All categories\n\nWeek,hangover: (United States),hangover drink: (United States)\n" \
               "2025-10-26,45,<1\n2025-11-02,90,3\n2025-11-09,40,2\n"
    res = auth_client.post("/api/insights/trends/upload", json={"csv": csv_text})
    assert res.status_code == 200 and res.get_json()["keywords"] == ["hangover", "hangover drink"]
    tr = auth_client.get("/api/insights/trends").get_json()
    hang = next(k for k in tr["keywords"] if k["keyword"] == "hangover")
    assert hang["points"][1] == {"date": "2025-11-02", "value": 90.0}
    drink = next(k for k in tr["keywords"] if k["keyword"] == "hangover drink")
    assert drink["points"][0]["value"] == 0.0
    assert auth_client.post("/api/insights/trends/upload", json={"csv": "hello,world\n1,2\n"}).status_code == 400


# ───────── AI (가짜 OpenAI) ─────────

def test_ai_rewrite_and_monthly_with_mocked_openai(app, auth_client, monkeypatch, fixed_today, sync_jobs):
    from tabs.insights import ai as insights_ai
    _seed_two_weeks(app)
    monkeypatch.setenv("OPENAI_API_KEY", "sk-test-openai-key-zzzz")
    seen = []

    def fake_openai(system, user, *, schema=None, schema_name="result", **kw):
        seen.append(schema_name)
        if schema_name == "weekly_actions":
            return {"items": [{"id": a["id"], "fact": "AI 사실 " + a["id"], "why": "AI 이유", "todo": "AI 할 일"}
                              for a in user["actions"]]}
        if schema_name == "review_topics":
            return {"items": [{"id": r["id"], "topics": ["가격"], "sentiment": "-", "summary_ko": "비싸요"} for r in user["reviews"]]}
        if schema_name == "monthly_review":
            return {"summary": ["한 줄", "두 줄", "세 줄", "네 줄"], "insights": [], "competitor_trends": ["Flyby 가격 인하"],
                    "changes": [], "personas": [{"name": "파티 다음 날 직장인", "oneLine": "아침 회의가 걱정", "needs": "간편함",
                                                 "message": "가볍게 챙기는 리추얼"}] * 4,
                    "swot": {"S": ["a"], "W": ["b"], "O": ["c"], "T": ["d"]}, "target": "21~34세", "positioning": "간편한 리추얼"}
        raise AssertionError(schema_name)

    monkeypatch.setattr(insights_ai, "openai_json", fake_openai)
    from tabs.insights import store
    with app.app_context():
        store.save_reviews([{"source": "weee", "product_key": "weee:X", "product_title": "X", "rating": 2.0,
                             "text": "too pricey for me"}], week_key(service.today()))
    res = auth_client.post("/api/insights/update", json={"items": ["reviews", "actions", "monthly"]})
    assert res.status_code == 202
    run = auth_client.get("/api/insights/update/status").get_json()["run"]
    assert run["status"] == "done", run["results"]
    assert seen == ["review_topics", "weekly_actions", "monthly_review"]
    week = auth_client.get("/api/insights/weekly").get_json()
    assert week["aiUsed"] and all(a["ai"] and a["fact"].startswith("AI 사실") for a in week["actions"])
    assert week["actions"][0]["base"]["fact"].startswith("경쟁 상품")  # 규칙 문장도 함께
    # 같은 데이터면 다시 부르지 않음 (저장된 문장)
    auth_client.post("/api/insights/update", json={"items": ["actions"]})
    assert seen.count("weekly_actions") == 1
    m = auth_client.get("/api/insights/monthly").get_json()
    assert m["ai"] and m["summary"] == ["한 줄", "두 줄", "세 줄"] and len(m["strategy"]["personas"]) == 3
    # 화면 'AI가 읽은 자료': 실행 때 넘긴 개수 그대로
    src = m["sources"]
    assert set(src) == {"shelf", "reviews", "questions", "expressions", "competitors", "trends", "lastMonth"}
    assert src["shelf"] > 0 and src["reviews"] >= 1 and src["lastMonth"] is False
    personas = auth_client.get("/api/insights/personas").get_json()
    assert personas[0] == {"name": "파티 다음 날 직장인", "oneLine": "아침 회의가 걱정", "needs": "간편함", "message": "가볍게 챙기는 리추얼"}
    rv = auth_client.get("/api/insights/data/reviews").get_json()
    assert rv["stats"]["ai"] == 1 and rv["items"][0]["summaryKo"] == "비싸요"


# ───────── 내려받기 ─────────

def test_reports_and_export(app, auth_client, fixed_today):
    from docx import Document
    from openpyxl import load_workbook
    _seed_two_weeks(app)
    res = auth_client.get("/api/insights/report/weekly.docx")
    assert res.status_code == 200 and "wordprocessingml" in res.headers["Content-Type"]
    assert "reboot-insights-weekly-2026-W43.docx" in res.headers["Content-Disposition"]
    text = "\n".join(p.text for p in Document(io.BytesIO(res.data)).paragraphs)
    assert "이번 주 액션" in text and "[긴급]" in text and "다가오는 시즌" in text and "Halloween" in text
    res = auth_client.get("/api/insights/report/monthly.docx")
    text = "\n".join(p.text for p in Document(io.BytesIO(res.data)).paragraphs)
    assert res.status_code == 200 and "월간 리포트" in text and "전략 업데이트가 없어요" in text
    res = auth_client.get("/api/insights/export.xlsx")
    assert res.status_code == 200 and "spreadsheetml" in res.headers["Content-Type"]
    wb = load_workbook(io.BytesIO(res.data))
    assert wb.sheetnames == ["진열대(주차별)", "구매자 리뷰", "VOC(유튜브·리뷰)", "검색 트렌드", "경쟁사 후보·점수", "경쟁사 분석", "이번 주 액션", "작업 로그"]
    assert wb["진열대(주차별)"].max_row == 1 + 4


# ───────── 파서 (수집기 단위) ─────────

def test_weee_markdown_parser():
    from tabs.insights.weee import parse_listing, parse_product_page
    items = parse_listing(WEEE_MD)
    assert [i["id"] for i in items] == ["12345", "67890", "11111"]
    assert items[0]["title"] == "Condition Hangover Drink 100ml x 10" and items[0]["price"] == 19.99 and items[0]["in_stock"] is True
    assert items[1]["in_stock"] is False and items[1]["price"] == 12.99
    page = parse_product_page("# Aldicom Hangover Jelly\n\n$12.99\n\nA refreshing jelly stick made with lemon and lime for your next good moment.\n",
                              "https://www.sayweee.com/en/product/Aldicom-Jelly/67890")
    assert page["id"] == "67890" and page["price"] == 12.99 and page["in_stock"] is True and "refreshing" in page["message"]


# 지금 Weee!: 카드 전체가 링크 [이미지 · 할인 · 가격 · 원가 · 단위 가격 · 상품명 · 배송](weee.com 주소)
WEEE_CARDS = """## 'hangover(27)'

[![weee_beverages_Oriental Raisin Tea Extract 10 ml*30 pack](https://img08.weeecdn.net/a.png)\\\\
\\\\
5% off\\\\
\\\\
$12.79\\\\
\\\\
$13.49\\\\
\\\\
$1.26/fl.oz\\\\
\\\\
Oriental Raisin Tea Extract 10 ml\\*30 pack\\\\
\\\\
![SNAP](https://img08.weeecdn.net/snap.png)\\\\
\\\\
SNAP](https://www.weee.com/en/product/Oriental-Raisin-Tea-Extract/112105?trace_id=1)

Add to cart

[6% off\\\\
\\\\
$13.99\\\\
\\\\
HOUSE Turmeric Hangover Granules 11 pack/ 1 box\\\\
\\\\
Sold out\\\\
\\\\
Ships from Sakura Charm](https://www.weee.com/en/product/HOUSE-Turmeric-Hangover-Granules-11-pack/2163335?trace_id=1)

Notify me
"""
WEEE_PAGE = """[Skip to content](https://www.weee.com/en/product/Easy-Tomorrow-Jelly-Stick/2081684#main-content)

$10 off on your first order

![Easy Tomorrow Jelly Stick Before/After Drink, 2 Sticks 1 case](https://img08.weeecdn.net/x.png)

$

7

99

Sold out

10 Sticks per Box (180g)

## Details

[Related](https://www.weee.com/en/product/Other/1) Sold out $3.00
"""


def test_weee_card_listing_and_split_price_page():
    from tabs.insights.weee import parse_listing, parse_product_page, relevant
    items = parse_listing(WEEE_CARDS)
    assert [i["id"] for i in items] == ["112105", "2163335"]
    assert items[0]["title"] == "Oriental Raisin Tea Extract 10 ml*30 pack" and items[0]["price"] == 12.79  # 할인가 (원가 · 단위 가격 아님)
    assert items[0]["in_stock"] is True and items[0]["link"].startswith("https://www.weee.com/")
    assert items[1]["price"] == 13.99 and items[1]["in_stock"] is False
    page = parse_product_page(WEEE_PAGE, "https://www.sayweee.com/en/product/Easy-Tomorrow-Jelly-Stick/2081684")
    assert page["price"] == 7.99 and page["in_stock"] is False and page["pack_text"].startswith("10 Sticks")  # '$10 off'는 가격 아님
    assert page["title"].startswith("Easy Tomorrow Jelly Stick")
    assert relevant("HOUSE Turmeric Hangover Granules") and relevant("Qone Jelly Zero Sugar") and not relevant("Coconut Jelly 4 oz")


def test_pellet_form_and_product_lines():
    from tabs.insights import lines
    from tabs.insights.competitors import rule_form, rule_form_of
    assert rule_form("Q.one Easy Tomorrow, Korean Herbal Pellets 10 Packs") == "pellet" and rule_form("상쾌환 20포") == "pellet"
    assert rule_form("환불 안내") == ""  # '환불'의 환은 제형 아님
    assert rule_form_of({"title": "Herbal Pellets 10 Packs", "message": "Do not chew"}) == "pellet"  # 상품명이 먼저
    assert lines.known_brand("K FOOD KOREAN FOOD", "Korean Food Easy Tomorrow Fresh Drink Mix")["name"] == "상쾌환"
    assert lines.known_brand("", "Qone Hangover Relief Jelly")["name"] == "상쾌환" and lines.known_brand("", "Flyby 30 Count") is None
    assert lines.stock_left("Only 12 left in stock - order soon.") == 12 and lines.stock_left("In Stock") is None

    def row(i, src, servings, ps, in_stock=True, left=None, form="pellet"):
        return {"id": i, "key": f"{src}:{i}", "source": src, "title": f"상쾌환 {servings}포", "servings": servings, "servingsSource": "parsed",
                "perServing": ps, "price": ps * servings, "in_stock": in_stock, "stockLeft": left, "own": False, "reviews": 10,
                "line": {"id": f"상쾌환|{form}", "brand": "상쾌환", "form": form, "label": "상쾌환 · 환", "confirmed": True}}
    rows = [row(1, "amazon", 10, 2.29), row(2, "amazon", 20, 2.20, left=13), row(3, "weee", 10, 2.5, in_stock=False),
            row(4, "weee", 20, 2.4, in_stock=False)]
    g = lines.build(rows)[0]
    assert g["channels"] == ["amazon", "weee"] and g["oosChannels"] == ["weee"]          # Weee!는 두 옵션 모두 품절
    assert [o["label"] for o in g["options"]] == ["10회분", "20회분"]                     # 용량이 달라도 한 제품
    assert g["options"][1]["cells"]["amazon"]["state"] == "low" and g["low"][0]["left"] == 13
    assert g["bulk"]["pct"] == -3.9 and not g["bulk"]["suspect"]
    rows[2]["in_stock"] = True
    g = lines.build(rows)[0]
    assert g["oosChannels"] == [] and g["partialOos"] == 1 and g["options"][0]["cells"]["weee"]["state"] == "in"


def test_amazon_search_parser_skips_ads():
    from tabs.insights.amazon import parse_search_item, stock_state
    item = parse_search_item({"asin": "b0comp0001", "title": " Flyby ", "price": "$29.99", "rating": "4.3", "reviews": "5,400"})
    assert item["asin"] == "B0COMP0001" and item["price"] == 29.99 and item["reviews"] == 5400 and item["title"] == "Flyby"
    assert stock_state("Currently unavailable.") is False and stock_state("Only 3 left in stock") is True and stock_state("") is None


# ───────── 경쟁사 분석 (판매 채널 · 칭찬/불만 · 이번 주 변화 · 웹 검색 보충) ─────────

YAMI_MD = """510 items

[![x](https://cdn.yamibuy.net/item/a.webp)](https://www.yami.com/us/en/p/morning-is-back-jelly/1026089171?bu_type=search)

[Morning Is Back](https://www.yami.com/us/en/b/morning-is-back/18336) [Blueberry Hangover Relief Jelly 30 Sticks](https://www.yami.com/us/en/p/morning-is-back-jelly/1026089171?bu_type=search "Blueberry")

30+ Sold

$18.99$19.99

Options: 2

[ALDICOM](https://www.yami.com/us/en/b/aldicom/24944) [Hangover Relief Liquid Supplement, 5 Sticks](https://www.yami.com/us/en/p/aldicom-a/1174080301?bu_type=search "A")

60+ Sold

$4.79$8.99

[ORIHIRO](https://www.yami.com/us/en/b/orihiro/663) [Collagen Beauty Drink \\* 10](https://www.yami.com/us/en/p/collagen/3030045071?x=1 "C")

$15.99

Options: 4

Sold by [TOKYO PLAZA](https://www.yami.com/us/en/s/tokyo-plaza/1)
"""


def test_item_form_from_amazon_product():
    from tabs.insights.amazon import _item_form
    from tabs.insights.competitors import rule_form_of
    pr = {"variants": [{"title": "Item Form", "items": [{"name": "Liquid"}, {"name": "Tablet", "selected": True}]}]}
    assert _item_form({}, pr) == "Tablet"
    assert _item_form({"product_details": {"item_form": "Powder"}}, {}) == "Powder"
    assert _item_form({"item_specifications": [{"name": "Item Form", "value": "Capsule"}]}, {}) == "Capsule"
    assert _item_form({}, {"variants": [{"title": "Size", "items": [{"name": "10 Count", "selected": True}]}]}) == ""
    x = {"title": "Blowfish for Hangovers - Fast Hangover Relief", "message": "Item Form: Tablet / Drink plenty of water"}
    assert rule_form_of(x) == "tablet"  # 상품명에 없으면 제품 선택 → 판매 문구('Drink')보다 앞
    assert rule_form_of(dict(x, title="Blowfish Liquid Shot")) == "liquid"  # 상품명이 먼저


def test_form_rules_and_yami_parser():
    from tabs.insights.competitors import rule_form
    from tabs.insights.yami import RELEVANT_RE, parse_search
    assert rule_form("Cheers Restore Hangover Liquid Shot 2oz") == "liquid"
    assert rule_form("Hangover Gummies 30 Count") == "gummy" and rule_form("Aldicom Hangover Jelly Stick") == "stick"  # 스틱 포장이면 스틱
    assert rule_form("Electrolyte Drink Mix Powder Packets") == "powder"  # 분말 음료는 분말
    assert rule_form("Flyby Recovery 30 Capsules") == "tablet" and rule_form("Hangover support") == ""
    assert rule_form("Vitamin C Effervescent 20 Count") == "tablet" and rule_form("Fizzy Hangover Drops") == "tablet"  # 발포정
    assert rule_form("HANGOVER CURE GINSENG AND PUERARIA 250g/10PCS") == "liquid" and rule_form("홍삼 스틱") == "stick"
    assert rule_form("Korean Red Ginseng Capsules") == "tablet"  # 다른 제형 단어가 있으면 그쪽
    assert rule_form("Alpitan Hangover Relief Powder 6 Sticks") == "stick" and rule_form("Hangover Relief Liquid, 5 Sticks") == "stick"
    assert rule_form("Q.one Herbal Pellets 10 Sticks") == "pellet"  # 환이 먼저
    items = parse_search(YAMI_MD)
    assert [i["id"] for i in items] == ["1026089171", "1174080301", "3030045071"]
    assert items[0]["brand"] == "Morning Is Back" and items[0]["price"] == 18.99 and items[0]["list_price"] == 19.99
    assert items[0]["sold"] == "30+" and items[1]["sold"] == "60+" and items[2]["sold"] == ""  # 'Sold by'는 판매량 아님
    assert items[2]["title"] == "Collagen Beauty Drink * 10" and not RELEVANT_RE.search(items[2]["title"])


def test_yami_collect_through_update(app, auth_client, monkeypatch, fixed_today, sync_jobs):
    monkeypatch.setenv("FIRECRAWL_API_KEY", "fc-test-key-zzzz")
    seen = []

    def fake_post(url, json=None, headers=None, timeout=90):
        seen.append(json["url"])
        return FakeRes(200, {"success": True, "data": {"markdown": YAMI_MD}})

    monkeypatch.setattr(http, "post", fake_post)
    assert auth_client.post("/api/insights/update", json={"items": ["yami"]}).status_code == 202
    run = auth_client.get("/api/insights/update/status").get_json()["run"]
    assert run["status"] == "done", run["results"]
    assert seen == ["https://www.yami.com/us/en/search?q=hangover", "https://www.yami.com/us/en/search?q=hangover%20jelly"]
    shelf = auth_client.get("/api/insights/data/shelf").get_json()["items"]
    yami = {x["key"]: x for x in shelf if x["source"] == "yami"}
    assert set(yami) == {"yami:1026089171", "yami:1174080301"}  # 숙취와 관계없는 상품은 뺌
    assert yami["yami:1026089171"]["sourceLabel"] == "Yami" and yami["yami:1026089171"]["servings"] == 30
    assert yami["yami:1174080301"]["own"] and "60+ Sold" in yami["yami:1174080301"]["availability"]


def _seed_competitors(app):
    from server.extensions import db
    from tabs.insights.models import Review, ShelfItem
    amazon = [  # (key, title, brand, queries, reviews, price, msg)
        ("amazon:B0SHOT0001", "Cheers Restore Hangover Liquid Shot 3 Count", "Cheers", "hangover liquid shot", 9000, 15.0,
         "DHM and electrolytes in a shot."),
        ("amazon:B0GUMMY001", "Flyby Hangover Gummies 30 Count", "Flyby", "hangover gummies", 5000, 30.0, "Herbal blend."),
        ("amazon:B0LIV00001", "Liquid I.V. Hydration Multiplier Drink Mix Powder", "Liquid I.V.", "electrolyte drink hangover",
         90000, 25.0, "Hydration."),
        ("amazon:B0OWN00001", "Aldicom Hangover Jelly Stick", "Aldicom", "korean hangover", 40, 20.0, ""),
    ]
    with app.app_context():
        for week in ("2026-W42", "2026-W43"):
            for key, title, brand, q, reviews, price, msg in amazon:
                if week == "2026-W42" and brand == "Flyby":
                    continue  # 이번 주 새 상품
                if week == "2026-W42" and brand == "Cheers":
                    price, msg = 18.0, "DHM shot."
                db.session.add(ShelfItem(week=week, source="amazon", key=key, title=title, brand=brand, queries=q, reviews=reviews,
                                         price=price, servings=parse_servings(title)[0], in_stock=True, message=msg,
                                         own=brand == "Aldicom", link=f"https://www.amazon.com/dp/{key.split(':')[1]}"))
        db.session.add(ShelfItem(week="2026-W43", source="yami", key="yami:1", title="Cheers Hangover Jelly 10 Sticks",
                                 brand="Cheers", queries="hangover", price=12.0, servings=10, in_stock=True,
                                 availability="30+ Sold", link="https://www.yami.com/us/en/p/cheers/1"))
        db.session.add(ShelfItem(week="2026-W43", source="yami", key="yami:2", title="Morning Is Back Hangover Drink 10 Bottles",
                                 brand="Morning Is Back", queries="hangover", price=14.99, servings=10, in_stock=True,
                                 link="https://www.yami.com/us/en/p/mib/2"))
        for i, (s, topics, text) in enumerate([("+", ["효과"], "Works fast"), ("+", ["효과"], "Great"), ("-", ["가격"], "Pricey")]):
            db.session.add(Review(uid=f"r{i}", source="amazon", product_key="amazon:B0SHOT0001", product_title="Cheers",
                                  brand="Cheers", rating=5 if s == "+" else 2, text=text, sentiment=s,
                                  topics=json.dumps(topics, ensure_ascii=False), summary_ko=f"요약 {text}", classified_by="rule"))
        db.session.commit()


def test_competitor_table_facts_and_weekly_changes(app, auth_client, fixed_today):
    from server.extensions import db
    from tabs.insights import competitors
    from tabs.insights.models import ProductPosition
    _seed_competitors(app)
    with app.app_context():  # 예전 버전이 저장한 제형 값(solid)이 남아 있어도 화면이 깨지지 않음
        db.session.add(ProductPosition(key="amazon:B0GUMMY001", brand="Flyby", form="solid", timing="post", classified_by="ai",
                                       sig=competitors._sig({"title": "Flyby Hangover Gummies 30 Count", "message": "Herbal blend."})))
        db.session.commit()
    res = auth_client.post("/api/insights/competitors/discover", json={})
    assert set(res.get_json()["added"]) == {"Cheers", "Flyby", "Liquid I.V.", "Morning Is Back"}
    d = auth_client.get("/api/insights/data/competitors").get_json()
    assert "matrix" not in d and "blueOcean" not in json.dumps(d)
    assert {c["brand"] for c in d["table"]} == {"Cheers", "Flyby", "Morning Is Back"}
    assert [c["brand"] for c in d["substitutes"]] == ["Liquid I.V."]  # 대체재 검색어로만 잡힘
    cheers = next(c for c in d["table"] if c["brand"] == "Cheers")
    assert [ch["label"] for ch in cheers["channels"]] == ["Amazon", "Yami"]
    assert cheers["channels"][0]["minPrice"] == 15.0 and cheers["channels"][1]["sold"] == "30+ Sold"
    assert cheers["perServing"] == 1.2 and {f["key"] for f in cheers["forms"]} == {"liquid", "stick"}  # Jelly Stick은 스틱
    fb = cheers["feedback"]
    assert fb["reviews"] == 3 and fb["praise"]["topic"] == "효과" and fb["praise"]["count"] == 2
    assert fb["complaint"]["topic"] == "가격" and fb["complaint"]["example"] == "요약 Pricey"
    assert d["own"]["own"] and d["own"]["channels"][0]["label"] == "Amazon"  # 알디콤 행
    kinds = {(c["brand"], c["kind"]) for c in d["changes"]}
    assert ("Cheers", "price") in kinds and ("Flyby", "new") in kinds and ("Cheers", "message") in kinds
    price = next(c for c in d["changes"] if c["kind"] == "price")
    assert price["label"] == "가격 인하" and "$18.00 → $15.00" in price["text"]
    assert price["group"] == "threat" and price["prevPrice"] == 18.0 and price["pct"] == -16.7 and price["impact"]
    msg = next(c for c in d["changes"] if c["kind"] == "message")  # 문구 비교: 새로 생긴 말 · 사라진 말
    assert msg["group"] == "watch" and msg["added"] == ["DHM and electrolytes in a shot"] and msg["removed"] == ["DHM shot"]
    assert next(c for c in d["changes"] if c["kind"] == "new")["group"] == "threat"
    assert not any(c["source"] == "Yami" for c in d["changes"])  # 지난주 Yami 수집이 없으면 비교 안 함
    assert d["changesDemo"] is False  # 실제 지난주가 있으면 예시를 쓰지 않음


def test_competitor_changes_demo_when_no_last_week(app, auth_client, fixed_today):
    """지난주 수집이 없으면 이번 주 실제 진열대로 만든 예시 지난주와 비교 — 화면 표시만, DB에는 안 남음."""
    from server.extensions import db
    from tabs.insights.models import ShelfItem
    _seed_competitors(app)
    with app.app_context():
        ShelfItem.query.filter(ShelfItem.week != "2026-W43").delete()
        db.session.commit()
        before = ShelfItem.query.count()
    d = auth_client.get("/api/insights/data/competitors").get_json()
    assert d["changesDemo"] is True and d["changes"]
    assert {c["kind"] for c in d["changes"]} <= {"new", "price", "oos", "back", "message"}
    again = auth_client.get("/api/insights/data/competitors").get_json()
    assert [(c["brand"], c["kind"]) for c in again["changes"]] == [(c["brand"], c["kind"]) for c in d["changes"]]  # 새로고침해도 같음
    with app.app_context():
        assert ShelfItem.query.count() == before  # 예시 지난주는 저장하지 않음


def test_competitor_research_fills_from_evidence_then_web(app, auth_client, monkeypatch, fixed_today, sync_jobs):
    from tabs.insights import ai as insights_ai
    from tabs.insights.competitors import NOT_FOUND
    _seed_competitors(app)
    monkeypatch.setenv("OPENAI_API_KEY", "sk-test-openai-key-zzzz")
    monkeypatch.setenv("FIRECRAWL_API_KEY", "fc-test-key-zzzz")
    posts, seen, version = [], [], {"n": 1}

    def fake_post(url, json=None, headers=None, timeout=90):
        posts.append(url)
        if url.endswith("/search"):
            q = json["query"].lower()
            if "official site" in q:
                hits = ([{"url": "https://www.amazon.com/cheers"}, {"url": "https://www.cheershealth.com/products"}]
                        if "cheers" in q else [{"url": "https://www.reddit.com/r/hangover"}])
            else:
                hits = [{"url": "https://bevnet.com/cheers-launch", "title": "Cheers launches new shot", "description": "Cheers expands to Target."},
                        {"url": "https://blog.example.com/best", "title": "10 Best Hangover Cures", "description": "Cheers and more"}]
            return FakeRes(200, {"success": True, "data": hits})
        if url.endswith("/scrape"):
            return FakeRes(200, {"success": True, "data": {"markdown": f"# Cheers\nScience-backed. Sold at CVS. v{version['n']}"}})
        raise AssertionError(url)

    def fake_openai(system, user, *, schema=None, schema_name="result", model=None, **kw):
        seen.append((schema_name, model))
        if schema_name == "product_forms":
            return {"items": [{"id": p["id"], "form": "other"} for p in user["products"]]}
        if schema_name == "competitor_analysis":
            ids = [e["id"] for e in user["evidence"]]
            return {"message": {"text": f"과학 기반 숙취 케어 v{version['n']}", "sources": [ids[0], "X9"]},
                    "ingredients": {"values": ["DHM"], "sources": []},  # 출처 없음 → 웹 검색
                    "retail": {"text": "CVS", "sources": ["W1"] if "W1" in ids else []},
                    "news": {"text": "Target 입점", "sources": ["N1"] if "N1" in ids else []},
                    "use": {"text": "젤리 제형의 간편함", "sources": []},
                    "avoid": {"text": "prevent 표현", "sources": []}}
        if schema_name == "competitor_web":
            out = {f: {"text": "", "values": [], "urls": []} for f in user["fields"]}
            if "ingredients" in out:
                out["ingredients"] = {"text": "", "values": ["DHM", "Milk Thistle"], "urls": ["https://www.cheershealth.com/science"]}
            return out
        raise AssertionError(schema_name)

    monkeypatch.setattr(http, "post", fake_post)
    monkeypatch.setattr(insights_ai, "openai_json", fake_openai)
    auth_client.post("/api/insights/competitors/discover", json={})
    assert auth_client.post("/api/insights/update", json={"items": ["competitors"]}).status_code == 202
    run = auth_client.get("/api/insights/update/status").get_json()["run"]
    assert run["status"] == "done", run["results"]
    assert "브랜드 분석 3개" in run["results"]["competitors"]["message"]
    assert ("competitor_web", insights_ai.WEB_MODEL) in seen  # 웹 검색은 검색 모델로
    d = auth_client.get("/api/insights/data/competitors").get_json()
    cheers = next(c for c in d["table"] if c["brand"] == "Cheers")
    a = cheers["analysis"]
    assert cheers["website"] == "https://www.cheershealth.com/"
    assert a["message"] == {"text": "과학 기반 숙취 케어 v1", "sources": ["P1"]}
    assert a["retail"]["sources"] == ["W1"] and a["news"]["sources"] == ["N1"]  # 'Best' 비교 글은 뉴스에서 뺌
    assert a["ingredients"]["values"] == ["DHM", "Milk Thistle"] and a["ingredients"]["sources"] == ["G1"]
    g1 = next(s for s in cheers["sources"] if s["id"] == "G1")
    assert g1["type"] == "AI 웹 검색" and g1["url"] == "https://www.cheershealth.com/science"
    mib = next(c for c in d["table"] if c["brand"] == "Morning Is Back")
    assert mib["analysis"]["retail"]["text"] == NOT_FOUND and mib["analysis"]["news"]["text"] == NOT_FOUND  # 지어내지 않음
    assert mib["analysis"]["use"]["text"] == "젤리 제형의 간편함"  # 판단 칸은 AI가 채움
    # 다시 조사 (홈페이지가 바뀜) → 바뀐 항목 기록, 홈페이지 찾기는 다시 안 함
    version["n"] = 2
    n_site = sum(1 for u in posts if u.endswith("/search"))
    assert auth_client.post(f"/api/insights/competitors/{cheers['id']}/research", json={}).status_code == 202
    assert sum(1 for u in posts if u.endswith("/search")) == n_site + 1  # 뉴스 검색 1회만
    cheers = next(c for c in auth_client.get("/api/insights/data/competitors").get_json()["table"] if c["brand"] == "Cheers")
    assert cheers["analysis"]["message"]["text"] == "과학 기반 숙취 케어 v2" and cheers["history"][0]["changed"] == ["핵심 메시지"]
    csv_text = auth_client.get("/api/insights/competitors.csv").data.decode("utf-8-sig")
    assert csv_text.splitlines()[0].startswith("브랜드,구분,상태,판매 채널 (최저가)") and "Amazon $15.00 · Yami $12.00" in csv_text
    rows = json.loads(auth_client.get("/api/insights/competitors.json").data)
    assert {r["brand"] for r in rows} == {"Cheers", "Flyby", "Liquid I.V.", "Morning Is Back"}
    assert "fc-test-key" not in csv_text and "sk-test" not in json.dumps(rows)


def test_competitor_edit_archive_and_website(app, auth_client, fixed_today):
    _seed_competitors(app)
    auth_client.post("/api/insights/competitors/discover", json={})
    d = auth_client.get("/api/insights/data/competitors").get_json()
    cid = next(c["id"] for c in d["table"] if c["brand"] == "Flyby")
    assert auth_client.patch(f"/api/insights/competitors/{cid}", json={"website": "flyby"}).status_code == 400
    out = auth_client.patch(f"/api/insights/competitors/{cid}", json={"website": "https://flybyrecovery.com/", "status": "archived"}).get_json()
    assert out["status"] == "archived" and out["website"] == "https://flybyrecovery.com/"
    d = auth_client.get("/api/insights/data/competitors").get_json()
    assert "Flyby" not in {c["brand"] for c in d["table"]} and [c["brand"] for c in d["archived"]] == ["Flyby"]
    assert auth_client.patch(f"/api/insights/competitors/{cid}", json={"status": "gone"}).status_code == 400
    assert auth_client.patch("/api/insights/competitors/9999", json={"status": "active"}).status_code == 404
    assert auth_client.post("/api/insights/competitors/9999/research", json={}).status_code == 404


# ───────── 채널 리뷰 (Olive Young · Yami) · CS ─────────

def _oy_unit(stars, text, date):
    full = '<div class="icon-star coral-50 left filled"></div><div class="icon-star coral-50 right filled"></div>' * stars
    gray = '<div class="icon-star gray-50 left filled"></div>'  # 항목별 별점 (전체 별점에 안 셈)
    return (f'<div class="product-review-unit"><div class="product-review-unit-main notranslate"><div class="review-star-rating">{full}</div>'
            f'<button>Report</button><div class="review-unit-option"><span>Aldicom 4 Sticks</span></div><ul><li>{gray}</li></ul>'
            f'<div class="review-unit-cont-comment -flip notranslate">{text}</div></div>'
            f'<span class="review-write-info-date notranslate">{date}</span><span>by. ab*****</span></div>')


def _oy_page(units, total):
    return (f"<html><h3>Global Reviews ({total})</h3><h3>Korean Reviews (800)</h3><div>4.6</div><div><b>{total}</b> reviews</div>"
            + "".join(units) + '<button class="review-list-more-btn">More (1/2)</button></html>')


def test_product_line_ignores_pack_size():
    from tabs.insights.channel_reviews import product_line
    assert product_line("Aldicom Hangover Helper 15g 4 Sticks") == "알디콤 A"
    assert product_line("Aldicom A Hangover Helper 15ml 1 Stick") == "알디콤 A"
    assert product_line("Aldicom V Hangover Helper 20ml 4 Sticks (4-day supply)") == "알디콤 V"
    assert product_line("Aldicom Plus Hangover Helper 21ml*4 Sticks") == "알디콤 Plus"
    assert product_line("ALDICOM Plus (4 boxs) 5 pack*4 boxes") == "알디콤 Plus"
    assert product_line("Aldicom Boost Daily Energy Gel Green Zest 30g * 8 Sticks") == "부스트 그린"
    assert product_line("Aldicom Boost Daily Energy Gel Pink Citrus 30g") == "부스트 핑크"


def test_oliveyoung_parser_reads_stars_text_date_without_names():
    from tabs.insights.channel_reviews import parse_oliveyoung
    p = parse_oliveyoung(_oy_page([_oy_unit(5, "Works great!<br>Buying again &amp; again", "2026/10/14"),
                                   _oy_unit(2, "The taste is too strong", "2026/09/01")], 12))
    assert p["total"] == 12 and p["avg"] == 4.6
    assert [r["rating"] for r in p["reviews"]] == [5.0, 2.0]
    assert p["reviews"][0]["text"] == "Works great! Buying again & again" and p["reviews"][1]["date"] == "2026/09/01"
    assert "ab***" not in json.dumps(p)


def test_oliveyoung_without_english_reviews_reads_nothing(monkeypatch):
    """영어 리뷰가 0개면 사이트가 한국 리뷰를 보여 줌 → 읽지 않음."""
    from tabs.insights import channel_reviews
    korean_shown = _oy_page([_oy_unit(5, "Translated Korean review", "2026/10/01")], 0)
    monkeypatch.setattr(channel_reviews, "firecrawl", lambda key, payload: {"rawHtml": korean_shown, "metadata": {"title": "Boost | OLIVE YOUNG"}})
    res = channel_reviews._oliveyoung("k", "https://global.oliveyoung.com/product/detail?prdtNo=GA1")
    assert res["reviews"] == [] and res["total"] == 0 and res["avg"] is None and res["calls"] == 1


def test_channel_reviews_amazon_marks_shelf_review_as_own(app, monkeypatch):
    """진열대(Amazon) 수집이 같은 리뷰를 경쟁사로 먼저 넣었어도 채널 리뷰가 자사로 바꿈 · SerpApi 키가 없으면 Amazon은 건너뜀."""
    from server.extensions import db
    from tabs.insights import channel_reviews, store
    from tabs.insights.models import Review
    monkeypatch.setattr(channel_reviews, "load_sources", lambda: [{"channel": "amazon", "asin": "B0DTYS4BP4"}])
    row = {"source": "amazon", "product_key": "amazon:B0DTYS4BP4", "product_title": "ALDICOM A", "brand": "", "own": False,
           "rating": 5.0, "title": "Works", "text": "No hangover", "date_text": "June 9, 2025"}
    monkeypatch.setattr(channel_reviews.amazon, "product", lambda key, asin: {
        "title": "ALDICOM A", "rating": 3.9, "reviews": 28, "review_rows": [dict(row, date_text=row["date_text"])]})
    with app.app_context():
        store.save_reviews([row], "2026-W43")
        with pytest.raises(http.CollectError):  # 키 없음 → Amazon만 있으면 읽은 게 없음
            channel_reviews.collect("fc", "2026-W43", None)
        res = channel_reviews.collect("fc", "2026-W43", "serp")
        assert res["reviews"] == 0 and res["serpapi"] == 1
        assert [r.own for r in Review.query.all()] == [True]
        db.session.remove()


def test_channel_reviews_collect_cs_and_translate(app, auth_client, monkeypatch, fixed_today, sync_jobs):
    from tabs.insights import ai as insights_ai
    from tabs.insights import channel_reviews
    monkeypatch.setenv("FIRECRAWL_API_KEY", "fc-test-key-zzzz")
    monkeypatch.setenv("SERPAPI_KEY", "serp-test-key-zzzz")
    monkeypatch.setattr(channel_reviews, "load_sources", lambda: [
        {"channel": "oliveyoung", "url": "https://global.oliveyoung.com/product/detail?prdtNo=GA240222422"},
        {"channel": "amazon", "asin": "B0DTYS4BP4"},
        {"channel": "yami", "url": "https://www.yami.com/us/en/p/aldicom-a/1174080301"}])
    first = [_oy_unit(5, f"Great hangover helper {i}", "2026/10/1%d" % i) for i in range(10)]
    more = first + [_oy_unit(1, "Box arrived crushed and one stick was leaking", "2026/08/01"),
                    _oy_unit(4, "Nice taste", "2026/07/01")]
    posts = []

    def fake_post(url, json=None, headers=None, timeout=90):
        posts.append(dict(json or {}))
        meta = {"title": "Aldicom Hangover Helper 15g 4 Sticks | OLIVE YOUNG Global"}
        if "oliveyoung" in json["url"]:
            raw = _oy_page(more if json.get("actions") else first, 12)
        else:
            raw, meta = '<section id="pdp-comment-section" class="commentSection_containerNoComment__x">Have your say</section>', {"title": "ALDICOM | Yami"}
        return FakeRes(200, {"success": True, "data": {"rawHtml": raw, "metadata": meta}})

    def fake_get(url, params=None, headers=None, timeout=25):
        assert params["engine"] == "amazon_product" and params["asin"] == "B0DTYS4BP4"
        return FakeRes(200, {"product_results": {"title": "ALDICOM A | Before Drinking, Travel Size Liquid 5-Pack", "rating": 3.9,
                                                 "reviews": 28},
                             "reviews_information": {"authors_reviews": [
                                 {"title": "Disgusting Taste", "text": "Such a disgusting taste, I almost vomited.", "rating": 1,
                                  "date": "August 6, 2025"},
                                 {"title": "Made Me a Tank", "text": "No hangover in the morning!!", "rating": 5, "date": "October 6, 2026"}]}})

    monkeypatch.setattr(http, "post", fake_post)
    monkeypatch.setattr(http, "get", fake_get)
    plan = {i["id"]: i for i in auth_client.get("/api/insights/update/plan").get_json()["items"]}
    assert plan["channel_reviews"]["ready"] and plan["channel_reviews"]["cost"] == {"firecrawl": 3, "serpapi": 1}
    assert auth_client.post("/api/insights/update", json={"items": ["channel_reviews"]}).status_code == 202
    run = auth_client.get("/api/insights/update/status").get_json()["run"]
    assert run["status"] == "done", run["results"]
    clicks = [a for a in posts[1]["actions"] if a["type"] == "executeJavascript"]
    assert len(clicks) == 1 and ".review-list-more-btn" in clicks[0]["script"]  # 12개 → 더 보기 1번 (한국 탭은 안 누름)

    d = auth_client.get("/api/insights/data/reviews").get_json()
    assert len(d["cs"]) == 14 and d["stats"]["oliveyoung"] == 12 and d["stats"]["amazon"] == 2 and d["stats"]["yami"] == 0
    assert {x["line"] for x in d["cs"]} == {"알디콤 A"} and {x["sourceLabel"] for x in d["cs"]} == {"Olive Young", "Amazon"}
    bad = next(x for x in d["cs"] if "crushed" in x["text"])
    assert bad["cs"] == "todo" and bad["csAuto"] and d["stats"]["csTodo"] == 2  # 상자 파손 + 아마존 1점
    amz = next(x for x in d["cs"] if x["source"] == "amazon" and x["rating"] == 1)
    assert amz["title"] == "Disgusting Taste" and amz["cs"] == "todo"
    ch = {c["channel"]: c for c in d["channels"]}
    assert ch["oliveyoung"]["total"] == 12 and ch["oliveyoung"]["got"] == 12 and ch["yami"]["total"] == 0
    assert ch["amazon"]["total"] == 28 and ch["amazon"]["avg"] == 3.9 and ch["amazon"]["got"] == 2

    # 다시 모아도 같은 리뷰는 한 번만
    auth_client.post("/api/insights/update", json={"items": ["channel_reviews"]})
    assert len(auth_client.get("/api/insights/data/reviews").get_json()["cs"]) == 14

    # CS 상태 · 메모
    res = auth_client.post(f"/api/insights/reviews/{bad['id']}/cs", json={"status": "done", "memo": "교환 안내함"})
    assert res.status_code == 200 and res.get_json()["cs"] == "done" and res.get_json()["csAuto"] is False
    assert auth_client.post(f"/api/insights/reviews/{bad['id']}/cs", json={"status": "nope"}).status_code == 400
    again = next(x for x in auth_client.get("/api/insights/data/reviews").get_json()["cs"] if x["id"] == bad["id"])
    assert again["cs"] == "done" and again["memo"] == "교환 안내함" and again["csBy"]

    # 번역: 키 없으면 안내, 있으면 한 번만 부르고 저장된 번역을 씀
    monkeypatch.delenv("OPENAI_API_KEY", raising=False)
    assert auth_client.post(f"/api/insights/reviews/{bad['id']}/translate", json={}).status_code == 400
    monkeypatch.setenv("OPENAI_API_KEY", "sk-test-openai-key-zzzz")
    seen = []

    def fake_openai(system, user, *, schema=None, schema_name="result", **kw):
        seen.append(schema_name)
        return {"text_ko": "상자가 찌그러져 왔고 한 포가 새고 있었어요"}

    monkeypatch.setattr(insights_ai, "openai_json", fake_openai)
    for _ in range(2):
        res = auth_client.post(f"/api/insights/reviews/{bad['id']}/translate", json={})
        assert res.status_code == 200 and res.get_json()["textKo"].startswith("상자가")
    assert seen == ["review_ko"]
    assert auth_client.post("/api/insights/reviews/999999/translate", json={}).status_code == 404


# ───────── 정확도 검증 · 퓨샷 ─────────

def test_comment_examples_are_consistent():
    """정답 100건: 예시 24건(근거 포함) · 채점 76건, 겹치지 않음 · 관련 없으면 질문 아님 · 주제는 질문일 때만."""
    from tabs.insights import ai as insights_ai, classify
    items = insights_ai.comment_examples()
    assert len(items) == 100 and len({x["no"] for x in items}) == 100
    ex = [x for x in items if x["use"] == "example"]
    assert len(ex) == 24 and all(x["reason"] for x in ex)
    assert {x["use"] for x in items} == {"example", "test"}
    for x in items:
        assert x["topic"] in set(classify.QUESTION_LABEL) | {"none"}, x["no"]
        assert not x["question"] or x["relevant"], x["no"]
        assert (x["topic"] != "none") == x["question"], x["no"]
        assert "@" not in x["text"] and "http" not in x["text"], x["no"]


def test_classify_comments_few_shot_prompt(monkeypatch):
    from tabs.insights import ai as insights_ai
    seen = []

    def fake_openai(system, user, *, schema=None, schema_name="result", **kw):
        seen.append((system, schema, kw["action"]))
        return {"items": [{"id": c["id"], "reason": "근거", "relevant": True, "question": False, "topic": "none"}
                          for c in user["comments"]]}

    monkeypatch.setattr(insights_ai, "openai_json", fake_openai)
    out = insights_ai.classify_comments([{"id": "c1", "text": "Where can I buy it?"}])
    insights_ai.classify_comments([{"id": "c1", "text": "Where can I buy it?"}], few_shot=False)
    assert out["c1"]["reason"] == "근거"
    (few, schema, action), (zero, _, zero_action) = seen
    sample = next(x for x in insights_ai.comment_examples() if x["use"] == "example")
    assert json.dumps(sample["text"], ensure_ascii=False) in few and sample["reason"] in few
    assert few.count("\n- ") == 24 and "예시" not in zero and "reason" in zero
    assert "예시 없이" in zero_action and "예시 없이" not in action
    props = schema["properties"]["items"]["items"]
    assert list(props["properties"])[:2] == ["id", "reason"] and "reason" in props["required"]  # 근거 먼저 (CoT)
    test_text = next(x["text"] for x in insights_ai.comment_examples() if x["use"] == "test")
    assert json.dumps(test_text, ensure_ascii=False) not in few  # 채점용은 예시에 없음


def test_accuracy_rule_only_then_ai(app, auth_client, monkeypatch, sync_jobs):
    from tabs.insights import accuracy, ai as insights_ai
    monkeypatch.delenv("OPENAI_API_KEY", raising=False)
    d = auth_client.get("/api/insights/data/accuracy").get_json()
    assert d["saved"] is None and len(d["examples"]) == 24 and d["testCounts"]["total"] == 76
    plan = {i["id"]: i for i in auth_client.get("/api/insights/update/plan").get_json()["items"]}
    assert plan["accuracy"]["manual"] and plan["accuracy"]["ready"] and plan["accuracy"]["cost"] == {"openai": 0}
    # AI 키 없음 → 규칙만
    assert auth_client.post("/api/insights/update", json={"items": ["accuracy"]}).status_code == 202
    run = auth_client.get("/api/insights/update/status").get_json()["run"]
    assert run["status"] == "done" and "규칙만" in run["results"]["accuracy"]["message"]
    s = auth_client.get("/api/insights/data/accuracy").get_json()["saved"]
    assert set(s["methods"]) == {"rule"} and s["winner"] == "rule" and s["tests"] == 76
    rule = s["methods"]["rule"]
    assert rule["n"] == 76 and 0 <= rule["exact"] <= rule["relevant"] <= 100
    assert all(w["text"] and "gold" in w and "pred" in w for w in rule["wrong"])

    # AI 키 있음 → 예시 포함은 정답 그대로, 예시 없이는 모두 '관련 없음'이라고 답하는 가짜 AI
    monkeypatch.setenv("OPENAI_API_KEY", "sk-test-openai-key-zzzz")
    gold = {str(x["no"]): x for x in insights_ai.comment_examples()}
    calls = []

    def fake_openai(system, user, *, schema=None, schema_name="result", **kw):
        few = "예시" in system
        calls.append(few)
        return {"items": [dict(id=c["id"], reason="", **({k: gold[c["id"]][k] for k in ("relevant", "question", "topic")} if few
                                                         else {"relevant": False, "question": False, "topic": "none"}))
                          for c in user["comments"]]}

    monkeypatch.setattr(insights_ai, "openai_json", fake_openai)
    plan = {i["id"]: i for i in auth_client.get("/api/insights/update/plan").get_json()["items"]}
    assert plan["accuracy"]["cost"] == {"openai": 4}
    auth_client.post("/api/insights/update", json={"items": ["accuracy"]})
    s = auth_client.get("/api/insights/data/accuracy").get_json()["saved"]
    assert calls == [False, False, True, True]  # 40건씩 2번 × 2방식
    assert s["winner"] == "few" and s["methods"]["few"]["exact"] == 100.0 and s["methods"]["few"]["wrong"] == []
    zero = s["methods"]["zero"]
    assert zero["relevant"] == round(26 * 100 / 76, 1)  # 정답이 '관련 없음'인 26건만 맞음
    with app.app_context():
        assert accuracy.score([{"no": 1, "relevant": True, "question": True, "topic": "buy"}],
                              [{"relevant": True, "question": False, "topic": "none"}])["exact"] == 0.0
